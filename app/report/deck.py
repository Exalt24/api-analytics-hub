"""Render a MonthlyReport as a PowerPoint deck, and read a deck back as text.

Every figure on a slide is a string `monthly.py` already formatted (the KPI rows carry
their own display text), so the deck does arithmetic nowhere. Chart series are the daily
values divided by the currency's minor unit for the axis only; the numbers a reader is
meant to quote are in the tables and titles.

Slides, in order: title, KPI summary, sales by day, orders by day, order value and buyers
with the best days, insights, data health. Slides are written for the client who reads
them, so they say "days with no data" and "data pull", not the pipeline's own words.

A day with orders is a point on the line, a day with no orders is a dot on the zero line,
and a day with no data is a break in the line with an X on the zero line. Those are three
different facts and the footnote of each chart says which are present.

`dump_deck` is the other half: it reads a saved file back through python-pptx so a test
(and a person without PowerPoint) can see exactly what a slide says, plus the file's
provenance line, which names the writer that produced the narrative.
"""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Mapping, Sequence

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.connectors.shopify import _ZERO_DECIMAL
from app.report.insights import Insights
from app.report.monthly import (
    KpiRow,
    MonthlyReport,
    MonthStats,
    claim_title,
    daily_series,
    days_text,
    fmt_day,
    format_money,
    insights_title,
    per_day_phrase,
    plural,
)
from app.report.paid import PaidReport, WINDOW_TEXT, paid_claim_title
from app.report.search import SearchReport, search_claim_title
from app.report.web import WebReport, web_claim_title

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x1B, 0x7F, 0x8C)
GREY = RGBColor(0x8A, 0x94, 0xA0)
ALERT = RGBColor(0xB4, 0x4A, 0x1E)
INK = RGBColor(0x2B, 0x2F, 0x36)
MUTED = RGBColor(0x5B, 0x64, 0x70)
BAND = RGBColor(0xF1, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
BODY_W = SLIDE_W - 2 * MARGIN
#: The smallest text on any slide, footnotes included.
FOOTNOTE_PT = 11

#: The line printed under the insights, whichever writer ran. Which writer it was is in the
#: file's provenance property and the text read back, not on the slide.
NARRATIVE_FOOTER = "Narrative written from the computed figures on the previous slides."


# ---------------------------------------------------------------- helpers

def _style_runs(frame, size: int, color: RGBColor, *, bold: bool = False, italic: bool = False) -> None:
    for para in frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size)
            run.font.name = FONT
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color


def _title(slide, text: str) -> None:
    box = slide.shapes.title
    box.left, box.top, box.width, box.height = MARGIN, Inches(0.3), BODY_W, Inches(1.1)
    box.text_frame.text = text
    box.text_frame.word_wrap = True
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for para in box.text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT
    _style_runs(box.text_frame, 28 if len(text) <= 58 else 24, NAVY, bold=True)


def _text(slide, left, top, width, height, lines: Sequence[str], size: int, color: RGBColor = INK,
          *, italic: bool = False, space_after: int = 0):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        para.space_after = Pt(space_after)
    _style_runs(frame, size, color, italic=italic)
    return box


def _new_slide(prs: Presentation, title: str, *, title_slide: bool = False):
    slide = prs.slides.add_slide(prs.slide_layouts[0 if title_slide else 5])
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx != 0 and not title_slide:
            ph._element.getparent().remove(ph._element)
    if not title_slide:
        _title(slide, title)
    return slide


def _footer(slide, lines: Sequence[str]) -> None:
    _text(slide, MARGIN, Inches(6.55), BODY_W, Inches(0.8), list(lines), FOOTNOTE_PT, MUTED)


def _table(slide, top, rows: Sequence[Sequence[str]], widths: Sequence[float], *, size: int = 16,
           row_h: float = 0.5, left_cols: int = 1, width=BODY_W, left=MARGIN):
    """A header row in navy, banded body rows, numbers right aligned."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, Inches(row_h * len(rows)))
    table = shape.table
    total = sum(widths)
    for i, w in enumerate(widths):
        table.columns[i].width = int(width * w / total)
    for r, row in enumerate(rows):
        table.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (BAND if r % 2 == 0 else WHITE)
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.LEFT if c < left_cols else PP_ALIGN.RIGHT
            _style_runs(cell.text_frame, size, WHITE if r == 0 else INK, bold=(r == 0))
    return table


def _units(currency: str | None) -> int:
    return 1 if currency and currency.upper() in _ZERO_DECIMAL else 100


def _major(value: int | None, currency: str | None) -> float | None:
    """Chart axis value only. The slide text never uses this."""
    return None if value is None else float(Decimal(value) / _units(currency))


def _unit_name(currency: str | None) -> str:
    return "$" if (currency or "").upper() == "USD" else (currency or "").upper()


def _note_text(report: MonthlyReport, row: KpiRow) -> str:
    """The Note cell: what is behind a figure, naming the month, never a bare "partial data"."""
    cur, prev = report.current, report.previous
    parts = []
    for note in row.notes:
        if note == "partial data":
            if cur.missing_days:
                parts.append("%s partial: %s missing" % (cur.short, plural(len(cur.missing_days), "day")))
            if not cur.complete:
                parts.append("%s partial: %d of %d days" % (cur.short, cur.days_expected, cur.days_in_month))
            if prev is not None and prev.missing_days:
                parts.append("%s partial: %s missing" % (prev.short, plural(len(prev.missing_days), "day")))
            if row.per_day is not None:
                parts.append("per reported day %s" % per_day_phrase(row.per_day))
        elif note == "zero baseline":
            parts.append("%s was zero" % prev.short)
        elif note == "no previous month":
            parts.append("no earlier month stored")
        elif note == "no previous value":
            parts.append("no %s figure" % prev.short)
        elif note == "not available":
            parts.append("not reported")
    return "; ".join(parts)


def _per_day_definition(report: MonthlyReport) -> list[str]:
    if not any(k.per_day is not None for k in report.kpis):
        return []
    return ["Per reported day divides a total by the days with orders plus the days with no orders, "
            "so months with different numbers of days with data can be compared."]


def _definitions(report: MonthlyReport) -> str:
    prev = report.previous
    change = ("Change %% is the change over %s's figure, n/a when that was zero. " % prev.short) if prev else ""
    return change + ("Buyers per active day is the average of each day's distinct buyers, over days with orders. "
                     "Average order value is gross sales divided by orders.")


# ---------------------------------------------------------------- charts

def _daily_chart(slide, report: MonthlyReport, metric: str, series_name: str, axis_title: str,
                 number_format: str, alt: str, *, compare: bool, top, height, left=MARGIN, width=BODY_W,
                 source: tuple[MonthStats, MonthStats | None] | None = None, currency: str | None = None,
                 legend: bool = True) -> None:
    """A line of daily values: a dot on zero for a day with no orders, a break and an X for no data.

    `source` draws another source's days (the website's, the ads', the search's) instead of the store's;
    `currency` is the money axis's currency when that source's is not the store's.
    """
    cc = currency if currency is not None else report.currency
    cur, prev = (source[0], source[1] if compare else None) if source else (
        report.current, report.previous if compare else None)
    days = max(cur.days_in_month, prev.days_in_month if prev else 0)
    to_axis = (lambda v: _major(v, cc)) if metric in ("net_sales", "ad_spend") else (lambda v: None if v is None else float(v))
    points = daily_series(cur, metric)
    values = [to_axis(v) for _, v in points]
    series = [(series_name, values, ACCENT)]
    if prev is not None:
        series.append((prev.name, [to_axis(v) for _, v in daily_series(prev, metric)], GREY))
    #: (series index, point indexes, colour): a dot on the zero line for a day with no orders, on every line
    dot_specs = [(0, [i for i, (d, v) in enumerate(points) if v is not None and d in cur.quiet_days], ACCENT)]
    if prev is not None:
        prev_points = daily_series(prev, metric)
        dot_specs.append((1, [i for i, (d, v) in enumerate(prev_points) if v is not None and d in prev.quiet_days], GREY))
    for stats, pts, color in ((cur, points, ALERT), (prev, prev_points if prev is not None else [], GREY)):
        if stats is not None and stats.missing_days:
            name = "Day with no data" if stats is cur else "%s day with no data" % stats.short
            series.append((name, [0.0 if v is None else None for _, v in pts], color))

    data = CategoryChartData()
    data.categories = [str(d) for d in range(1, days + 1)]
    for name, vals, _ in series:
        data.add_series(name, vals)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.LINE, left, top, width, height, data)
    gframe._element.nvGraphicFramePr.cNvPr.set("descr", alt)
    chart = gframe.chart
    chart.has_title = False
    chart.font.size = Pt(12)
    chart.font.name = FONT
    chart.font.color.rgb = INK
    chart.has_legend = len(series) > 1 and legend
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    axis = chart.value_axis
    axis.has_major_gridlines = True
    axis.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xDE, 0xE4)
    axis.tick_labels.number_format = number_format
    axis.tick_labels.number_format_is_linked = False
    axis.format.line.fill.background()
    axis.has_title = True
    axis.axis_title.text_frame.text = axis_title
    _style_runs(axis.axis_title.text_frame, 12, MUTED)
    chart.category_axis.format.line.color.rgb = GREY
    for plot_series, (name, _, color) in zip(chart.plots[0].series, series):
        plot_series.smooth = False
        if name.lower().endswith("day with no data"):
            plot_series.format.line.fill.background()
            plot_series.marker.style = XL_MARKER_STYLE.X
            plot_series.marker.size = 10
            plot_series.marker.format.line.color.rgb = color
            continue
        plot_series.format.line.color.rgb = color
        plot_series.format.line.width = Pt(2.5)
        plot_series.marker.style = XL_MARKER_STYLE.NONE
    for s_index, idxs, dot_color in dot_specs:
        for i in idxs:
            marker = chart.plots[0].series[s_index].points[i].marker
            marker.style = XL_MARKER_STYLE.CIRCLE
            marker.size = 9
            marker.format.fill.solid()
            marker.format.fill.fore_color.rgb = dot_color
            marker.format.line.color.rgb = dot_color


def _chart_notes(stats: MonthStats, what: str, prev: MonthStats | None = None) -> list[str]:
    lines = ["Day of the month across, %s up." % what]
    if stats.quiet_days or (prev is not None and prev.quiet_days):
        lines.append("A dot on the zero line is a day with no orders, on either line.")
    if stats.missing_days:
        lines.append("A break in the line with an X on the zero line is a day with no data (%s). %d of %d days are plotted."
                     % (days_text(stats.missing_days), stats.reported_days, stats.days_expected))
    if prev is not None and prev.missing_days:
        lines.append("A break in the grey line with a grey X is a day with no data in %s (%s)."
                     % (prev.short, days_text(prev.missing_days)))
    return lines


# ---------------------------------------------------------------- slides

def _title_slide(prs, report: MonthlyReport, data_note: str) -> None:
    slide = _new_slide(prs, "", title_slide=True)
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = MARGIN, Inches(2.2), BODY_W, Inches(1.4)
    title.text_frame.text = "Monthly performance report"
    for para in title.text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT
    _style_runs(title.text_frame, 44, NAVY, bold=True)
    sub = slide.placeholders[1]
    sub.left, sub.top, sub.width, sub.height = MARGIN, Inches(3.7), BODY_W, Inches(1.0)
    sub.text_frame.text = "%s, %s" % (report.label, report.current.name) if report.label else report.current.name
    for para in sub.text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT
    _style_runs(sub.text_frame, 24, ACCENT)
    if data_note:
        _text(slide, MARGIN, Inches(6.5), BODY_W, Inches(0.6), [data_note], 14, MUTED, italic=True)


def _kpi_slide(prs, report: MonthlyReport) -> None:
    cur, prev = report.current, report.previous
    slide = _new_slide(prs, claim_title(report, "net_sales"))
    head = ["Metric", cur.name, prev.name if prev else "Previous month", "Change", "Change %", "Note"]
    rows = [head] + [[k.label, k.current_display, k.previous_display, k.change_display, k.pct_display,
                      _note_text(report, k)] for k in report.kpis]
    extra = _extra_kpi_rows(report)
    if extra:
        # The summary table carries a headline row or two for each added source, so it is tighter than the store's alone.
        rows += extra
        _table(slide, Inches(1.45), rows, [2.2, 1.6, 1.6, 1.4, 1.0, 5.8], size=12, row_h=min(0.58, 4.6 / len(rows)))
    else:
        _table(slide, Inches(1.6), rows, [3.0, 2.0, 2.0, 1.8, 1.3, 3.7], size=14, row_h=0.58)
    _footer(slide, ["Source: the store's daily order records."] + _extra_kpi_footer(report)
            + [_definitions(report)] + _per_day_definition(report))


def _sales_slide(prs, report: MonthlyReport) -> None:
    cur, cc = report.current, report.currency
    slide = _new_slide(prs, "Net sales by day in %s" % cur.short)
    unit = _unit_name(cc)
    prev = report.previous
    _daily_chart(
        slide, report, "net_sales", cur.name, "Net sales (%s)" % unit,
        '"$"#,##0' if unit == "$" else "#,##0",
        "Line chart of net sales for each day of %s%s" % (cur.name, ", with %s as a grey line" % prev.name if prev else ""),
        compare=True, top=Inches(1.5), height=Inches(4.9),
    )
    _footer(slide, _chart_notes(cur, "net sales in %s" % ("dollars" if unit == "$" else unit), prev))


def _orders_slide(prs, report: MonthlyReport) -> None:
    cur = report.current
    slide = _new_slide(prs, "Orders by day in %s" % cur.short)
    _daily_chart(
        slide, report, "orders", cur.name, "Orders", "#,##0",
        "Line chart of orders for each day of %s" % cur.name,
        compare=False, top=Inches(1.5), height=Inches(4.9),
    )
    _footer(slide, _chart_notes(cur, "orders"))


def _value_slide(prs, report: MonthlyReport) -> None:
    cur, prev = report.current, report.previous
    slide = _new_slide(prs, "Order value, buyers and best days in %s" % cur.short)
    head = ["Metric", cur.name, prev.name if prev else "Previous month", "Change", "Change %", "Note"]
    rows = [head] + [[k.label, k.current_display, k.previous_display, k.change_display, k.pct_display,
                      _note_text(report, k)] for k in report.kpis if k.key in ("avg_order_value", "buyers_per_day")]
    _table(slide, Inches(1.5), rows, [3.0, 2.0, 2.0, 1.8, 1.3, 3.7], size=14, row_h=0.55)
    if report.top_days:
        _text(slide, MARGIN, Inches(3.35), BODY_W, Inches(0.4), ["Best days by net sales"], 18, NAVY)
        top = [["Rank", "Day", "Net sales"]] + [
            [str(i + 1), fmt_day(day), format_money(value, report.currency or "")]
            for i, (day, value) in enumerate(report.top_days)
        ]
        _table(slide, Inches(3.8), top, [1.0, 2.2, 2.8], size=15, row_h=0.5, left_cols=2, width=Inches(6.0))
    _footer(slide, [_definitions(report)])


def _insights_size(report: MonthlyReport, lines: Sequence[str]) -> tuple[int, int]:
    """The font size and the space after each line that keep the narrative on the slide.

    A store alone, or with a website, keeps the sizes it always had. With ads or search there are more lines,
    so the size is the largest of a few that fit when each line is estimated to wrap at about 1,750 / size
    characters across the slide (a measure, not a promise: the render is looked at).
    """
    if report.paid is None and report.search is None:
        size = 18 if report.web is None or len(lines) <= 8 else 15
        return size, 8 if size == 18 else 5
    for size in (18, 16, 14, 13, 12):
        per_line = int(1750 / size)
        wrapped = sum(max(1, -(-len(line) // per_line)) for line in lines)
        if wrapped * size * 1.2 + len(lines) * 4 <= 4.7 * 72:
            return size, 4
    return 12, 3


def _insights_slide(prs, report: MonthlyReport, insights: Insights) -> None:
    slide = _new_slide(prs, insights_title(report))
    # More lines when a second source is on the report, so the text is a little smaller to stay on the slide.
    size, space = _insights_size(report, insights.lines)
    _text(slide, MARGIN, Inches(1.6), BODY_W, Inches(4.8), list(insights.lines), size, INK, space_after=space)
    _footer(slide, [NARRATIVE_FOOTER])


def _health_title(report: MonthlyReport) -> str:
    cur = report.current
    if cur.missing_days:
        return "Data health: %s is missing %d of %d days" % (cur.short, len(cur.missing_days), cur.days_in_month)
    if not cur.complete:
        return "Data health: %d of %d days of %s are counted so far" % (cur.days_expected, cur.days_in_month, cur.short)
    return "Data health: every day of %s is covered" % cur.short


def _health_slide(prs, report: MonthlyReport) -> None:
    cur, prev = report.current, report.previous
    slide = _new_slide(prs, _health_title(report))

    def span(stats: MonthStats | None) -> str:
        if stats is None:
            return "n/a"
        return "%s to %s" % (fmt_day(stats.data_days[0]), fmt_day(stats.data_days[-1])) if stats.data_days else "none"

    def col(stats: MonthStats | None) -> list[str]:
        if stats is None:
            return ["n/a"] * 6
        return [str(stats.days_in_month), str(len(stats.data_days)), str(len(stats.quiet_days)),
                str(len(stats.missing_days)), str(stats.rows_used), span(stats)]

    labels = ["Days in the month", "Days with orders", "Days with no orders", "Days with no data",
              "Stored daily figures used", "First and last day with orders"]
    a, b = col(cur), col(prev)
    rows = [["Count", cur.name, prev.name if prev else "Previous month"]] + [[l, x, y] for l, x, y in zip(labels, a, b)]
    _table(slide, Inches(1.5), rows, [4.2, 3.2, 3.2], size=15, row_h=0.46)
    notes = ["Days with no data in %s: %s." % (cur.short, days_text(cur.missing_days))]
    if prev is not None:
        notes.append("Days with no data in %s: %s." % (prev.short, days_text(prev.missing_days)))
    else:
        notes.append("No earlier month is stored, so there is nothing to compare against.")
    if not cur.complete:
        notes.append("%s is not over: %d of %d days are counted." % (cur.short, cur.days_expected, cur.days_in_month))
    if cur.missing_days or not cur.complete or (prev is not None and prev.missing_days):
        notes.append("A month with days missing is understated, so its changes are marked partial.")
    _text(slide, MARGIN, Inches(4.95), BODY_W, Inches(1.6), notes, 15, INK, space_after=4)
    _footer(slide, ["Days with orders, days with no orders and days with no data add up to the days in the month.",
                    "A day with no orders is one where a data pull ran and found none. A day with no data is one no data pull covered."])


# ---------------------------------------------------------------- the website source

_HALF_W = (BODY_W - Inches(0.3)) // 2


def _web_chart_notes(report: MonthlyReport) -> list[str]:
    web, shop = report.web, report.current
    lines = ["Day of the month across. Left: website sessions. Right: orders. The grey line is the month before."]
    for name, stats, quiet in (("Sessions", web.current, "no sessions reported"), ("Orders", shop, "no orders")):
        bits = []
        if stats.quiet_days:
            bits.append("a dot on the zero line is a day with %s" % quiet)
        if stats.missing_days:
            bits.append("a break with an X is a day with no data (%s)" % days_text(stats.missing_days))
        if bits:
            lines.append("%s: %s." % (name, ", ".join(bits)))
    return lines


def _sessions_orders_slide(prs, report: MonthlyReport) -> None:
    web, cur = report.web, report.current
    slide = _new_slide(prs, "Sessions and orders by day in %s" % cur.short)
    top, height = Inches(1.5), Inches(4.5)
    _daily_chart(
        slide, report, "sessions", web.current.name, "Sessions", "#,##0",
        "Line chart of website sessions for each day of %s%s" % (
            cur.name, ", with %s as a grey line" % web.previous.name if web.previous else ""),
        compare=True, top=top, height=height, left=MARGIN, width=_HALF_W, source=(web.current, web.previous),
    )
    _daily_chart(
        slide, report, "orders", cur.name, "Orders", "#,##0",
        "Line chart of orders for each day of %s%s" % (
            cur.name, ", with %s as a grey line" % report.previous.name if report.previous else ""),
        compare=True, top=top, height=height, left=MARGIN + _HALF_W + Inches(0.3), width=_HALF_W,
    )
    _footer(slide, _web_chart_notes(report))


def _web_note_text(web: WebReport, row: KpiRow) -> str:
    """The Note cell of a website figure, naming the month and the days behind it."""
    cur, prev = web.current, web.previous
    parts = []
    for note in row.notes:
        if note == "partial data":
            if cur.missing_days:
                parts.append("%s: %s missing" % (cur.short, plural(len(cur.missing_days), "day")))
            if not cur.complete:
                parts.append("%s: %d of %d days" % (cur.short, cur.days_expected, cur.days_in_month))
            if prev is not None and prev.missing_days:
                parts.append("%s: %s missing" % (prev.short, plural(len(prev.missing_days), "day")))
            if row.per_day is not None:
                parts.append("per reported day %s" % per_day_phrase(row.per_day))
        elif note == "zero baseline":
            parts.append("%s was zero" % prev.short)
        elif note == "no previous month":
            parts.append("no earlier month stored")
        elif note == "no previous value":
            parts.append("no %s figure" % prev.short)
        elif note == "not available":
            parts.append("not reported")
    if row.key == "order_rate" and row.current is not None:
        parts.append("%d days in both" % web.joint_current)
    return "; ".join(parts)


def _web_per_day_definition(web: WebReport) -> list[str]:
    if not any(k.per_day is not None for k in web.kpis):
        return []
    return ["Per reported day divides a total by the days with traffic plus the days with no traffic."]


def _web_slide(prs, report: MonthlyReport) -> None:
    web = report.web
    wcur, wprev = web.current, web.previous
    slide = _new_slide(prs, web_claim_title(web, "sessions"))
    head = ["Metric", wcur.name, wprev.name if wprev else "Previous month", "Change", "Change %", "Note"]
    rows = [head] + [[k.label, k.current_display, k.previous_display, k.change_display, k.pct_display,
                      _web_note_text(web, k)] for k in web.kpis]
    _table(slide, Inches(1.4), rows, [2.6, 1.8, 1.8, 1.5, 1.2, 5.0], size=12, row_h=0.42)
    if web.weeks:
        _text(slide, MARGIN, Inches(4.1), BODY_W, Inches(0.35), ["Sessions against orders by week"], 16, NAVY)
        weeks = [["Week", "Days counted", "Sessions", "Orders", "Order rate"]] + [
            [w.label, w.days, w.sessions, w.orders, w.rate] for w in web.weeks]
        _table(slide, Inches(4.5), weeks, [2.6, 1.8, 1.6, 1.4, 1.6], size=12, row_h=0.32, left_cols=1,
               width=Inches(9.0))
    _footer(slide, [
        "Source: the website's Google Analytics 4 daily figures. Website days follow the property's own time zone, store days follow UTC.",
        "Users per active day is the average of each day's distinct users. Order rate is orders divided by sessions over the days both sources reported, not a Google Analytics conversion rate.",
    ] + _web_per_day_definition(web))


def _health_title_multi(report: MonthlyReport) -> str:
    cur, web = report.current, report.web.current
    if cur.missing_days and web.missing_days:
        return "Data health: %s is missing %s of store data and %s of website data" % (
            cur.short, plural(len(cur.missing_days), "day"), plural(len(web.missing_days), "day"))
    if cur.missing_days:
        return "Data health: %s is missing %d of %d days of store data" % (cur.short, len(cur.missing_days), cur.days_in_month)
    if web.missing_days:
        return "Data health: %s is missing %d of %d days of website data" % (cur.short, len(web.missing_days), web.days_in_month)
    if not cur.complete:
        return "Data health: %d of %d days of %s are counted so far" % (cur.days_expected, cur.days_in_month, cur.short)
    return "Data health: every day of %s is covered in both sources" % cur.short


def _health_slide_multi(prs, report: MonthlyReport) -> None:
    cur, prev, web = report.current, report.previous, report.web
    wcur, wprev = web.current, web.previous
    slide = _new_slide(prs, _health_title_multi(report))

    def col(stats: MonthStats | None) -> list[str]:
        if stats is None:
            return ["n/a"] * 5
        return [str(stats.days_in_month), str(len(stats.data_days)), str(len(stats.quiet_days)),
                str(len(stats.missing_days)), str(stats.rows_used)]

    labels = ["Days in the month", "Days with data", "Days with nothing to report", "Days with no data",
              "Stored daily figures used"]
    cols = [col(cur), col(prev), col(wcur), col(wprev)]
    head = ["Count", "Store, %s" % cur.short, "Store, %s" % (prev.short if prev else "earlier"),
            "Website, %s" % cur.short, "Website, %s" % (wprev.short if wprev else "earlier")]
    rows = [head] + [[labels[i]] + [c[i] for c in cols] for i in range(len(labels))]
    _table(slide, Inches(1.4), rows, [3.8, 2.0, 2.0, 2.0, 2.0], size=14, row_h=0.44)
    notes = ["Days with no data in %s, store: %s." % (cur.short, days_text(cur.missing_days)),
             "Days with no data in %s, website: %s." % (cur.short, days_text(wcur.missing_days))]
    if prev is not None:
        notes.append("Days with no data in %s, store: %s." % (prev.short, days_text(prev.missing_days)))
    if wprev is not None:
        notes.append("Days with no data in %s, website: %s." % (wprev.short, days_text(wprev.missing_days)))
    if not cur.complete:
        notes.append("%s is not over: %d of %d days are counted." % (cur.short, cur.days_expected, cur.days_in_month))
    partial = (cur.missing_days or wcur.missing_days or not cur.complete
               or (prev is not None and prev.missing_days) or (wprev is not None and wprev.missing_days))
    if partial:
        notes.append("A month with days missing is understated, so its changes are marked partial.")
    _text(slide, MARGIN, Inches(4.2), BODY_W, Inches(2.3), notes, 14, INK, space_after=3)
    _footer(slide, ["For each source, days with data, days with nothing to report and days with no data add up to the days in the month.",
                    "A day with nothing to report is one where a data pull ran and found none. A day with no data is one no data pull covered."])


# ---------------------------------------------------------------- the paid media and search sources

def _source_note_text(src, row: KpiRow, *, flagged: frozenset[str] = frozenset(), provisional: Sequence = (),
                      coverage: Mapping[str, int] | None = None) -> str:
    """The Note cell of an ad or search figure, naming the month and the days behind it.

    `flagged` are the figures that can still change in the `provisional` days, and `coverage` maps a figure
    that exists only on some days to the number of days that carry it.
    """
    cur, prev = src.current, src.previous
    parts = []
    for note in row.notes:
        if note == "partial data":
            if cur.missing_days:
                parts.append("%s: %s missing" % (cur.short, plural(len(cur.missing_days), "day")))
            if not cur.complete:
                parts.append("%s: %d of %d days" % (cur.short, cur.days_expected, cur.days_in_month))
            if prev is not None and prev.missing_days:
                parts.append("%s: %s missing" % (prev.short, plural(len(prev.missing_days), "day")))
            if row.per_day is not None:
                parts.append("per reported day %s" % per_day_phrase(row.per_day))
        elif note == "zero baseline":
            parts.append("%s was zero" % prev.short)
        elif note == "no previous month":
            parts.append("no earlier month stored")
        elif note == "no previous value":
            parts.append("no %s figure" % prev.short)
        elif note == "not available":
            parts.append("not reported")
    if coverage and row.key in coverage and row.current is not None and coverage[row.key] < cur.reported_days:
        parts.append("on %d of %d days" % (coverage[row.key], cur.reported_days))
    if provisional and row.key in flagged and row.current is not None:
        parts.append("last %s provisional" % plural(len(provisional), "day"))
    return "; ".join(parts)


_PAID_COVERED = frozenset({"ad_purchases", "ad_purchase_value", "roas"})
_PAID_HEADLINES = ("ad_spend", "roas")
_SEARCH_HEADLINES = ("search_clicks", "position")


def _paid_note(paid: PaidReport, row: KpiRow) -> str:
    cur = paid.current
    return _source_note_text(paid, row, flagged=_PAID_COVERED, provisional=paid.provisional,
                             coverage={"ad_purchases": cur.purchase_days, "ad_purchase_value": cur.value_days,
                                       "roas": cur.joint_days})


def _search_note(search: SearchReport, row: KpiRow) -> str:
    return _source_note_text(search, row, flagged=frozenset(k.key for k in search.kpis), provisional=search.provisional)


def _extra_kpi_rows(report: MonthlyReport) -> list[list[str]]:
    """A headline row or two for the ads and for search, in the summary table beside the store's figures."""
    rows: list[list[str]] = []
    if report.paid is not None:
        for key in _PAID_HEADLINES:
            k = report.paid.kpi(key)
            rows.append([k.label, k.current_display, k.previous_display, k.change_display, k.pct_display, _paid_note(report.paid, k)])
    if report.search is not None:
        for key in _SEARCH_HEADLINES:
            k = report.search.kpi(key)
            rows.append([k.label, k.current_display, k.previous_display, k.change_display, k.pct_display, _search_note(report.search, k)])
    return rows


def _extra_kpi_footer(report: MonthlyReport) -> list[str]:
    parts = []
    if report.paid is not None:
        parts.append("Ad rows: Meta, in the ad account's currency%s, purchases in a %s window." % (
            " (%s)" % report.paid.currency if report.paid.currency else "", WINDOW_TEXT))
    if report.search is not None:
        parts.append("Search rows: Search Console, Pacific time, a lower position is nearer the top.")
    return [" ".join(parts)] if parts else []


def _source_footer(src, noun_quiet: str, first: str, side: str, extra: Sequence[str] = ()) -> list[str]:
    """Footnote lines of a source slide: where the figures come from, which days are dots and gaps, and what may still change."""
    lines = [first]
    cur = src.current
    bits = []
    prev = src.previous
    if cur.quiet_days or (prev is not None and prev.quiet_days):
        bits.append("a dot on the zero line is a day with %s, on either line" % noun_quiet)
    if cur.missing_days:
        bits.append("a break with an X is a day with no data (%s)" % days_text(cur.missing_days))
    if prev is not None and prev.missing_days:
        bits.append("the grey line has no data on %s" % days_text(prev.missing_days))
    joined = "; ".join(bits)
    lines.append("Day of the month across. %s The grey line is the month before.%s" % (
        side, (" " + joined[:1].upper() + joined[1:] + ".") if bits else ""))
    return lines + list(extra)


def _paid_slide(prs, report: MonthlyReport) -> None:
    paid = report.paid
    cur, prev = paid.current, paid.previous
    slide = _new_slide(prs, paid_claim_title(paid, "ad_spend"))
    left_w = Inches(5.9)
    head = ["Metric", cur.short, prev.short if prev else "Before", "Change %"]
    rows = [head] + [[k.label, k.current_display, k.previous_display, k.pct_display] for k in paid.kpis]
    _table(slide, Inches(1.5), rows, [2.35, 1.65, 1.65, 1.05], size=12, row_h=0.46, width=left_w)
    x = MARGIN + left_w + Inches(0.3)
    w = BODY_W - left_w - Inches(0.3)
    unit = _unit_name(paid.currency)
    _daily_chart(
        slide, report, "ad_spend", cur.name, "Ad spend (%s)" % unit, '"$"#,##0' if unit == "$" else "#,##0",
        "Line chart of ad spend for each day of %s%s" % (cur.name, ", with %s as a grey line" % prev.name if prev else ""),
        compare=True, top=Inches(1.45), height=Inches(2.45), left=x, width=w, source=(cur, prev),
        currency=paid.currency)
    _daily_chart(
        slide, report, "ad_clicks", cur.name, "Clicks", "#,##0",
        "Line chart of ad clicks for each day of %s%s" % (cur.name, ", with %s as a grey line" % prev.name if prev else ""),
        compare=True, top=Inches(3.95), height=Inches(2.45), left=x, width=w, source=(cur, prev), legend=False)
    extra = []
    if paid.provisional:
        extra.append("Provisional days (a click can lead to a purchase after the day is over): %s." % days_text(paid.provisional))
    if any(k.per_day is not None for k in paid.kpis):
        extra.append("Per reported day divides a total by the days with ad delivery plus the days with none.")
    _footer(slide, _source_footer(
        paid, "no ad delivery",
        "Source: Meta Marketing API, in the ad account's currency%s and time zone; store days follow UTC. Purchases are Meta's own count in a %s window, not the store's orders."
        % (" (%s)" % paid.currency if paid.currency else "", WINDOW_TEXT),
        "Top chart: ad spend. Bottom chart: clicks.", extra))


def _top_table(slide, left, title: str, rows, first_col: str) -> None:
    _text(slide, left, Inches(3.9), Inches(3.85), Inches(0.35), [title], 15, NAVY)
    body = [[first_col, "Clicks", "Impr.", "Pos."]] + [[r.shown_key, format(r.clicks, ",d"), format(r.impressions, ",d"), r.position]
                                                       for r in rows]
    if len(body) == 1:
        body.append(["no rows stored", "", "", ""])
    _table(slide, Inches(4.25), body, [2.2, 0.6, 0.8, 0.55], size=11, row_h=0.33, left_cols=1, width=Inches(3.85), left=left)


def _search_slide(prs, report: MonthlyReport) -> None:
    search = report.search
    cur, prev = search.current, search.previous
    slide = _new_slide(prs, search_claim_title(search, "search_clicks"))
    top, height = Inches(1.4), Inches(2.45)
    _daily_chart(
        slide, report, "search_clicks", cur.name, "Search clicks", "#,##0",
        "Line chart of search clicks for each day of %s%s" % (cur.name, ", with %s as a grey line" % prev.name if prev else ""),
        compare=True, top=top, height=height, left=MARGIN, width=_HALF_W, source=(cur, prev))
    _daily_chart(
        slide, report, "search_impressions", cur.name, "Search impressions", "#,##0",
        "Line chart of search impressions for each day of %s%s" % (cur.name, ", with %s as a grey line" % prev.name if prev else ""),
        compare=True, top=top, height=height, left=MARGIN + _HALF_W + Inches(0.3), width=_HALF_W, source=(cur, prev), legend=False)
    _text(slide, MARGIN, Inches(3.9), Inches(4.0), Inches(0.35), ["Search figures"], 15, NAVY)
    head = ["Metric", cur.short[:3], prev.short[:3] if prev else "Before", "Chg %"]
    rows = [head] + [[k.label, k.current_display, k.previous_display, k.pct_display] for k in search.kpis]
    _table(slide, Inches(4.25), rows, [2.2, 0.85, 0.85, 0.8], size=11, row_h=0.33, width=Inches(4.0))
    _top_table(slide, MARGIN + Inches(4.15), "Top queries", search.top_queries, "Query")
    _top_table(slide, MARGIN + Inches(8.15), "Top pages", search.top_pages, "Page")
    extra = ["Top queries and pages add up each day's top rows over %d days, so a key outside a day's top rows is not counted: they are lower bounds, not totals."
             % search.top_row_days]
    if search.provisional:
        extra.append("Provisional days (Search Console is still collecting them): %s." % days_text(search.provisional))
    _footer(slide, _source_footer(
        search, "no search traffic",
        "Source: Google Search Console. Search days are Pacific time, store days follow UTC. A lower average position is nearer the top of the results.",
        "Left: search clicks. Right: search impressions.", extra))


def _sources(report: MonthlyReport) -> list[tuple[str, MonthStats, MonthStats | None]]:
    """(name, this month, the month before) of every source on the report, the store first."""
    out = [("Store", report.current, report.previous)]
    if report.web is not None:
        out.append(("Website", report.web.current, report.web.previous))
    if report.paid is not None:
        out.append(("Ads", report.paid.current, report.paid.previous))
    if report.search is not None:
        out.append(("Search", report.search.current, report.search.previous))
    return out


def _health_title_sources(report: MonthlyReport) -> str:
    cur = report.current
    sources = _sources(report)
    short = [name.lower() for name, c, _ in sources if c.missing_days]
    if short:
        names = short[0] if len(short) == 1 else ", ".join(short[:-1]) + " and " + short[-1]
        return "Data health: %s is missing days of %s data" % (cur.short, names)
    if not cur.complete:
        return "Data health: %d of %d days of %s are counted so far" % (cur.days_expected, cur.days_in_month, cur.short)
    return "Data health: every day of %s is covered in all %d sources" % (cur.short, len(sources))


def _health_slide_sources(prs, report: MonthlyReport) -> None:
    """One row for each source, then the days with no data for each source and month, and the days that may still change."""
    cur = report.current
    slide = _new_slide(prs, _health_title_sources(report))
    sources = _sources(report)
    provisional = {"Ads": report.paid.provisional if report.paid else [], "Search": report.search.provisional if report.search else []}
    head = ["Source, %s" % cur.short, "Days with data", "Days with nothing to report", "Days with no data",
            "Days that may still change", "Stored daily figures used"]
    rows = [head] + [[name, str(len(c.data_days)), str(len(c.quiet_days)), str(len(c.missing_days)),
                      str(len(provisional[name])) if name in provisional else "n/a", str(c.rows_used)] for name, c, _ in sources]
    _table(slide, Inches(1.4), rows, [2.3, 1.5, 2.2, 1.7, 2.0, 2.2], size=13, row_h=0.46)
    notes = []
    for name, c, p in sources:
        notes.append("Days with no data in %s, %s: %s." % (c.short, name.lower(), days_text(c.missing_days)))
        if p is not None and p.missing_days:
            notes.append("Days with no data in %s, %s: %s." % (p.short, name.lower(), days_text(p.missing_days)))
    for name in ("Ads", "Search"):
        if provisional[name]:
            notes.append("May still change, %s: %s." % (name.lower(), days_text(provisional[name])))
    if not cur.complete:
        notes.append("%s is not over: %d of %d days are counted." % (cur.short, cur.days_expected, cur.days_in_month))
    if any(c.missing_days or not c.complete or (p is not None and p.missing_days) for _, c, p in sources):
        notes.append("A month with days missing is understated, so its changes are marked partial.")
    _text(slide, MARGIN, Inches(1.65 + 0.46 * len(rows)), BODY_W, Inches(2.4), notes, 12, INK, space_after=2)
    _footer(slide, ["For each source, days with data, days with nothing to report and days with no data add up to the days in the month.",
                    "A day with nothing to report is one where a data pull ran and found none. A day with no data is one no data pull covered. "
                    "A day that may still change is one close to the report date: purchases and search figures are restated for a few days."])


# ---------------------------------------------------------------- public

def _set_app_properties(prs: Presentation) -> None:
    """Replace the library template's extended properties (a Mac PowerPoint, 0 slides, 4:3)."""
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        "<Application>api-analytics-hub</Application><PresentationFormat>Widescreen</PresentationFormat>"
        "<Slides>%d</Slides></Properties>" % len(prs.slides)
    )
    package = prs.part.package
    for part in package.iter_parts():
        if str(part.partname) == "/docProps/app.xml":
            part._blob = xml.encode("utf-8")
    # The template's thumbnail is a picture of a blank slide, not of this deck.
    for rid, rel in list(package._rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            package._rels.pop(rid)


def build_deck(report: MonthlyReport, insights: Insights, *, data_note: str = "") -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _title_slide(prs, report, data_note)
    _kpi_slide(prs, report)
    _sales_slide(prs, report)
    if report.web is None:
        _orders_slide(prs, report)
    else:
        # A second source needs its own figures and its own days, so the orders chart moves beside the
        # sessions chart and the website table is new.
        _sessions_orders_slide(prs, report)
    _value_slide(prs, report)
    if report.web is not None:
        _web_slide(prs, report)
    if report.paid is not None:
        _paid_slide(prs, report)
    if report.search is not None:
        _search_slide(prs, report)
    _insights_slide(prs, report, insights)
    if report.paid is not None or report.search is not None:
        _health_slide_sources(prs, report)
    elif report.web is not None:
        _health_slide_multi(prs, report)
    else:
        _health_slide(prs, report)
    props = prs.core_properties
    props.title = "Monthly performance report, %s" % report.current.name
    props.author = props.last_modified_by = "api-analytics-hub"
    props.keywords = props.subject = ""
    #: Which writer produced the narrative and why, kept off the slides.
    props.comments = insights.label
    # The library's template carries a 2013 timestamp, which reads as a stale file.
    props.created = props.modified = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    _set_app_properties(prs)
    return prs


def save_deck(report: MonthlyReport, insights: Insights, path: str | Path, *, data_note: str = "") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_deck(report, insights, data_note=data_note).save(str(path))
    return path


def dump_deck(source: str | Path | Presentation) -> str:
    """Every slide as text: the title, each text box, each table row and each chart series."""
    prs = Presentation(str(source)) if isinstance(source, (str, Path)) else source
    out: list[str] = []
    for n, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ""
        out.append("--- slide %d: %s" % (n, title))
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_chart:
                chart = shape.chart
                for s in chart.plots[0].series:
                    values = list(s.values)
                    out.append("chart series %r (%d points): %s" % (s.name, len(values), values))
            elif shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(cell.text for cell in row.cells))
            elif shape.has_text_frame and shape.text_frame.text.strip():
                out.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
    provenance = prs.core_properties.comments
    if provenance:
        out.append("--- provenance (file property, not on a slide): %s" % provenance)
    return "\n".join(out)
