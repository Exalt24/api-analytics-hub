"""Replay of the blind corpus for the Search Console connector and the search half of the report (output round 1).

tests/corpus/hub_search_corpus_blind_1.jsonl holds 40 cases (24 must-pass and 16 must-fail) written by a reviewer who read docs/GSC_SOURCES.md and the code. Each
carries `promise` (a quote from the spec), `why` (who is harmed on an ordinary day), `source` (real or invented) and `expected`. The adversary also recorded
`actual`, what the code did at 0e30795 when it wrote the case.

Every case is replayed here against the real code, offline: the connector through httpx.MockTransport with the raw Search Console bodies the case names, the
report from stored rows, and case 39 through the real sync runner into Postgres (skipped, never failed, when no Postgres is reachable). NOT_REPLAYABLE is empty on
purpose: all 40 are replayed. For a case the code agrees with the live outcome must equal `expected`. The four cases the old code failed (35 to 38, findings F2 and
F3: five page addresses and three long queries that printed as one text, and an average position that fell 6.0 to 9.0 and was called up 50.0%) have an
`expected` written as a description, so each has its own check below and the recorded `actual` must no longer be what the code does.

KNOWN_LIMIT with no corpus row: F9, a page address over 2,048 characters (an ad landing page with tracking strings) is refused by the top row length check,
scored 1x3 with no spec quote; it was kept out of the corpus by its author and its code is not changed.
"""
from __future__ import annotations

import ast
import asyncio
import json
import os
import sys
import tempfile
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import httpx  # noqa: E402

from app.connectors.gsc import GSCConnector  # noqa: E402
from app.report.deck import save_deck  # noqa: E402
from app.report.insights import template_insights  # noqa: E402
from app.report.monthly import build_report  # noqa: E402
from app.report.search import attach_search  # noqa: E402
from report_fixtures import day_rows  # noqa: E402

CORPUS = ROOT / "tests" / "corpus" / "hub_search_corpus_blind_1.jsonl"
CASES = [json.loads(line) for line in CORPUS.read_text(encoding="utf-8").splitlines() if line.strip()]

#: case number -> why it is not replayed offline. Empty: all 40 are replayed.
NOT_REPLAYABLE: dict[int, str] = {}

#: case number -> why the code is deliberately left as it is. Empty: the only rejected finding of this surface, F9, has no corpus row (see the docstring).
KNOWN_LIMITS: dict[int, str] = {}

#: the cases the old code failed and this change fixed; each is checked by its own predicate below
FIXED = {35, 36, 37, 38}

SITE = "sc-domain:acmeroofing.com"

# ---------------------------------------------------------------- the connector, the adversary's harness reproduced


def route_key(body: dict) -> str:
    dim = body["dimensions"][0]
    return "date" if dim == "date" else "%s:%s" % (dim, body["startDate"])


def make_transport(responses: dict, log: list):
    """responses: route -> list of {"status", "body" | "raw", "headers"}; a route's last entry repeats. Missing route: 200 {}."""
    used: dict[str, int] = {}

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        key = route_key(body)
        log.append({"path": request.url.raw_path.decode(), "auth": request.headers.get("authorization"), "body": body})
        seq = responses.get(key)
        if not seq:
            return httpx.Response(200, json={})
        i = used.get(key, 0)
        used[key] = i + 1
        r = seq[min(i, len(seq) - 1)]
        if "raw" in r:
            return httpx.Response(r.get("status", 200), content=r["raw"].encode(), headers=r.get("headers", {}))
        return httpx.Response(r.get("status", 200), content=json.dumps(r["body"]).encode(),
                              headers={"content-type": "application/json", **r.get("headers", {})})

    return httpx.MockTransport(handler)


async def run_connector(inp: dict):
    log: list = []
    slept: list = []

    async def sleep(s):
        slept.append(s)

    client = httpx.AsyncClient(transport=make_transport(inp.get("responses", {}), log))
    conn = GSCConnector(site_url=inp["site_url"], access_token=inp.get("token", "tok-adv-SECRET"), client=client,
                        top_n=inp.get("top_n", 10), sleep=sleep, rng=lambda: 0.0,
                        clock=lambda: datetime(2026, 11, 2, 12, tzinfo=timezone.utc))
    points, rows, cursors = [], [], []
    error = None
    try:
        async for page in conn.fetch(window_start=date.fromisoformat(inp["window"][0]), window_end=date.fromisoformat(inp["window"][1])):
            points += page.points
            rows += page.rows
            cursors.append(page.cursor)
    except Exception as exc:  # noqa: BLE001
        error = "%s: %s" % (type(exc).__name__, exc)
    finally:
        await client.aclose()
    return {"points": points, "rows": rows, "cursors": cursors, "error": error, "log": log, "slept": slept}


def connector(inp):
    try:
        return asyncio.run(run_connector(inp))
    except Exception as exc:  # noqa: BLE001  - constructor errors
        return {"error": "%s: %s" % (type(exc).__name__, exc), "points": [], "rows": [], "log": [], "slept": []}


def summary_fetch(res):
    if res["error"]:
        return "error " + res["error"].split(":")[0]
    by = {}
    for p in res["points"]:
        by.setdefault(p.observed_on.isoformat(), {})[p.metric_key.replace("search_", "")] = p.value_numeric
    prov = sorted({p.observed_on.isoformat() for p in res["points"] if (p.raw or {}).get("provisional")})
    return "days %s; provisional %s; top rows %d" % (json.dumps(by, sort_keys=True), prov, len(res["rows"]))


def path_and_body(res):
    if res["error"]:
        return "error " + res["error"]
    first = res["log"][0]
    b = first["body"]
    return "path %s; startDate %s endDate %s dataState %s type %s rowLimit %r startRow %r; auth %s" % (
        first["path"], b["startDate"], b["endDate"], b["dataState"], b["type"], b["rowLimit"], b["startRow"],
        "Bearer present" if (first["auth"] or "").startswith("Bearer ") else "missing")


def top_bodies(res):
    if res["error"]:
        return "error " + res["error"]
    tops = [e["body"] for e in res["log"] if e["body"]["dimensions"] != ["date"]]
    return "top calls %s; rows %s" % ([(b["dimensions"][0], b["startDate"], b["endDate"], b["rowLimit"]) for b in tops],
                                      [(r.observed_on.isoformat(), r.dimension, r.key, r.clicks, r.impressions, r.position_milli) for r in res["rows"]])


def err_kind(res):
    return ("error " + res["error"].split(":")[0] + "; waits %d" % len(res["slept"])) if res["error"] else "no error; " + summary_fetch(res)


def secret_leak(res):
    return "token in error: %s" % ("tok-adv-SECRET" in (res["error"] or ""))


def ctor(res):
    return ("error " + res["error"].split(":")[0]) if res["error"] else "built"


#: the summary each connector case is reported in, by case number (the corpus lists them in this order)
CONNECTOR_SUMMARY = {**{n: summary_fetch for n in range(0, 6)}, 6: path_and_body, 7: path_and_body, 8: top_bodies,
                      **{n: err_kind for n in (9, 10, 11, 12)}, 13: secret_leak, **{n: err_kind for n in range(14, 29)}, 29: ctor}

# ---------------------------------------------------------------- the report


def _store():
    rows = []
    for m, last in ((9, 30), (10, 31)):
        for d in range(1, last + 1):
            rows += day_rows(date(2026, m, d), 10, 10000, 0, 8)
    return rows


STORE = _store()
COV = [(date(2026, 9, 1), date(2026, 11, 1))]


def srows(spec):
    out = []
    for iso, (c, i, p) in spec.items():
        day = date.fromisoformat(iso)
        out += [{"observed_on": day, "metric_key": "search_clicks", "value_numeric": c, "currency": None},
                {"observed_on": day, "metric_key": "search_impressions", "value_numeric": i, "currency": None},
                {"observed_on": day, "metric_key": "search_position", "value_numeric": p, "currency": None}]
    return out


def search_report(inp):
    rows = srows(inp["rows"])
    top = [{"observed_on": date.fromisoformat(t[0]), "dimension": t[1], "key": t[2], "clicks": t[3], "impressions": t[4],
            "position_milli": t[5]} for t in inp.get("top", [])]
    cov = [(date.fromisoformat(a), date.fromisoformat(b)) for a, b in inp["covered"]]
    as_of = date.fromisoformat(inp["as_of"])
    report = build_report(STORE, 2026, 10, label="Store", covered=COV, as_of=as_of)
    attach_search(report, rows, top, label="Search", covered=cov, as_of=as_of)
    return report


def _kpis(r):
    return "; ".join("%s %s (prev %s, %s)" % (k.label, k.current_display, k.previous_display, k.pct_display) for k in r.search.kpis)


def _deck_top_pages(r):
    from pptx import Presentation
    with tempfile.TemporaryDirectory() as tmp:
        prs = Presentation(str(save_deck(r, template_insights(r), Path(tmp) / "d.pptx")))
        for slide in prs.slides:
            for s in slide.shapes:
                if s.has_table and s.table.rows[0].cells[0].text == "Page":
                    return "Top pages cells: %s" % [row.cells[0].text for row in list(s.table.rows)[1:]]
    return "no Top pages table"


REPORT_CALLS = {
    "kpis": _kpis,
    "position": lambda r: "Average position %s (exact %s)" % (r.search.kpi("position").current_display,
                                                              r.search.current.position_exact.quantize(Decimal("0.001"))),
    "days": lambda r: "data %d, quiet %s, missing %s" % (len(r.search.current.data_days), [d.isoformat() for d in r.search.current.quiet_days],
                                                         [d.isoformat() for d in r.search.current.missing_days]),
    "provisional": lambda r: "provisional %s" % [d.isoformat() for d in r.search.provisional],
    "top_queries": lambda r: "; ".join("%s %d/%d pos %s" % (t.shown_key, t.clicks, t.impressions, t.position) for t in r.search.top_queries),
    "top_pages shown keys": lambda r: "shown keys %s (distinct %d)" % ([t.shown_key for t in r.search.top_pages], len({t.shown_key for t in r.search.top_pages})),
    "top_queries shown keys": lambda r: "shown keys %s (distinct %d)" % ([t.shown_key for t in r.search.top_queries],
                                                                         len({t.shown_key for t in r.search.top_queries})),
    "template narrative, position line": lambda r: [ln for ln in template_insights(r).lines if ln.startswith("Average position")][0],
    "saved deck, Top pages table first column": _deck_top_pages,
}

# ---------------------------------------------------------------- case 39: the real runner into Postgres

DSN = os.environ.get("ANALYTICS_TEST_DSN", "postgresql://analytics:analytics@127.0.0.1:5433/analytics")

try:
    import psycopg
except ImportError:  # pragma: no cover
    psycopg = None


def _reachable() -> bool:
    try:
        with psycopg.connect(DSN, connect_timeout=3):
            return True
    except Exception:
        return False


needs_postgres = pytest.mark.skipif(psycopg is None or not _reachable(), reason="Postgres not reachable, run docker compose up -d")


def _sql(q, params=()):
    with psycopg.connect(DSN, autocommit=True) as c, c.cursor() as cur:
        cur.execute(q, params)
        return cur.fetchall() if cur.description else None


def _world():
    t = uuid.uuid4()
    with psycopg.connect(DSN, autocommit=True) as c, c.cursor() as cur:
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(t),))
        cur.execute("insert into tenant (id, slug, name) values (%s, %s, %s)", (t, "adv-" + t.hex[:8], "adv"))
        cur.execute("insert into platform_connection (tenant_id, platform, external_account, display_name, credentials_enc) "
                    "values (%s, 'gsc', %s, 'adv search', %s) returning id", (t, SITE, b"ciphertext"))
        return t, cur.fetchone()[0]


async def _sync(t, cid, inp):
    from app.db import Database
    from app.sync.runner import SyncRunner

    db = Database(DSN, max_size=2, role="sync_worker")
    await db.connect()
    client = httpx.AsyncClient(transport=make_transport(inp["responses"], []))

    async def no_sleep(s):
        return None

    conn = GSCConnector(site_url=SITE, access_token="tok-adv-SECRET", client=client, top_n=inp.get("top_n", 10), sleep=no_sleep)
    try:
        return await SyncRunner(db).run_once(tenant_id=str(t), connection_id=str(cid), connector=conn, window_start=date.fromisoformat(inp["window"][0]),
                                             window_end=date.fromisoformat(inp["window"][1]), trigger="backfill")
    finally:
        await client.aclose()
        await db.close()


def db_actual(inp):
    t, cid = _world()
    outs = []
    for step in inp["syncs"]:
        try:
            o = asyncio.run(_sync(t, cid, step))
            outs.append("%s %s" % (o.status, o.error_code or ""))
        except Exception as exc:  # noqa: BLE001
            outs.append("raised %s" % type(exc).__name__)
    tops = _sql("select observed_on::text, dimension, left(key, 60), clicks from search_top_row where connection_id = %s order by 1, 2, 3", (cid,))
    snaps = _sql("select count(*) from metric_snapshot where connection_id = %s", (cid,))[0][0]
    runs = _sql("select status, error_code from sync_run where connection_id = %s order by started_at", (cid,))
    return "syncs %s; runs %s; snapshots %d; top %s" % (outs, runs, snaps, tops)


@pytest.fixture(scope="module")
def database():
    os.environ["DATABASE_URL"] = DSN
    migrations = sorted(p.name for p in (ROOT / "migrations").glob("*.sql"))
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("drop schema if exists public cascade; create schema public")
        for name in migrations:
            cur.execute((ROOT / "migrations" / name).read_text(encoding="utf-8"))
        cur.execute("grant usage on schema public to sync_worker")
    yield DSN


# ---------------------------------------------------------------- the replay


def actual(index: int) -> str:
    inp = CASES[index]["input"]
    kind = inp["kind"]
    if kind == "gsc_connector":
        return CONNECTOR_SUMMARY[index](connector(inp))
    if kind == "search_report":
        return REPORT_CALLS[inp["call"]](search_report(inp))
    raise AssertionError("case %d has kind %r, which has its own replay" % (index, kind))


def test_the_corpus_is_the_one_that_was_written():
    assert len(CASES) == 40
    assert sum(1 for c in CASES if c["expect"] == "pass") == 24 and sum(1 for c in CASES if c["expect"] == "fail") == 16
    for n, c in enumerate(CASES):
        assert c["source"].split(" ")[0] in ("real", "invented"), n
        assert len(c["why"]) >= 10 and c["expected"] and c["promise"], n


def test_no_case_is_skipped_silently():
    assert NOT_REPLAYABLE == {} and KNOWN_LIMITS == {}
    assert {c["input"]["kind"] for c in CASES} == {"gsc_connector", "search_report", "gsc_sync_db"}
    assert [n for n, c in enumerate(CASES) if c["input"]["kind"] == "gsc_sync_db"] == [39]
    assert sorted(CONNECTOR_SUMMARY) == [n for n, c in enumerate(CASES) if c["input"]["kind"] == "gsc_connector"]


REPLAYED = [n for n, c in enumerate(CASES) if c["input"]["kind"] in ("gsc_connector", "search_report")]


@pytest.mark.parametrize("index", [n for n in REPLAYED if n not in FIXED])
def test_the_code_gives_the_outcome_the_corpus_expects(index):
    case = CASES[index]
    assert actual(index) == case["expected"], "case %d (%s): %s" % (index, case["expect"], case["why"])
    assert actual(index) == case["actual"]


def test_case_35_five_page_addresses_of_one_site_print_as_five_different_texts():
    got = actual(35)
    assert got.endswith("(distinct 5)"), got
    assert got != CASES[35]["actual"] and "(distinct 1)" in CASES[35]["actual"]


def test_case_36_three_long_queries_that_start_alike_print_as_three_different_texts_and_the_short_one_whole():
    got = actual(36)
    assert got.endswith("(distinct 4)") and "'roof repair near me'" in got, got
    assert "(distinct 2)" in CASES[36]["actual"]


def test_case_37_the_narrative_says_a_position_that_went_from_6_to_9_got_worse():
    line = actual(37)
    assert line == "Average position was 9.0 in October, worse by 50.0% (further from the top) from September (6.0)."
    assert "up 50.0%" in CASES[37]["actual"] and "up 50.0%" not in line


def test_case_38_the_saved_deck_shows_five_different_page_cells():
    cells = ast.literal_eval(actual(38).partition(": ")[2])      # the list of five strings `_deck_top_pages` wrote
    assert len(cells) == 5 and len(set(cells)) == 5, cells
    assert actual(38) != CASES[38]["actual"]


@needs_postgres
def test_case_39_a_resync_of_a_day_replaces_that_days_top_rows_as_a_set(database):
    case = CASES[39]
    assert db_actual(case["input"]) == case["expected"] == case["actual"]
