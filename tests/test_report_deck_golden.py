"""Snapshots of whole decks for a handful of reports, so the wording, the geometry and the chart settings of the ad and search slides
are pinned as they were when someone looked at the rendered pages.

Each scenario builds a deck from hand made rows (no database) and is compared with a file in tests/fixtures/deck_golden: the text of
every slide, the facts the narrative check holds a model to, and one SHA-256 for the exact XML of each slide and each chart. A
change that moves a box, resizes a font, renames a series or rewords a note turns exactly one of those lines red and the failure
shows which slide. That is a characterisation test, not a statement that the layout is right: the images in docs/monthly-report-sample
are where it was looked at. After an intended change, look at the pages again and then rewrite the files with

    UPDATE_DECK_GOLDEN=1 python -m pytest tests/test_report_deck_golden.py

The hashes depend on python-pptx 1.0.2 and lxml (both pinned in requirements.txt).
"""
from __future__ import annotations

import difflib
import hashlib
import json
import os
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.report.deck import dump_deck, save_deck  # noqa: E402
from app.report.insights import Insights, report_facts, template_insights  # noqa: E402
from app.report.monthly import build_report  # noqa: E402
from app.report.paid import attach_paid  # noqa: E402
from app.report.search import attach_search  # noqa: E402
from app.report.web import attach_web  # noqa: E402
from report_fixtures import AS_OF, COVERED, SEPTEMBER_COVERED, both_months  # noqa: E402
from sources_fixtures import (  # noqa: E402
    PAID_COVERED,
    SEARCH_COVERED,
    paid_day,
    paid_both,
    paid_october,
    paid_september,
    search_both,
    search_day,
    top_october,
)
from test_report_web import WEB_COVERED, WEB_ROWS  # noqa: E402

GOLDEN = Path(__file__).resolve().parent / "fixtures" / "deck_golden"
UPDATE = os.environ.get("UPDATE_DECK_GOLDEN") == "1"


def build(*, web=False, web_rows=WEB_ROWS, web_covered=WEB_COVERED, paid=None, search=None, tops=None, as_of=AS_OF, covered=COVERED, paid_covered=PAID_COVERED,
          search_covered=SEARCH_COVERED, store=None):
    report = build_report(both_months() if store is None else store, 2026, 10, label="Example Outfitters", covered=covered, as_of=as_of)
    if web:
        attach_web(report, web_rows, label="Example site", covered=web_covered, as_of=as_of)
    if paid is not None:
        attach_paid(report, paid, label="Example ads", covered=paid_covered, as_of=as_of)
    if search is not None:
        attach_search(report, search, top_october() if tops is None else tops, label="Example search", covered=search_covered, as_of=as_of)
    return report


LONG = ("Ad spend came to EUR 180.00 in October against EUR 150.00 in September, with the return on ad spend at 2.9 against 3.4, "
        "which is the figure to watch while the purchase counts are still arriving inside the click window")
LONG_NARRATIVE = Insights(tuple([LONG] * 3 + [LONG[:120]] * 3), "template", "template")
NINE_LINES = Insights(tuple("Line number %d of the narrative." % n for n in range(9)), "template", "template")


def scenarios():
    sept_zero = paid_day(date(2026, 9, 1), 0, 0, 0, 0, 0) + paid_day(date(2026, 9, 2), 0, 0, 0, 0, 0)
    sept_no_purchases = paid_day(date(2026, 9, 1), 10000, 1000, 50) + paid_day(date(2026, 9, 2), 5000, 500, 20)
    october_only = paid_october()
    search_october_only = search_day(date(2026, 10, 1), 12, 150, 5000) + search_day(date(2026, 10, 2), 8, 50, 9000)
    october_runs = [(date(2026, 10, 1), date(2026, 11, 1))]
    full_runs = [(date(2026, 9, 1), date(2026, 11, 1))]
    oct_full_sept_hole = [(date(2026, 9, 1), date(2026, 9, 10)), (date(2026, 9, 12), date(2026, 11, 1))]
    zero_clicks = paid_day(date(2026, 9, 1), 10000, 1000, 50, 5, 40000) + paid_day(date(2026, 10, 1), 12000, 1200, 0, 0, 0)
    full_october = [r for n in range(3, 32) for r in paid_day(date(2026, 10, n), 1000, 100, 5, 1, 4000)]
    big_top = [dict(r, clicks=r["clicks"] * 400, impressions=r["impressions"] * 900) for r in top_october()]
    return {
        "all_four_sources": (build(web=True, paid=paid_both(), search=search_both()), None),
        "paid_only": (build(paid=paid_both()), None),
        "search_only": (build(search=search_both()), None),
        "paid_no_earlier_month": (build(paid=october_only, search=search_october_only), None),
        "paid_zero_baseline": (build(paid=sept_zero + paid_october()), None),
        "paid_earlier_month_without_purchases": (build(paid=sept_no_purchases + paid_october()), None),
        "paid_and_search_only_october_covered": (build(paid=october_only, search=search_october_only, paid_covered=october_runs,
                                                       search_covered=october_runs), None),
        "every_day_covered": (build(paid=paid_both(), search=search_both(), covered=full_runs, paid_covered=full_runs,
                                    search_covered=full_runs), None),
        "paid_and_search_only_september": (build(paid=paid_september(), search=search_day(date(2026, 9, 1), 10, 100, 7500),
                                                 paid_covered=SEPTEMBER_COVERED, search_covered=SEPTEMBER_COVERED), None),
        "provisional_days_everywhere": (build(web=True, paid=paid_both(), search=search_both(), as_of=date(2026, 11, 2)), None),
        "ads_covered_but_no_rows": (build(paid=[], covered=full_runs, paid_covered=full_runs), None),
        "web_without_rows": (build(web=True, web_rows=[], web_covered=[]), None),
        "web_one_day_missing": (build(web=True, web_covered=[(date(2026, 9, 1), date(2026, 10, 5)), (date(2026, 10, 6), date(2026, 11, 1))]), None),
        "web_every_day_covered": (build(web=True, web_covered=full_runs, covered=full_runs), None),
        "provisional_days_more_than_listed": (build(paid=paid_both(), search=search_both(), as_of=date(2026, 11, 1), covered=full_runs,
                                                    paid_covered=full_runs, search_covered=full_runs), None),
        "web_many_days_missing": (build(web=True, web_covered=[(date(2026, 9, 1), date(2026, 10, 1)), (date(2026, 10, 20), date(2026, 11, 1))]), None),
        "september_has_a_hole": (build(paid=paid_both(), search=search_both(), covered=oct_full_sept_hole, paid_covered=oct_full_sept_hole,
                                       search_covered=oct_full_sept_hole), None),
        "paid_zero_clicks": (build(paid=zero_clicks), None),
        "paid_full_october_and_quiet_september": (build(paid=paid_both() + full_october, paid_covered=full_runs), None),
        "big_top_rows": (build(search=search_both(), tops=big_top), None),
        "month_in_progress_no_runs": (build(web=True, paid=paid_both(), search=search_both(), as_of=date(2026, 10, 20), covered=(),
                                            paid_covered=(), search_covered=()), None),
        "long_narrative": (build(paid=paid_both(), search=search_both()), LONG_NARRATIVE),
        "web_nine_lines": (build(web=True), NINE_LINES),
        "store_nine_lines": (build(), NINE_LINES),
        "web_only": (build(web=True), None),
    }


def xml_hash(element) -> str:
    return hashlib.sha256(etree.tostring(element, method="c14n")).hexdigest()[:16]


def snapshot(report, insights, tmp_path) -> str:
    path = save_deck(report, insights or template_insights(report), tmp_path / "deck.pptx")
    from pptx import Presentation
    prs = Presentation(str(path))
    out = ["## text of every slide", dump_deck(prs), "## notes"]
    for n, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide:
            out.append("slide %d notes: %s" % (n, slide.notes_slide.notes_text_frame.text))
    out.append("## facts the narrative check holds a model to")
    out.append(json.dumps(report_facts(report), indent=1, sort_keys=True, default=lambda v: str(v) if isinstance(v, (Decimal, date)) else repr(v)))
    out.append("## exact xml, one hash for each slide and each chart")
    for n, slide in enumerate(prs.slides, 1):
        out.append("slide %d %s" % (n, xml_hash(slide._element)))
        for shape in slide.shapes:
            if shape.has_chart:
                out.append("slide %d chart %s %s" % (n, shape.name, xml_hash(shape.chart._chartSpace)))
    return "\n".join(out) + "\n"


@pytest.mark.parametrize("name", sorted(scenarios()))
def test_the_deck_is_what_it_was_when_the_pages_were_looked_at(name, tmp_path):
    report, insights = scenarios()[name]
    got = snapshot(report, insights, tmp_path)
    target = GOLDEN / (name + ".txt")
    if UPDATE:
        GOLDEN.mkdir(parents=True, exist_ok=True)
        target.write_text(got, encoding="utf-8", newline="\n")
    assert target.exists(), "no snapshot yet: run with UPDATE_DECK_GOLDEN=1 and look at the pages"
    want = target.read_text(encoding="utf-8").replace("\r\n", "\n")
    if got != want:
        diff = "\n".join(list(difflib.unified_diff(want.splitlines(), got.splitlines(), "snapshot", "now", lineterm="", n=0))[:40])
        pytest.fail("the deck changed:\n" + diff)
