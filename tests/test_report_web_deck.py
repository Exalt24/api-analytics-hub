"""The multi source deck, saved and read back, and the narrative check over the website figures.

A tenant with no website renders exactly as before: the saved read-back of the store-only sample deck
(tests/fixtures/store_only_report_2026-10.txt, taken from the commit before the website was added) is
compared with a fresh build, character for character. A tenant with a website gets eight slides, and the
figure check covers the new metrics from both sides.
"""
from __future__ import annotations

import asyncio
import json
import re
import sys
from datetime import date
from pathlib import Path

import httpx
import pytest
from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.connectors.ga4 import METRICS  # noqa: E402
from app.report.deck import dump_deck, save_deck  # noqa: E402
from app.report.demo_data import DEMO_RUNS, WEB_SYNC_WINDOWS, demo_points, demo_web_days  # noqa: E402
from app.report.insights import (  # noqa: E402
    ReportInsightWriter,
    check_model_text,
    report_facts,
    template_insights,
    ungrounded_numbers,
    write_insights,
)
from app.report.monthly import build_report  # noqa: E402
from app.report.web import attach_web  # noqa: E402
from report_fixtures import AS_OF, COVERED, both_months  # noqa: E402
from test_report_web import WEB_COVERED, WEB_ROWS  # noqa: E402

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "store_only_report_2026-10.txt"
SAMPLE_AS_OF = date(2026, 11, 2)
NOTE = "Sample report: the orders are invented, not a real store."


def run(coro):
    return asyncio.run(coro)


def multi_report():
    report = build_report(both_months(), 2026, 10, label="Example Outfitters", covered=COVERED, as_of=AS_OF)
    return attach_web(report, WEB_ROWS, label="Example site", covered=WEB_COVERED, as_of=AS_OF)


def saved(tmp_path, report=None, note=""):
    report = report or multi_report()
    insights = template_insights(report)
    path = save_deck(report, insights, tmp_path / "deck.pptx", data_note=note)
    return Presentation(str(path)), report, insights, path


def tables(slide):
    return [[[c.text for c in row.cells] for row in s.table.rows] for s in slide.shapes if s.has_table]


def charts(slide):
    return [s.chart for s in slide.shapes if s.has_chart]


def all_text(prs):
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                out.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(out)


# ---------------------------------------------------------------- a store with no website is unchanged

def sample_store_only_dump(tmp_path):
    rows = [{"observed_on": p.observed_on, "metric_key": p.metric_key, "value_numeric": p.value_numeric,
             "currency": p.currency} for p in demo_points()]
    covered = [(r.window_start, r.window_end) for r in DEMO_RUNS if r.status in ("succeeded", "partial")]
    report = build_report(rows, 2026, 10, label="Example Outfitters", covered=covered, as_of=SAMPLE_AS_OF)
    assert report.web is None
    insights = run(write_insights(report))
    return dump_deck(save_deck(report, insights, tmp_path / "so.pptx", data_note=NOTE))


def test_a_store_with_no_website_renders_exactly_as_it_did_before_the_website_existed(tmp_path):
    saved_text = FIXTURE.read_text(encoding="utf-8").replace("\r\n", "\n").strip()
    assert sample_store_only_dump(tmp_path) == saved_text
    assert saved_text.count("--- slide ") == 7


# ---------------------------------------------------------------- the multi source deck

def test_with_a_website_the_deck_is_eight_slides_and_the_titles_say_what_each_one_is(tmp_path):
    prs, report, _, _ = saved(tmp_path)
    titles = [s.shapes.title.text_frame.text for s in prs.slides]
    assert len(titles) == 8
    assert titles[0] == "Monthly performance report"
    assert titles[2] == "Net sales by day in October" and titles[3] == "Sessions and orders by day in October"
    assert titles[4] == "Order value, buyers and best days in October"
    assert titles[5] == ("October sessions total is 36.7% above September, but October is missing 2 days of website data: "
                         "per reported day it is up 41.4%")
    assert titles[7] == "Data health: October is missing 2 days of store data and 2 days of website data"


def test_the_website_table_holds_the_hand_computed_figures_and_their_notes(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    table = tables(prs.slides[5])[0]
    assert table[0] == ["Metric", "October 2026", "September 2026", "Change", "Change %", "Note"]
    assert table[1] == ["Sessions", "410", "300", "+110", "+36.7%", "October: 2 days missing; per reported day up 41.4%"]
    assert table[2][:5] == ["Users per active day", "150.0", "115.0", "+35.0", "+30.4%"]
    assert table[2][5] == "October: 2 days missing"
    assert table[5] == ["Order rate", "4.9%", "5.0%", "-0.1 pts", "-2.4%", "October: 2 days missing; 27 days in both"]


def test_the_weekly_table_sets_sessions_against_orders_and_says_how_many_days_it_counted(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    weeks = tables(prs.slides[5])[1]
    assert weeks[0] == ["Week", "Days counted", "Sessions", "Orders", "Order rate"]
    assert weeks[1] == ["Oct 1 to 7", "5 of 7", "410", "20", "4.9%"]
    assert weeks[2] == ["Oct 8 to 14", "5 of 7", "0", "0", "n/a"]


def test_the_sessions_chart_has_a_gap_for_a_day_with_no_website_data_and_the_orders_chart_keeps_its_own(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    sessions, orders = charts(prs.slides[3])
    cur = {s.name: list(s.values) for s in sessions.plots[0].series}
    assert cur["October 2026"][:6] == [250.0, 160.0, 0.0, 0.0, None, None], "days 3 and 4 are quiet (zero), 5 and 6 missing (a gap)"
    assert cur["Day with no data"][4] == 0.0 and cur["Day with no data"][3] is None
    assert cur["September 2026"][:2] == [200.0, 100.0]
    ord_cur = {s.name: list(s.values) for s in orders.plots[0].series}
    assert ord_cur["October 2026"][:2] == [12.0, 8.0] and ord_cur["October 2026"][9] is None, "the store's own missing days"
    text = "\n".join(s.text_frame.text for s in prs.slides[3].shapes if s.has_text_frame)
    assert "Sessions: a dot on the zero line is a day with no sessions reported, a break with an X is a day with no data (Oct 5, Oct 6)." in text
    assert "Orders: a dot on the zero line is a day with no orders, a break with an X is a day with no data (Oct 10, Oct 11)." in text


def test_the_data_health_slide_names_the_days_with_no_data_for_each_source(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    table = tables(prs.slides[7])[0]
    assert table[0] == ["Count", "Store, October", "Store, September", "Website, October", "Website, September"]
    by_label = {row[0]: row[1:] for row in table[1:]}
    assert by_label["Days with data"] == ["2", "2", "2", "2"]
    assert by_label["Days with no data"] == ["2", "0", "2", "0"]
    assert by_label["Days with nothing to report"] == ["27", "28", "27", "28"]
    text = "\n".join(s.text_frame.text for s in prs.slides[7].shapes if s.has_text_frame)
    assert "Days with no data in October, store: Oct 10, Oct 11." in text
    assert "Days with no data in October, website: Oct 5, Oct 6." in text


def test_each_sources_days_add_up_to_the_days_in_the_month(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    table = tables(prs.slides[7])[0]
    by_label = {row[0]: [int(v) for v in row[1:]] for row in table[1:] if row[0] != "Stored daily figures used"}
    for i in range(4):
        assert by_label["Days with data"][i] + by_label["Days with nothing to report"][i] + by_label["Days with no data"][i] \
            == by_label["Days in the month"][i]


def test_the_health_title_names_which_source_is_missing_days():
    from app.report.deck import _health_title_multi

    full = [(date(2026, 9, 1), date(2026, 10, 1)), (date(2026, 10, 1), date(2026, 11, 1))]

    def title(store_covered, web_covered):
        report = build_report(both_months(), 2026, 10, covered=store_covered, as_of=AS_OF)
        return _health_title_multi(attach_web(report, WEB_ROWS, covered=web_covered, as_of=AS_OF))

    assert title(COVERED, WEB_COVERED) == "Data health: October is missing 2 days of store data and 2 days of website data"
    assert title(COVERED, full) == "Data health: October is missing 2 of 31 days of store data"
    assert title(full, WEB_COVERED) == "Data health: October is missing 2 of 31 days of website data"
    assert title(full, full) == "Data health: every day of October is covered in both sources"


def test_the_deck_uses_plain_words_no_iso_date_and_no_text_under_eleven_points(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    everything = all_text(prs)
    assert not re.search(r"\d{4}-\d{2}-\d{2}", everything) and "USD" not in everything
    for word in ("snapshot", "sync", "succeeded", "runReport", "propertyQuota", "keyEvents", "totalUsers", "metric_key"):
        assert word.lower() not in everything.lower(), word
    sizes = []
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                sizes += [run.font.size for para in frame.paragraphs for run in para.runs]
    assert sizes and all(size is not None and size >= Pt(11) for size in sizes)


def test_every_figure_the_template_prints_about_the_website_is_a_computed_figure(tmp_path):
    prs, report, insights, _ = saved(tmp_path)
    facts = report_facts(report)
    for line in insights.lines:
        assert ungrounded_numbers(line, facts) == [], line
    joined = "\n".join(insights.lines)
    assert ("Sessions total was 410 in October, 36.7% above September (300), but October is missing 2 days of website data: "
            "per reported day, sessions are up 41.4% (14.14 against 10.00).") in joined
    assert "Order rate was 4.9% in October, down 2.4% from September (5.0%), and October is missing 2 days of website data." in joined
    assert "2 days have no website data (Oct 5, Oct 6), so the website totals are understated." in joined


def test_the_deck_file_still_reads_back_with_its_provenance_line(tmp_path):
    _, _, insights, path = saved(tmp_path)
    assert dump_deck(path).rstrip().endswith("--- provenance (file property, not on a slide): " + insights.label)


# ---------------------------------------------------------------- the figure check over the website metrics

FACTS = report_facts(multi_report())


def test_the_website_facts_are_in_the_same_shape_as_the_stores_and_carry_no_year():
    assert FACTS["Sessions"] == "410" and FACTS["Sessions previous"] == "300" and FACTS["Sessions change percent"] == "+36.7%"
    assert FACTS["Sessions per reported day"] == "14.14"
    assert FACTS["days with no website data"] == 2 and FACTS["dates with no website data"] == "Oct 5, Oct 6"
    assert FACTS["Order rate"] == "4.9%" and FACTS["Order rate change"] == "-0.1 pts"
    assert not any("2026" in str(v) for v in FACTS.values())


@pytest.mark.parametrize("sentence", [
    "Sessions were 410 in October, up 36.7% from September.",
    "Visits were 410 in October.",
    "Page views were 1,100 in October, up 29.4% from September.",
    "Views were 1,100 in October.",
    "Users per active day was 150.0 in October.",
    "Visitors averaged 150.0 per active day.",
    "Key events were 20 in October, up 5 from September.",
    "Order rate was 4.9% in October, down 2.4% from September.",
    "2 days have no website data (Oct 5, Oct 6), so the website totals are understated.",
])
def test_a_website_figure_beside_its_own_metric_passes(sentence):
    assert ungrounded_numbers(sentence, FACTS) == []


@pytest.mark.parametrize("sentence,flagged", [
    ("Page views were 410 in October.", "410"),                      # sessions' figure on page views
    ("Sessions were 1,100 in October.", "1,100"),                    # page views' figure on sessions
    ("Sessions were 411 in October.", "411"),                        # invented
    ("Sessions rose 36.8% from September.", "36.8%"),                # a percent nobody computed
    ("Visits and conversions were 20 in October.", "20"),            # conversions is not a metric of the report, and 20 is key events'
    ("Net sales were 410 in October.", "410"),                       # sessions' figure on the store's metric
    ("Sessions were up 410%.", "410%"),
    ("Order rate was 5.0% last month and 4.9% now.", None),          # control: both figures are facts, written as they are
    ("Order rate was 5.0 percent last month.", "5.0"),               # the same figure written another way is refused
    ("Sessions were 410 in October and twenty on the best day.", "twenty"),
])
def test_a_website_figure_beside_the_wrong_metric_or_not_computed_is_refused(sentence, flagged):
    bad = ungrounded_numbers(sentence, FACTS)
    if flagged is None:
        assert bad == [], bad
    else:
        assert flagged in bad, (sentence, bad)


def test_a_sentence_about_no_metric_cannot_carry_a_website_figure():
    assert ungrounded_numbers("Traffic rose 36.7% from September.", FACTS) == ["36.7%"]


def test_the_store_metrics_still_refuse_each_others_figures_with_the_website_in_the_facts():
    # 410 is the sessions figure: it must not pass for the store's orders
    assert "410" in ungrounded_numbers("Orders were 410 in October.", FACTS)
    assert ungrounded_numbers("Orders were 20 in October.", FACTS) == []


GOOD_MODEL = (
    "Sessions were 410 in October, up 36.7% from September.\n"
    "Page views were 1,100 in October.\n"
    "2 days have no website data, so the website totals are understated."
)


def replying(content):
    def handler(request):
        return httpx.Response(200, json={"choices": [{"message": {"content": content}}]})
    return ReportInsightWriter("test-key", client=httpx.AsyncClient(transport=httpx.MockTransport(handler)))


def test_a_model_reply_about_the_website_is_used_when_every_figure_is_a_fact():
    insights = run(write_insights(multi_report(), replying(GOOD_MODEL)))
    assert insights.source == "model" and list(insights.lines) == GOOD_MODEL.splitlines()


def test_a_model_reply_with_a_wrong_website_figure_falls_back_to_the_template():
    bad = GOOD_MODEL.replace("1,100", "410")
    insights = run(write_insights(multi_report(), replying(bad)))
    assert insights.source == "template_fallback" and "410" in insights.rejected_reason
    assert any("Sessions total was 410" in line for line in insights.lines)


def test_the_model_is_sent_the_website_facts_and_never_a_row_or_a_year():
    seen = []

    def handler(request):
        seen.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"message": {"content": GOOD_MODEL}}]})

    writer = ReportInsightWriter("test-key", client=httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    run(write_insights(multi_report(), writer))
    user = seen[0]["messages"][1]["content"]
    assert "Sessions: 410" in user and "dates with no website data: Oct 5, Oct 6" in user
    assert "observed_on" not in user and "value_numeric" not in user and "2026" not in user
    assert "website" in seen[0]["messages"][0]["content"]


def test_check_model_text_refuses_a_conversions_claim_because_no_such_metric_is_in_the_facts():
    lines, reason = check_model_text("Conversions were 20 in October.", FACTS)
    assert lines == [] and "20" in reason


# ---------------------------------------------------------------- the sample, end to end without a database

def test_the_sample_builds_through_the_same_functions_and_has_eight_slides(tmp_path):
    canon = dict(METRICS)
    rows = [{"observed_on": p.observed_on, "metric_key": p.metric_key, "value_numeric": p.value_numeric,
             "currency": p.currency} for p in demo_points()]
    web_rows = [{"observed_on": d, "metric_key": canon[n], "value_numeric": v, "currency": None}
                for d, vals in demo_web_days().items() for n, v in vals.items()]
    covered = [(r.window_start, r.window_end) for r in DEMO_RUNS if r.status in ("succeeded", "partial")]
    report = build_report(rows, 2026, 10, label="Example Outfitters", covered=covered, as_of=SAMPLE_AS_OF)
    attach_web(report, web_rows, label="Example site", covered=list(WEB_SYNC_WINDOWS), as_of=SAMPLE_AS_OF)
    insights = template_insights(report)
    text = dump_deck(save_deck(report, insights, tmp_path / "s.pptx", data_note=NOTE))
    assert text.count("--- slide ") == 8
    assert report.web.current.missing_days == [date(2026, 10, 3), date(2026, 10, 4), date(2026, 10, 5)]
