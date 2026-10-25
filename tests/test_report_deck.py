"""The deck, saved and read back through python-pptx.

Each assertion is about what a person sees on a slide: how many slides, what the titles say
(with the missing day caveat inside them), that the table cells hold the hand computed
figures from report_fixtures.py, that a day with no orders and a day with no data are drawn
differently, that the slides use plain words and no pipeline jargon, and that the file
itself carries honest properties. The deck is saved to disk and reopened, so a pass means
the file says it.
"""
from __future__ import annotations

import re
import sys
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.report.deck import NARRATIVE_FOOTER, dump_deck, save_deck
from app.report.insights import report_facts, template_insights, ungrounded_numbers
from app.report.monthly import build_report
from report_fixtures import AS_OF, COVERED, OCTOBER_COVERED, SEPTEMBER_COVERED, both_months, october_rows

FULL = SEPTEMBER_COVERED + [(date(2026, 10, 1), date(2026, 11, 1))]


def saved(tmp_path, rows=None, covered=COVERED, note=""):
    report = build_report(both_months() if rows is None else rows, 2026, 10, label="Example Outfitters",
                          covered=covered, as_of=AS_OF)
    insights = template_insights(report)
    path = save_deck(report, insights, tmp_path / "deck.pptx", data_note=note)
    return Presentation(str(path)), report, insights, path


def table_rows(slide, index=0):
    tables = [s.table for s in slide.shapes if s.has_table]
    return [[cell.text for cell in row.cells] for row in tables[index].rows]


def charts(slide):
    return [s.chart for s in slide.shapes if s.has_chart]


def slide_text(prs, n):
    """All the text on one slide (1 based): title, text boxes and table cells."""
    out = []
    for shape in prs.slides[n - 1].shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
        if shape.has_table:
            out.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(out)


def test_the_deck_has_seven_slides_with_the_planned_titles(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    assert [s.shapes.title.text_frame.text for s in prs.slides] == [
        "Monthly performance report",
        "October net sales total is 50.0% above September, but October is missing 2 days: per reported day it is up 55.2%",
        "Net sales by day in October",
        "Orders by day in October",
        "Order value, buyers and best days in October",
        "October refunds total is 200.0% above September, but October is missing 2 days: per reported day it is up 210.3%",
        "Data health: October is missing 2 of 31 days",
    ]


def test_a_complete_month_has_titles_without_a_caveat(tmp_path):
    prs, _, _, _ = saved(tmp_path, covered=FULL)
    assert prs.slides[1].shapes.title.text_frame.text == "Net sales rose 50.0% from September"
    assert prs.slides[6].shapes.title.text_frame.text == "Data health: every day of October is covered"


def test_the_kpi_table_holds_the_hand_computed_figures_and_names_the_month_in_each_note(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    note = "October partial: 2 days missing"
    assert table_rows(prs.slides[1]) == [
        ["Metric", "October 2026", "September 2026", "Change", "Change %", "Note"],
        ["Orders", "20", "15", "+5", "+33.3%", note + "; per reported day up 37.9%"],
        ["Gross sales", "$240.00", "$150.00", "+$90.00", "+60.0%", note + "; per reported day up 65.5%"],
        ["Refunds", "$30.00", "$10.00", "+$20.00", "+200.0%", note + "; per reported day up 210.3%"],
        ["Net sales", "$210.00", "$140.00", "+$70.00", "+50.0%", note + "; per reported day up 55.2%"],
        ["Average order value", "$12.00", "$10.00", "+$2.00", "+20.0%", note],
        ["Buyers per active day", "8.0", "6.5", "+1.5", "+23.1%", note],
    ]


def test_a_complete_month_has_no_per_reported_day_notes_or_definition(tmp_path):
    prs, _, _, _ = saved(tmp_path, covered=FULL)
    assert all(row[5] == "" for row in table_rows(prs.slides[1])[1:])
    assert "reported day" not in slide_text(prs, 2) and "reported day" not in slide_text(prs, 6)


def test_the_kpi_slide_defines_per_reported_day_when_it_shows_it(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    assert ("Per reported day divides a total by the days with orders plus the days with no orders"
            in slide_text(prs, 2))


def test_the_slides_that_show_buyers_define_it(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    for n in (2, 5):
        assert "Buyers per active day is the average of each day's distinct buyers, over days with orders" in slide_text(prs, n)


def test_the_value_slide_carries_the_notes_and_a_best_days_table_of_its_own_width(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    slide = prs.slides[4]
    assert table_rows(slide, 0)[1] == ["Average order value", "$12.00", "$10.00", "+$2.00", "+20.0%", "October partial: 2 days missing"]
    assert table_rows(slide, 1) == [["Rank", "Day", "Net sales"], ["1", "Oct 1", "$150.00"], ["2", "Oct 2", "$60.00"]]
    widths = [s.width for s in slide.shapes if s.has_table]
    assert widths[1] < widths[0] * 0.6


def test_the_first_month_keeps_a_note_for_both_the_baseline_and_the_missing_days(tmp_path):
    prs, _, _, _ = saved(tmp_path, rows=october_rows(), covered=OCTOBER_COVERED)
    rows = table_rows(prs.slides[1])
    assert rows[0][2] == "Previous month" and rows[1][2:5] == ["n/a", "n/a", "n/a"]
    assert rows[1][5] == "no earlier month stored; October partial: 2 days missing"
    assert prs.slides[1].shapes.title.text_frame.text == "Net sales came to $210.00 in October (October is missing 2 days)"
    (chart,) = charts(prs.slides[2])
    assert [s.name for s in chart.plots[0].series] == ["October 2026", "Day with no data"]
    assert "No earlier month is stored" in dump_deck(prs)


# ------------------------------------------------------------ the charts

def test_the_sales_chart_has_a_gap_for_each_missing_day_and_zero_for_quiet_days(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    (chart,) = charts(prs.slides[2])
    october, september, missing = chart.plots[0].series
    assert october.name == "October 2026" and september.name == "September 2026"
    values = list(october.values)
    assert len(values) == 31
    assert values[0] == 150.0 and values[1] == 60.0     # (18000 - 3000) / 100 and 6000 / 100
    assert values[2] == 0.0                              # a day with no orders
    assert values[9] is None and values[10] is None      # October 10 and 11 were never pulled
    assert list(september.values)[:2] == [90.0, 50.0]
    assert len(list(september.values)) == 30             # September has 30 days: no padded 31st point


def test_a_day_with_no_data_is_marked_apart_from_a_day_with_no_orders(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    (chart,) = charts(prs.slides[2])
    marker = chart.plots[0].series[2]
    assert marker.name == "Day with no data"
    assert [i for i, v in enumerate(marker.values) if v == 0.0] == [9, 10]
    assert [v for i, v in enumerate(marker.values) if i not in (9, 10)] == [None] * 29
    assert marker._element.xpath("c:marker/c:symbol/@val") == ["x"]


def dotted(series):
    """The point indexes of a chart series that carry their own circle marker."""
    return [int(i) for i in series._element.xpath("c:dPt[c:marker/c:symbol/@val='circle']/c:idx/@val")]


def test_a_day_with_no_orders_is_a_dot_on_both_lines_not_only_the_current_one(tmp_path):
    prs, report, _, _ = saved(tmp_path)
    (chart,) = charts(prs.slides[2])
    october, september = chart.plots[0].series[0], chart.plots[0].series[1]
    oct_quiet = [d.day - 1 for d in report.current.quiet_days]
    sept_quiet = [d.day - 1 for d in report.previous.quiet_days]
    assert len(oct_quiet) == 27 and len(sept_quiet) == 28          # the fixture's days with no orders
    assert dotted(october) == oct_quiet
    assert dotted(september) == sept_quiet
    for index in sept_quiet:
        assert list(september.values)[index] == 0.0                  # a dot sits on a zero point
    assert dotted(chart.plots[0].series[2]) == []


def test_each_dot_takes_its_own_lines_colour(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    (chart,) = charts(prs.slides[2])
    colours = [[c.get("val") for c in series._element.xpath("c:dPt")[0].xpath("c:marker/c:spPr/a:solidFill/a:srgbClr")]
               for series in list(chart.plots[0].series)[:2]]
    assert colours == [["1B7F8C"], ["8A94A0"]]


def test_the_orders_chart_has_dots_for_its_one_line(tmp_path):
    prs, report, _, _ = saved(tmp_path)
    (chart,) = charts(prs.slides[3])
    assert dotted(chart.plots[0].series[0]) == [d.day - 1 for d in report.current.quiet_days]


def test_a_previous_month_day_with_no_data_gets_its_own_grey_x_and_a_footnote(tmp_path):
    hole = [(date(2026, 9, 1), date(2026, 9, 9)), (date(2026, 9, 10), date(2026, 11, 1))]
    prs, _, _, _ = saved(tmp_path, covered=hole)
    (chart,) = charts(prs.slides[2])
    series = chart.plots[0].series
    assert [s.name for s in series] == ["October 2026", "September 2026", "September day with no data"]
    assert [i for i, v in enumerate(series[2].values) if v == 0.0] == [8]
    assert "A break in the grey line with a grey X is a day with no data in September (Sep 9)." in slide_text(prs, 3)


def test_the_chart_footnote_says_which_marks_are_present_and_how_many_days_are_plotted(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    text = slide_text(prs, 3)
    assert "A dot on the zero line is a day with no orders, on either line." in text
    assert "a day with no data (Oct 10, Oct 11). 29 of 31 days are plotted." in text
    complete, _, _, _ = saved(tmp_path, covered=FULL)
    assert "no data" not in slide_text(complete, 3)


def test_the_plotted_points_agree_with_the_data_health_slide(tmp_path):
    prs, report, _, _ = saved(tmp_path)
    (chart,) = charts(prs.slides[3])
    plotted = [v for v in chart.plots[0].series[0].values if v is not None]
    health = table_rows(prs.slides[6])
    with_orders, no_orders, no_data = (int(health[i][1]) for i in (2, 3, 4))
    assert len(plotted) == with_orders + no_orders == 29
    assert with_orders + no_orders + no_data == int(health[1][1]) == 31


def test_the_daily_charts_label_their_value_axis(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    for n, label in ((3, "Net sales ($)"), (4, "Orders")):
        (chart,) = charts(prs.slides[n - 1])
        assert chart.value_axis.axis_title.text_frame.text == label


def test_every_chart_has_alt_text(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart:
                assert len(shape._element.nvGraphicFramePr.cNvPr.get("descr", "")) > 20


# ------------------------------------------------------------ the insights and the data health slides

def test_the_insights_slide_carries_the_checked_lines_each_with_its_caveat(tmp_path):
    prs, _, insights, _ = saved(tmp_path)
    paragraphs = [p.text for s in prs.slides[5].shapes if s.has_text_frame for p in s.text_frame.paragraphs if p.text]
    for line in insights.lines:
        assert line in paragraphs
    changes = [line for line in insights.lines if "September" in line]
    assert len(changes) == 5
    assert all("October is missing 2 days" in line for line in changes)
    totals = [line for line in changes if "total was" in line]
    assert len(totals) == 3 and all("per reported day" in line for line in totals)
    assert any(line.startswith("The biggest move was Refunds, 200.0% above September in total") for line in insights.lines)


def test_the_insights_slide_footer_is_neutral_and_the_writer_is_kept_in_the_file_property(tmp_path):
    prs, _, insights, path = saved(tmp_path)
    text = slide_text(prs, 6)
    assert NARRATIVE_FOOTER in text
    assert "model" not in text.lower() and "template" not in text.lower()
    assert prs.core_properties.comments == insights.label
    assert "No model was used" in dump_deck(path)


def test_every_number_in_the_insight_text_on_the_slide_is_a_computed_fact(tmp_path):
    prs, report, insights, _ = saved(tmp_path)
    facts = report_facts(report)
    for line in insights.lines:
        assert ungrounded_numbers(line, facts) == []


def test_the_data_health_slide_states_the_counts_and_names_the_missing_days(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    slide = prs.slides[6]
    assert table_rows(slide) == [
        ["Count", "October 2026", "September 2026"],
        ["Days in the month", "31", "30"],
        ["Days with orders", "2", "2"],
        ["Days with no orders", "27", "28"],
        ["Days with no data", "2", "0"],
        ["Stored daily figures used", "12", "12"],
        ["First and last day with orders", "Oct 1 to Oct 2", "Sep 1 to Sep 2"],
    ]
    text = slide_text(prs, 7)
    assert "Days with no data in October: Oct 10, Oct 11." in text
    assert "understated" in text


# ------------------------------------------------------------ plain words, plain dates, readable size

def test_no_slide_uses_pipeline_words(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    everything = "\n".join(slide_text(prs, n) for n in range(1, 8))
    for word in ("snapshot", "sync", "succeeded", "partial sync", "seed", "Demo store"):
        assert word.lower() not in everything.lower(), word


def test_no_slide_shows_an_iso_date_or_the_usd_code(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    everything = "\n".join(slide_text(prs, n) for n in range(1, 8))
    assert not re.search(r"\d{4}-\d{2}-\d{2}", everything)
    assert "USD" not in everything and "$" in everything


def test_no_table_has_an_empty_header_cell_and_the_title_slide_names_the_store_and_the_note(tmp_path):
    prs, _, _, _ = saved(tmp_path, note="Sample report: the orders are invented, not a real store.")
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                assert all(cell.text.strip() for cell in shape.table.rows[0].cells)
    assert slide_text(prs, 1) == ("Monthly performance report\nExample Outfitters, October 2026\n"
                                  "Sample report: the orders are invented, not a real store.")


def test_no_text_on_any_slide_is_smaller_than_eleven_points(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    sizes = []
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                sizes += [run.font.size for para in frame.paragraphs for run in para.runs]
    assert sizes and all(size is not None and size >= Pt(11) for size in sizes)


# ------------------------------------------------------------ the file itself

def test_no_slide_keeps_an_empty_placeholder_or_the_library_authors_name(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    for slide in prs.slides:
        for ph in slide.placeholders:
            assert ph.text_frame.text.strip(), "empty placeholder on a slide"
    assert prs.core_properties.author == "api-analytics-hub"
    assert "Canny" not in prs.core_properties.author + prs.core_properties.last_modified_by


def test_the_deck_is_stamped_with_this_build_not_the_librarys_template_date(tmp_path):
    prs, _, _, _ = saved(tmp_path)
    assert abs(datetime.now(timezone.utc).replace(tzinfo=None) - prs.core_properties.created) < timedelta(days=1)


def test_the_extended_properties_are_honest_not_the_templates_mac_four_by_three(tmp_path):
    _, _, _, path = saved(tmp_path)
    with zipfile.ZipFile(path) as z:
        app = z.read("docProps/app.xml").decode("utf-8")
        names = z.namelist()
    assert "Macintosh" not in app and "4:3" not in app
    assert "<Slides>7</Slides>" in app and "<Application>api-analytics-hub</Application>" in app
    assert "Widescreen" in app
    assert not [n for n in names if "thumbnail" in n]


def test_the_text_read_back_has_one_block_per_slide_and_the_provenance_line(tmp_path):
    _, _, _, path = saved(tmp_path)
    text = dump_deck(path)
    assert len(re.findall(r"^--- slide \d+: ", text, re.M)) == 7
    assert "$210.00" in text
    assert text.splitlines()[-1].startswith("--- provenance (file property, not on a slide): Written by a fixed template")
