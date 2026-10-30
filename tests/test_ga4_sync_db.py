"""GA4 through the real sync runner into real Postgres, then out through the monthly report.

The whole path, against the stand-in server (never a live property): GA4Connector over HTTP, the
SyncRunner as the sync_worker role, the dated snapshots, row level security as the request role, and
the report with both sources. It also proves what only a database can: a second sync rewrites the same
rows instead of adding more, a failed run is recorded and is not coverage, the stored rows of one tenant
never reach another's report, and the schema itself refuses a currency on a count and a metric nobody
registered.

Skips cleanly when Postgres is not running (`docker compose up -d`). Like the other database suites it
drops and recreates the `public` schema of the database it points at.
"""
from __future__ import annotations

import asyncio
import os
import subprocess
import sys
import uuid
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

psycopg = pytest.importorskip("psycopg", reason="psycopg not installed")
pytest.importorskip("asyncpg")

DSN = os.environ.get("ANALYTICS_TEST_DSN", "postgresql://analytics:analytics@127.0.0.1:5433/analytics")


def _reachable() -> bool:
    try:
        with psycopg.connect(DSN, connect_timeout=3):
            return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(not _reachable(), reason="Postgres not reachable, run docker compose up -d")

import httpx  # noqa: E402
from pptx import Presentation  # noqa: E402

from app.connectors.base import RateLimited  # noqa: E402
from app.connectors.ga4 import GA4Connector  # noqa: E402
from app.connectors.ga4_fake import QUOTA_BODY, Fault, FakeGA4  # noqa: E402
from app.db import Database  # noqa: E402
from app.report.deck import dump_deck  # noqa: E402
from app.report.monthly import ReportError  # noqa: E402
from app.report.service import generate, load_inputs  # noqa: E402
from app.sync.runner import SyncRunner  # noqa: E402
from report_fixtures import AS_OF, both_months  # noqa: E402
from test_report_service_db import RUNS, _seed  # noqa: E402

TOKEN = "tok-db-SECRET"
PROPERTY = "424242"
OTHER_PROPERTY = "515151"

WEB = {
    date(2026, 9, 1): {"sessions": 200, "totalUsers": 150, "screenPageViews": 600, "keyEvents": 10},
    date(2026, 9, 2): {"sessions": 100, "totalUsers": 80, "screenPageViews": 250, "keyEvents": 5},
    date(2026, 10, 1): {"sessions": 250, "totalUsers": 180, "screenPageViews": 700, "keyEvents": 12},
    date(2026, 10, 2): {"sessions": 160, "totalUsers": 120, "screenPageViews": 400, "keyEvents": 8},
}
#: The website runs the world makes: all of September, and October except the 5th and 6th.
WINDOWS = ((date(2026, 9, 1), date(2026, 10, 1)), (date(2026, 10, 1), date(2026, 10, 5)), (date(2026, 10, 7), date(2026, 11, 1)))
TITLE = ("October sessions total is 36.7% above September, but October is missing 2 days of website data: "
         "per reported day it is up 41.4%")


def run(coro):
    return asyncio.run(coro)


@pytest.fixture(scope="module")
def server():
    with FakeGA4(token=TOKEN) as fake:
        yield fake


def connect_row(cur, tenant, platform, account, name):
    cur.execute(
        "insert into platform_connection (tenant_id, platform, external_account, display_name, credentials_enc) "
        "values (%s, %s, %s, %s, %s) returning id", (tenant, platform, account, name, b"ciphertext"))
    return cur.fetchone()[0]


@pytest.fixture(scope="module")
def world(server):
    """Tenant A has a store and a website, B a website with huge figures, C a store only, D two websites, E a website only."""
    a, b, c, d, e = (uuid.uuid4() for _ in range(5))
    ids = {}
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("drop schema if exists public cascade; create schema public")
        for name in ("001_init.sql", "002_control_plane.sql", "006_ga4_metrics.sql"):
            cur.execute((ROOT / "migrations" / name).read_text(encoding="utf-8"))
        # A schema recreated by `create schema public` has none of the grants initdb's public schema has, and 001 grants
        # usage to app_request only. The worker needs it to see the functions, as it has on a database made by initdb.
        cur.execute("grant usage on schema public to sync_worker")
        ids["a_shop"] = _seed(cur, a, "tenant-a", both_months())
        ids["c_shop"] = _seed(cur, c, "tenant-c", both_months())
        for tenant, label in ((a, "tenant-a"), (b, "tenant-b"), (c, None), (d, "tenant-d"), (e, "tenant-e")):
            cur.execute("select set_config('app.tenant_id', %s, false)", (str(tenant),))
            if tenant in (b, d, e):
                cur.execute("insert into tenant (id, slug, name) values (%s, %s, %s)", (tenant, label, label))
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(a),))
        ids["a_web"] = connect_row(cur, a, "ga4", "properties/" + PROPERTY, "tenant-a site")
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(b),))
        ids["b_web"] = connect_row(cur, b, "ga4", "properties/" + OTHER_PROPERTY, "tenant-b site")
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(d),))
        ids["d_web1"] = connect_row(cur, d, "ga4", "properties/1", "d site one")
        ids["d_web2"] = connect_row(cur, d, "ga4", "properties/2", "d site two")
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(e),))
        ids["e_web"] = connect_row(cur, e, "ga4", "properties/3", "e site")
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(d),))
        ids["d_shop"] = connect_row(cur, d, "shopify", "d.myshopify.com", "d store")
    server.set_days(PROPERTY, WEB)
    server.set_days(OTHER_PROPERTY, {date(2026, 10, 1): {"sessions": 9_999_999, "totalUsers": 9_000_000,
                                                          "screenPageViews": 99_999_999, "keyEvents": 999_999}})
    return {"a": a, "b": b, "c": c, "d": d, "e": e, **ids}


def worker_db():
    return Database(DSN, max_size=2, role="sync_worker")


def request_db():
    return Database(DSN, max_size=2)


def connector(server, property_id=PROPERTY, token=TOKEN, **kw):
    client = httpx.AsyncClient(trust_env=False, verify=False, timeout=10)

    async def no_sleep(seconds):
        return None

    kw.setdefault("sleep", no_sleep)
    return GA4Connector(property_id=property_id, access_token=token, base_url=server.url, client=client, **kw)


async def sync(world, key, tenant_key, server, start, end, property_id=PROPERTY, **kw):
    db = worker_db()
    await db.connect()
    conn = connector(server, property_id, **kw)
    try:
        return await SyncRunner(db).run_once(tenant_id=str(world[tenant_key]), connection_id=str(world[key]),
                                             connector=conn, window_start=start, window_end=end, trigger="backfill")
    finally:
        await conn.aclose()
        await db.close()


def sql(query, params=()):
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute(query, params)
        return cur.fetchall() if cur.description else None


def snapshot_count(connection_id):
    return sql("select count(*) from metric_snapshot where connection_id = %s", (connection_id,))[0][0]


# ---------------------------------------------------------------- the sync

def test_a_sync_stores_four_metrics_a_day_as_integers_with_no_currency_and_records_the_run(world, server):
    out = run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 10, 1)))
    assert out.status == "succeeded" and out.claimed and out.rows_written == 8 and out.api_calls == 1
    rows = sql("select metric_key, observed_on, value_numeric, currency from metric_snapshot "
               "where connection_id = %s order by observed_on, metric_key", (world["a_web"],))
    assert len(rows) == 8 and all(r[3] is None for r in rows)
    assert {r[0] for r in rows} == {"sessions", "users", "page_views", "key_events"}
    assert {(r[0], r[1].day, r[2]) for r in rows if r[1].day == 1} == {
        ("sessions", 1, 200), ("users", 1, 150), ("page_views", 1, 600), ("key_events", 1, 10)}
    run_row = sql("select status, window_start, window_end, rows_written, api_calls, error_code from sync_run "
                  "where connection_id = %s", (world["a_web"],))
    assert run_row == [("succeeded", date(2026, 9, 1), date(2026, 10, 1), 8, 1, None)]


def test_syncing_the_same_window_again_rewrites_the_same_rows_and_never_adds_more(world, server):
    before = snapshot_count(world["a_web"])
    stamp = sql("select max(synced_at) from metric_snapshot where connection_id = %s", (world["a_web"],))[0][0]
    out = run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 10, 1)))
    assert out.status == "succeeded" and out.rows_written == 8
    assert snapshot_count(world["a_web"]) == before == 8
    assert sql("select max(synced_at) from metric_snapshot where connection_id = %s", (world["a_web"],))[0][0] > stamp


def test_a_figure_google_restates_replaces_the_stored_one_and_touches_nothing_else(world, server):
    server.set_days(PROPERTY, {**WEB, date(2026, 9, 2): {**WEB[date(2026, 9, 2)], "sessions": 130}})
    try:
        run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 10, 1)))
        got = dict(sql("select metric_key, value_numeric from metric_snapshot where connection_id = %s "
                       "and observed_on = '2026-09-02'", (world["a_web"],)))
        assert got == {"sessions": 130, "users": 80, "page_views": 250, "key_events": 5}
        assert snapshot_count(world["a_web"]) == 8
    finally:
        server.set_days(PROPERTY, WEB)
        run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 10, 1)))
    assert dict(sql("select metric_key, value_numeric from metric_snapshot where connection_id = %s "
                    "and observed_on = '2026-09-02'", (world["a_web"],)))["sessions"] == 100


@pytest.mark.xfail(strict=True, reason="known limit: a day Google later reports as all zero sends no row, so the old figures stay")
def test_a_day_google_later_leaves_out_keeps_its_old_figures_which_is_the_known_limit(world, server):
    server.set_days(PROPERTY, {d: v for d, v in WEB.items() if d != date(2026, 9, 2)})
    try:
        run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 10, 1)))
        assert sql("select count(*) from metric_snapshot where connection_id = %s and observed_on = '2026-09-02'",
                   (world["a_web"],))[0][0] == 0
    finally:
        server.set_days(PROPERTY, WEB)
        run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 10, 1)))


def test_the_october_windows_leave_the_5th_and_6th_unsynced(world, server):
    for start, end in WINDOWS[1:]:
        assert run(sync(world, "a_web", "a", server, start, end)).status == "succeeded"
    covered = sql("select window_start, window_end from sync_run where connection_id = %s and status = 'succeeded' "
                  "order by window_start, started_at", (world["a_web"],))
    assert (date(2026, 10, 5), date(2026, 10, 7)) not in covered
    assert (date(2026, 10, 1), date(2026, 10, 5)) in covered and (date(2026, 10, 7), date(2026, 11, 1)) in covered


def test_a_rejected_token_is_a_failed_run_with_its_code_and_writes_nothing(world, server):
    before = snapshot_count(world["a_web"])
    ok = run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 9, 3)))        # positive control
    assert ok.status == "succeeded"
    bad = run(sync(world, "a_web", "a", server, date(2026, 11, 1), date(2026, 11, 3), token="wrong"))
    assert bad.status == "failed" and bad.error_code == "auth_expired" and bad.rows_written == 0
    assert "wrong" not in (bad.error_detail or "") and TOKEN not in (bad.error_detail or "")
    assert snapshot_count(world["a_web"]) == before
    row = sql("select status, error_code, error_detail from sync_run where connection_id = %s and window_start = '2026-11-01'",
              (world["a_web"],))[0]
    assert row[0] == "failed" and row[1] == "auth_expired" and "wrong" not in (row[2] or "")


def test_a_property_the_token_cannot_read_is_a_failed_run_that_says_to_reconnect(world, server):
    bad = run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 9, 3), property_id="999"))
    assert bad.status == "failed" and bad.error_code == "auth_revoked"


def test_quota_exhaustion_fails_the_run_as_rate_limited_and_the_next_run_succeeds(world, server):
    server.faults.extend([Fault(429, QUOTA_BODY)] * 3)
    failed = run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 9, 3), max_attempts=3))
    assert failed.status == "failed" and failed.error_code == "rate_limited" and failed.api_calls == 0
    assert "gave up after 3 attempts" in failed.error_detail
    assert not server.faults
    again = run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 9, 3)))
    assert again.status == "succeeded"


def test_progress_is_committed_per_page_so_a_stop_halfway_keeps_the_first_page_and_its_cursor(world, server):
    sql("delete from metric_snapshot where connection_id = %s and observed_on = '2026-09-02'", (world["a_web"],))
    server.tokens_per_hour = 101                     # the first answer leaves 100, the floor, so page two is not asked for
    try:
        out = run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 9, 3), page_size=1, quota_floor=100))
    finally:
        server.tokens_per_hour = 40_000
    assert out.status == "failed" and out.error_code == "rate_limited" and out.rows_written == 4
    cursor = sql("select resume_cursor from sync_run where connection_id = %s order by started_at desc limit 1",
                 (world["a_web"],))[0][0]
    assert cursor == "1"
    assert sql("select count(*) from metric_snapshot where connection_id = %s and observed_on = '2026-09-01'",
               (world["a_web"],))[0][0] == 4
    assert sql("select count(*) from metric_snapshot where connection_id = %s and observed_on = '2026-09-02'",
               (world["a_web"],))[0][0] == 0, "the second page was never asked for"
    run(sync(world, "a_web", "a", server, date(2026, 9, 1), date(2026, 9, 3)))          # the next run finishes the job
    assert sql("select count(*) from metric_snapshot where connection_id = %s and observed_on = '2026-09-02'",
               (world["a_web"],))[0][0] == 4


# ---------------------------------------------------------------- the database's own rules

def test_the_schema_refuses_a_currency_on_a_count_and_a_metric_nobody_registered(world):
    def insert(metric, currency):
        with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
            cur.execute("select set_config('app.tenant_id', %s, false)", (str(world["a"]),))
            cur.execute(
                "insert into metric_snapshot (tenant_id, connection_id, metric_key, observed_on, value_numeric, currency) "
                "values (%s, %s, %s, '2026-09-20', 5, %s)", (world["a"], world["a_web"], metric, currency))

    with pytest.raises(psycopg.errors.CheckViolation):
        insert("sessions", "USD")
    with pytest.raises(psycopg.errors.ForeignKeyViolation):
        insert("conversions", None)
    insert("sessions", None)                                                       # the control: the same row, no currency
    sql("delete from metric_snapshot where connection_id = %s and observed_on = '2026-09-20'", (world["a_web"],))


def test_the_migration_is_safe_to_run_twice_and_registers_four_counts():
    sql((ROOT / "migrations" / "006_ga4_metrics.sql").read_text(encoding="utf-8"))
    got = {k: (u, kind) for k, u, kind in sql("select key, unit, kind from metric where key in "
                                              "('sessions','users','page_views','key_events')")}
    assert got == {"sessions": ("count", "delta"), "users": ("count", "gauge"),
                   "page_views": ("count", "delta"), "key_events": ("count", "delta")}


def test_a_tenant_sees_only_its_own_website_rows_and_the_other_tenants_huge_figures_never_cross(world, server):
    run(sync(world, "b_web", "b", server, date(2026, 10, 1), date(2026, 10, 2), property_id=OTHER_PROPERTY))

    def as_tenant(tenant):
        with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
            cur.execute("set role app_request")
            if tenant:
                cur.execute("select set_config('app.tenant_id', %s, false)", (str(tenant),))
            cur.execute("select connection_id, value_numeric from metric_snapshot where metric_key = 'sessions'")
            return cur.fetchall()

    a_rows, b_rows = as_tenant(world["a"]), as_tenant(world["b"])
    assert b_rows and {r[0] for r in b_rows} == {world["b_web"]} and max(r[1] for r in b_rows) == 9_999_999
    assert world["b_web"] not in {r[0] for r in a_rows} and max(r[1] for r in a_rows) < 9_999_999
    assert as_tenant(None) == [], "no tenant set means no rows, not an error"
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("set role app_request")
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(world["a"]),))
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            cur.execute(
                "insert into metric_snapshot (tenant_id, connection_id, metric_key, observed_on, value_numeric) "
                "values (%s, %s, 'sessions', '2026-10-09', 1)", (world["b"], world["b_web"]))


# ---------------------------------------------------------------- out through the report

async def _generate(world, tenant_key, tmp_path, **kw):
    db = request_db()
    await db.connect()
    try:
        return await generate(db, str(world[tenant_key]), 2026, 10, tmp_path / "deck.pptx", as_of=AS_OF, **kw)
    finally:
        await db.close()


def test_from_google_to_the_deck_the_figures_are_the_ones_computed_by_hand(world, server, tmp_path):
    # the October rows for 1 and 2 arrive through the same connector and runner
    for start, end in WINDOWS:
        run(sync(world, "a_web", "a", server, start, end))
    path, report, insights = run(_generate(world, "a", tmp_path, connection_id=str(world["a_shop"])))
    prs = Presentation(str(path))
    assert len(prs.slides) == 8
    assert prs.slides[5].shapes.title.text_frame.text == TITLE
    assert report.web.kpi("sessions").current_display == "410" and report.web.kpi("page_views").current_display == "1,100"
    assert report.web.current.missing_days == [date(2026, 10, 5), date(2026, 10, 6)]
    assert report.kpi("net_sales").current_display == "$210.00", "the store half is the store's own"
    assert "9,999,999" not in dump_deck(path) and "99,999,999" not in dump_deck(path)


def test_a_failed_website_run_over_a_hole_is_not_coverage_so_those_days_stay_missing(world, server, tmp_path):
    for start, end in WINDOWS:
        run(sync(world, "a_web", "a", server, start, end))
    run(sync(world, "a_web", "a", server, date(2026, 10, 5), date(2026, 10, 7), token="wrong"))   # fails, covers the hole
    inputs = run(_inputs(world, "a"))
    assert (date(2026, 10, 5), date(2026, 10, 7)) not in inputs.web.covered
    _, report, _ = run(_generate(world, "a", tmp_path))
    assert report.web.current.missing_days == [date(2026, 10, 5), date(2026, 10, 6)]


async def _inputs(world, tenant_key, **kw):
    db = request_db()
    await db.connect()
    try:
        return await load_inputs(db, str(world[tenant_key]), 2026, 10, **kw)
    finally:
        await db.close()


def test_a_tenant_with_no_website_gets_the_store_only_deck_with_seven_slides(world, tmp_path):
    path, report, _ = run(_generate(world, "c", tmp_path))
    assert report.web is None and len(Presentation(str(path)).slides) == 7
    assert run(_inputs(world, "c")).web is None


def test_two_websites_with_no_name_are_refused_and_naming_one_works(world, tmp_path):
    with pytest.raises(ReportError, match="2 active website connections, name one with --web-connection"):
        run(_inputs(world, "d"))
    inputs = run(_inputs(world, "d", web_connection_id=str(world["d_web1"])))
    assert inputs.web is not None and inputs.web.label == "d site one" and inputs.label == "d store"


def test_a_website_connection_id_from_another_tenant_or_a_store_is_not_found(world):
    with pytest.raises(ReportError, match="no active website connection found with that id"):
        run(_inputs(world, "a", web_connection_id=str(world["b_web"])))
    with pytest.raises(ReportError, match="no active website connection found with that id"):
        run(_inputs(world, "a", web_connection_id=str(world["a_shop"])))
    with pytest.raises(ReportError, match="no active connection"):
        run(_inputs(world, "a", connection_id=str(world["a_web"])))


def test_a_tenant_with_only_a_website_has_no_store_to_report_on(world):
    with pytest.raises(ReportError, match="no active connection found for this tenant"):
        run(_inputs(world, "e"))


def _cli(args, extra_env=None):
    env = dict(os.environ, DATABASE_URL=DSN, **(extra_env or {}))
    for name in ("GROQ_API_KEY", "LLM_API_KEY", "GA4_ACCESS_TOKEN", "GA4_PROPERTY_ID", "GA4_API_BASE_URL"):
        if not extra_env or name not in extra_env:
            env.pop(name, None)
    return subprocess.run([sys.executable, str(ROOT / "scripts" / args[0]), *args[1:]], capture_output=True, text=True,
                          encoding="utf-8", timeout=90, env=env, cwd=ROOT)


def test_the_report_command_prints_that_the_website_is_in_the_deck(world, server, tmp_path):
    for start, end in WINDOWS:
        run(sync(world, "a_web", "a", server, start, end))
    done = _cli(["monthly_report.py", "--month", "2026-10", "--tenant", str(world["a"]), "--connection", str(world["a_shop"]),
                 "--web-connection", str(world["a_web"]), "--out", str(tmp_path / "cli.pptx"), "--as-of", AS_OF.isoformat(),
                 "--dump"])
    assert done.returncode == 0, done.stderr
    assert "(October 2026 with website figures, insights: template)" in done.stdout
    assert TITLE in done.stdout
    bad = _cli(["monthly_report.py", "--month", "2026-10", "--tenant", str(world["a"]), "--web-connection", str(world["b_web"]),
                "--out", str(tmp_path / "x.pptx"), "--as-of", AS_OF.isoformat()])
    assert bad.returncode == 2 and "no report: no active website connection found with that id" in bad.stderr


def test_the_sync_command_runs_against_the_stand_in_and_never_prints_the_token(world, server):
    env = {"GA4_PROPERTY_ID": PROPERTY, "GA4_ACCESS_TOKEN": TOKEN, "GA4_API_BASE_URL": server.url}
    done = _cli(["sync_ga4.py", "--tenant", str(world["a"]), "--connection", str(world["a_web"]), "--days", "3",
                 "--end", "2026-09-04"], env)
    assert done.returncode == 0, done.stderr + done.stdout
    assert "sync succeeded: window 2026-09-01 to 2026-09-04, 8 rows written, 1 api calls" in done.stdout
    assert TOKEN not in done.stdout + done.stderr
    wrong = _cli(["sync_ga4.py", "--tenant", str(world["b"]), "--connection", str(world["a_web"]), "--days", "3",
                  "--end", "2026-09-04"], env)
    assert wrong.returncode == 2 and "no active ga4 connection with that id" in wrong.stderr
    store = _cli(["sync_ga4.py", "--tenant", str(world["a"]), "--connection", str(world["a_shop"]), "--days", "3"], env)
    assert store.returncode == 2
    missing = _cli(["sync_ga4.py", "--tenant", str(world["a"]), "--connection", str(world["a_web"])],
                   {"GA4_PROPERTY_ID": PROPERTY})
    assert missing.returncode == 2 and "GA4_ACCESS_TOKEN" in missing.stderr and "Traceback" not in missing.stderr
    failed = _cli(["sync_ga4.py", "--tenant", str(world["a"]), "--connection", str(world["a_web"]), "--days", "3",
                   "--end", "2026-09-04"], {**env, "GA4_ACCESS_TOKEN": "wrong-one"})
    assert failed.returncode == 1 and "error auth_expired" in failed.stdout and "wrong-one" not in failed.stdout + failed.stderr
