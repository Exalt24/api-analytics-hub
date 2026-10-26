"""The deck with ads and search on it, saved and read back, and the narrative check over the new figures.

A store alone, or with a website, renders exactly as it did before (tests/test_report_web_deck.py holds the saved
read-back of the store-only sample and the website's tests). Here: the slides that appear with ads and with search
and in what order, the summary table that carries a headline row for each, the paid media slide and the search
slide read back through python-pptx (tables, chart series in the ad account's currency, gaps for days with no
data, top rows, the footnotes), the data health slide with one row for each source, plain words and no text under
11 points, and the figure check extended to the new metrics from both sides: a figure beside its own metric in
its own currency passes, a figure beside another metric, in another currency or not computed is refused.
"""
from __future__ import annotations

import asyncio
import json
import re
import sys
from datetime import date
from pathlib import Path

import httpx
import pytest
from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from app.report.deck import dump_deck, save_deck  # noqa: E402
from app.report.insights import (  # noqa: E402
    ReportInsightWriter,
    check_model_text,
    report_facts,
    template_insights,
    ungrounded_numbers,
    write_insights,
)
from app.report.monthly import build_report  # noqa: E402
from app.report.paid import attach_paid  # noqa: E402
from app.report.search import attach_search  # noqa: E402
from app.report.web import attach_web  # noqa: E402
from sources_fixtures import (  # noqa: E402
    AS_OF,
    COVERED,
    PAID_COVERED,
    SEARCH_COVERED,
    both_months,
    paid_both,
    search_both,
    top_october,
)
from test_report_web import WEB_COVERED, WEB_ROWS  # noqa: E402


def run(coro):
    return asyncio.run(coro)


def report_with(paid=True, search=True, web=False, *, paid_rows=None, tops=None, as_of=AS_OF, covered=COVERED, paid_covered=PAID_COVERED,
                search_covered=SEARCH_COVERED):
    report = build_report(both_months(), 2026, 10, label="Example Outfitters", covered=covered, as_of=as_of)
    if web:
        attach_web(report, WEB_ROWS, label="Example site", covered=WEB_COVERED, as_of=as_of)
    if paid:
        attach_paid(report, paid_both() if paid_rows is None else paid_rows, label="Example ads", covered=paid_covered, as_of=as_of)
    if search:
        attach_search(report, search_both(), top_october() if tops is None else tops, label="Example search",
                      covered=search_covered, as_of=as_of)
    return report


def saved(tmp_path, report=None, **kw):
    report = report or report_with(**kw)
    insights = template_insights(report)
    path = save_deck(report, insights, tmp_path / "deck.pptx")
    return Presentation(str(path)), report, insights, path


def titles(prs):
    return [s.shapes.title.text_frame.text for s in prs.slides]


def tables(slide):
    return [[[c.text for c in row.cells] for row in s.table.rows] for s in slide.shapes if s.has_table]


def charts(slide):
    return [s for s in slide.shapes if s.has_chart]


def texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and not s.has_chart]


def all_text(prs):
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                out.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(out)


# ---------------------------------------------------------------- which slides, in what order

@pytest.mark.parametrize("kw,count,after_value", [
    ({"paid": True, "search": False}, 8, ["ad spend"]),
    ({"paid": False, "search": True}, 8, ["search clicks"]),
    ({"paid": True, "search": True}, 9, ["ad spend", "search clicks"]),
    ({"paid": True, "search": True, "web": True}, 10, ["sessions", "ad spend", "search clicks"]),
])
def test_each_added_source_adds_its_own_slide_after_the_order_value_slide_and_before_the_narrative(tmp_path, kw, count, after_value):
    prs, *_ = saved(tmp_path, **kw)
    got = titles(prs)
    assert len(got) == count
    assert got[0] == "Monthly performance report" and got[2] == "Net sales by day in October"
    value_at = next(i for i, t in enumerate(got) if t.startswith("Order value, buyers and best days"))
    middle = got[value_at + 1: -2]
    assert len(middle) == len(after_value) and all(word in title.lower() for word, title in zip(after_value, middle))
    assert got[-2].startswith(("October", "Net sales", "Refunds")) and got[-1].startswith("Data health")


def test_without_a_website_the_orders_chart_stays_on_its_own_slide(tmp_path):
    prs, *_ = saved(tmp_path, paid=True, search=True)
    assert titles(prs)[3] == "Orders by day in October" and len(charts(prs.slides[3])) == 1


def test_a_store_alone_and_a_store_with_only_a_website_keep_their_slide_counts_and_their_old_health_slide(tmp_path):
    alone, *_ = saved(tmp_path, paid=False, search=False)
    web, *_ = saved(tmp_path, paid=False, search=False, web=True)
    assert len(alone.slides) == 7 and len(web.slides) == 8
    assert titles(alone)[-1].startswith("Data health: October is missing") and "Source, October" not in all_text(alone)
    assert "Store, October" in all_text(web) and "Source, October" not in all_text(web)


# ---------------------------------------------------------------- the summary table

def test_the_summary_table_gains_a_headline_row_for_each_added_source_and_keeps_the_stores_rows(tmp_path):
    prs, report, *_ = saved(tmp_path)
    rows = tables(prs.slides[1])[0]
    assert [r[0] for r in rows] == ["Metric", "Orders", "Gross sales", "Refunds", "Net sales", "Average order value",
                                    "Buyers per active day", "Ad spend", "Return on ad spend", "Search clicks", "Average position"]
    assert rows[7][:5] == ["Ad spend", "EUR 180.00", "EUR 150.00", "+EUR 30.00", "+20.0%"]
    assert rows[8][:5] == ["Return on ad spend", "4.50", "3.33", "+1.17", "+35.0%"]
    assert rows[9][:5] == ["Search clicks", "20", "30", "-10", "-33.3%"]
    assert rows[10][:5] == ["Average position", "6.0", "6.5", "-0.5", "-7.7%"]
    plain = tables(saved(tmp_path, paid=False, search=False)[0].slides[1])[0]
    assert rows[:7] == plain[:7], "the store's own rows are not touched by the extra ones"


def test_the_extra_rows_name_the_days_behind_them_and_how_many_days_carry_a_purchase_figure(tmp_path):
    rows = tables(saved(tmp_path)[0].slides[1])[0]
    assert rows[7][5] == "October: 2 days missing; per reported day up 24.1%"
    assert rows[8][5] == "October: 2 days missing; on 1 of 29 days"
    assert rows[10][5] == "October: 2 days missing"


def test_the_summary_footer_says_which_currency_and_window_the_ad_rows_are_in(tmp_path):
    prs, *_ = saved(tmp_path)
    assert "Ad rows: Meta, in the ad account's currency (EUR), purchases in a 7 day click window. Search rows: Search Console, " \
           "Pacific time, a lower position is nearer the top." in texts(prs.slides[1])[-1]
    only_paid = saved(tmp_path, search=False)[0]
    footer = texts(only_paid.slides[1])[-1]
    assert "Ad rows: Meta" in footer and "Search rows" not in footer


def test_a_provisional_figure_says_so_in_its_note(tmp_path):
    report = report_with(as_of=date(2026, 11, 2), covered=COVERED, paid_covered=[(date(2026, 9, 1), date(2026, 11, 1))],
                         search_covered=[(date(2026, 9, 1), date(2026, 11, 1))])
    rows = tables(saved(tmp_path, report)[0].slides[1])[0]
    assert "last 6 days provisional" in rows[8][5]            # the return on ad spend: a purchase can still arrive
    assert "provisional" not in rows[7][5], "ad spend itself is not flagged"
    assert "last 2 days provisional" in rows[9][5] and "last 2 days provisional" in rows[10][5]


# ---------------------------------------------------------------- the paid media slide

def test_the_paid_slide_holds_the_hand_computed_figures_in_the_ad_accounts_currency(tmp_path):
    prs, *_ = saved(tmp_path)
    slide = prs.slides[5]
    assert titles(prs)[5] == ("October ad spend total is 20.0% above September, but October is missing 2 days of ad data: "
                              "per reported day it is up 24.1%")
    assert tables(slide)[0] == [
        ["Metric", "October", "September", "Change %"],
        ["Ad spend", "EUR 180.00", "EUR 150.00", "+20.0%"],
        ["Impressions", "1,600", "1,500", "+6.7%"],
        ["Clicks", "70", "70", "0.0%"],
        ["Click through rate", "4.38%", "4.67%", "-6.3%"],
        ["Cost per click", "EUR 2.57", "EUR 2.14", "+20.0%"],
        ["Purchases", "6", "7", "-14.3%"],
        ["Purchase value", "EUR 540.00", "EUR 500.00", "+8.0%"],
        ["Return on ad spend", "4.50", "3.33", "+35.0%"],
    ]
    assert "USD" not in all_text(prs) and "$" not in " ".join(sum(tables(slide)[0], []))


def test_the_spend_chart_is_in_major_units_of_the_ad_currency_with_a_gap_for_the_days_nobody_looked(tmp_path):
    prs, *_ = saved(tmp_path)
    spend, clicks = [s.chart for s in charts(prs.slides[5])]
    series = {s.name: list(s.values) for s in spend.plots[0].series}
    assert series["October 2026"][:3] == [120.0, 60.0, 0.0] and series["October 2026"][9] is None and series["October 2026"][10] is None
    assert series["Day with no data"][9] == 0.0 and series["Day with no data"][10] == 0.0
    assert spend.value_axis.axis_title.text_frame.text == "Ad spend (EUR)"
    assert spend.value_axis.tick_labels.number_format == "#,##0", "an euro axis is not given a dollar sign"
    assert clicks.value_axis.axis_title.text_frame.text == "Clicks"
    assert {s.name: list(s.values) for s in clicks.plots[0].series}["October 2026"][:3] == [60.0, 10.0, 0.0]
    assert spend.has_legend and not clicks.has_legend


def test_a_dollar_ad_account_gets_a_dollar_axis_and_dollar_figures(tmp_path):
    rows = [dict(r, currency="USD" if r["currency"] else None) for r in paid_both()]
    prs, *_ = saved(tmp_path, paid_rows=rows)
    spend = charts(prs.slides[5])[0].chart
    assert spend.value_axis.tick_labels.number_format == '"$"#,##0' and spend.value_axis.axis_title.text_frame.text == "Ad spend ($)"
    assert tables(prs.slides[5])[0][1][1] == "$180.00"


def test_the_paid_slide_footnotes_say_where_the_figures_come_from_and_what_is_not_the_stores(tmp_path):
    prs, *_ = saved(tmp_path)
    text = "\n".join(texts(prs.slides[5]))
    assert "Source: Meta Marketing API, in the ad account's currency (EUR) and time zone; store days follow UTC. " \
           "Purchases are Meta's own count in a 7 day click window, not the store's orders." in text
    assert "A dot on the zero line is a day with no ad delivery, on either line; a break with an X is a day with no data (Oct 10, Oct 11)." in text
    assert "Per reported day divides a total by the days with ad delivery plus the days with none." in text


def test_the_paid_alt_text_names_the_metric_and_the_month(tmp_path):
    prs, *_ = saved(tmp_path)
    alts = [c._element.nvGraphicFramePr.cNvPr.get("descr") for c in charts(prs.slides[5])]
    assert alts == ["Line chart of ad spend for each day of October 2026, with September 2026 as a grey line",
                    "Line chart of ad clicks for each day of October 2026, with September 2026 as a grey line"]


def test_provisional_ad_days_are_named_on_the_paid_slide(tmp_path):
    report = report_with(search=False, as_of=date(2026, 11, 2), paid_covered=[(date(2026, 9, 1), date(2026, 11, 1))])
    text = "\n".join(texts(saved(tmp_path, report)[0].slides[5]))
    assert "Provisional days (a click can lead to a purchase after the day is over): Oct 26, Oct 27, Oct 28, Oct 29, Oct 30, Oct 31." in text


def test_an_ad_account_with_only_one_month_stored_has_a_plain_headline_and_no_grey_line(tmp_path):
    rows = [r for r in paid_both() if r["observed_on"].month == 10]
    prs, *_ = saved(tmp_path, search=False, paid_rows=rows, paid_covered=[(date(2026, 10, 1), date(2026, 11, 1))])
    assert titles(prs)[5] == "Ad spend came to EUR 180.00 in October"
    assert [s.name for s in charts(prs.slides[5])[0].chart.plots[0].series] == ["October 2026"]


# ---------------------------------------------------------------- the search slide

def test_the_search_slide_holds_the_hand_computed_figures_and_the_top_rows(tmp_path):
    prs, *_ = saved(tmp_path, paid=False)
    slide = prs.slides[5]
    assert titles(prs)[5] == ("October search clicks total is 33.3% below September, but October is missing 2 days of search data: "
                              "per reported day it is down 31.0%")
    kpi, queries, pages = tables(slide)
    assert kpi == [["Metric", "Oct", "Sep", "Chg %"], ["Search clicks", "20", "30", "-33.3%"], ["Search impressions", "200", "300", "-33.3%"],
                   ["Search click through rate", "10.00%", "10.00%", "0.0%"], ["Average position", "6.0", "6.5", "-7.7%"]]
    assert queries == [["Query", "Clicks", "Impr.", "Pos."], ["shoes", "9", "80", "5.5"], ["socks", "5", "30", "4.0"], ["boots", "4", "40", "10.0"]]
    assert pages == [["Page", "Clicks", "Impr.", "Pos."], ["/a", "14", "140", "6.7"], ["/b", "4", "10", "2.0"]]
    clicks, impressions = [c.chart for c in charts(slide)]
    assert {s.name: list(s.values) for s in clicks.plots[0].series}["October 2026"][:2] == [12.0, 8.0]
    assert {s.name: list(s.values) for s in impressions.plots[0].series}["September 2026"][:2] == [100.0, 200.0]
    assert clicks.has_legend and not impressions.has_legend


def test_a_month_with_no_top_rows_stored_says_so_in_the_table(tmp_path):
    prs, *_ = saved(tmp_path, paid=False, tops=[])
    _, queries, pages = tables(prs.slides[5])
    assert queries[1] == ["no rows stored", "", "", ""] and pages[1] == ["no rows stored", "", "", ""]


def test_the_search_footnotes_say_the_top_rows_are_lower_bounds_and_which_days_are_provisional(tmp_path):
    report = report_with(paid=False, as_of=date(2026, 11, 2), search_covered=[(date(2026, 9, 1), date(2026, 11, 1))])
    text = "\n".join(texts(saved(tmp_path, report)[0].slides[5]))
    assert "Top queries and pages add up each day's top rows over 2 days, so a key outside a day's top rows is not counted: " \
           "they are lower bounds, not totals." in text
    assert "Provisional days (Search Console is still collecting them): Oct 30, Oct 31." in text
    assert "Search days are Pacific time, store days follow UTC. A lower average position is nearer the top of the results." in text


def test_a_long_query_is_cut_in_the_cell_and_whole_in_the_stored_figures(tmp_path):
    from sources_fixtures import top
    long_key = "a very long query that people really did type into a search box one day"
    prs, report, *_ = saved(tmp_path, paid=False, tops=[top(date(2026, 10, 1), "query", long_key, 5, 9, 4000)])
    cell = tables(prs.slides[5])[1][1][0]
    assert cell.endswith("...") and len(cell) <= 30 and report.search.top_queries[0].key == long_key


# ---------------------------------------------------------------- the data health slide

def test_the_health_slide_has_a_row_for_each_source_and_names_the_days_with_no_data_for_each(tmp_path):
    prs, *_ = saved(tmp_path, web=True)
    slide = prs.slides[-1]
    assert titles(prs)[-1] == "Data health: October is missing days of store, website, ads and search data"
    table = tables(slide)[0]
    assert table[0] == ["Source, October", "Days with data", "Days with nothing to report", "Days with no data",
                        "Days that may still change", "Stored daily figures used"]
    assert [row[0] for row in table[1:]] == ["Store", "Website", "Ads", "Search"]
    assert table[3] == ["Ads", "2", "27", "2", "0", "8"] and table[4] == ["Search", "2", "27", "2", "0", "6"]
    assert table[1][4] == "n/a" and table[2][4] == "n/a", "the store and the website have no provisional rule, so not a zero"
    text = "\n".join(texts(slide))
    for source in ("store", "ads", "search"):
        assert "Days with no data in October, %s: Oct 10, Oct 11." % source in text
    assert "Days with no data in October, website:" in text


def test_the_days_of_every_source_add_up_to_the_days_in_the_month(tmp_path):
    prs, report, *_ = saved(tmp_path, web=True)
    for _, row in zip(range(4), tables(prs.slides[-1])[0][1:]):
        assert int(row[1]) + int(row[2]) + int(row[3]) == 31


def test_the_health_title_names_only_the_sources_that_are_missing_days(tmp_path):
    full = [(date(2026, 9, 1), date(2026, 11, 1))]
    only_ads = report_with(covered=full, search_covered=full)
    assert titles(saved(tmp_path, only_ads)[0])[-1] == "Data health: October is missing days of ads data"
    none = report_with(covered=full, paid_covered=full, search_covered=full)
    assert titles(saved(tmp_path, none)[0])[-1] == "Data health: every day of October is covered in all 3 sources"
    two = report_with(covered=full)
    assert titles(saved(tmp_path, two)[0])[-1] == "Data health: October is missing days of ads and search data"


def test_a_month_in_progress_says_how_many_days_are_counted_and_lists_provisional_days(tmp_path):
    mid = report_with(as_of=date(2026, 10, 12), covered=[(date(2026, 9, 1), date(2026, 10, 12))],
                      paid_covered=[(date(2026, 9, 1), date(2026, 10, 12))], search_covered=[(date(2026, 9, 1), date(2026, 10, 12))])
    prs, *_ = saved(tmp_path, mid)
    assert titles(prs)[-1] == "Data health: 11 of 31 days of October are counted so far"
    text = "\n".join(texts(prs.slides[-1]))
    assert "October is not over: 11 of 31 days are counted." in text and "May still change, ads: Oct 5" in text
    assert "May still change, search: Oct 9, Oct 10, Oct 11." in text


# ---------------------------------------------------------------- plain words and sizes

def test_the_deck_uses_plain_words_no_iso_date_and_no_text_under_eleven_points(tmp_path):
    prs, *_ = saved(tmp_path, web=True)
    everything = all_text(prs)
    assert not re.search(r"\d{4}-\d{2}-\d{2}", everything) and "USD" not in everything
    for word in ("snapshot", "sync", "succeeded", "runReport", "keyEvents", "metric_key", "searchAnalytics", "omni_purchase", "7d_click",
                 "action_values", "dataState", "first_incomplete", "ad_spend", "search_clicks", "act_"):
        assert word.lower() not in everything.lower(), word
    sizes = []
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for frame in frames:
                sizes += [run.font.size for para in frame.paragraphs for run in para.runs]
    assert sizes and all(size is not None and size >= Pt(11) for size in sizes)


def test_every_shape_stays_inside_the_slide(tmp_path):
    prs, *_ = saved(tmp_path, web=True)
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (number, shape.name)
            assert shape.left + shape.width <= prs.slide_width + 10 and shape.top + shape.height <= prs.slide_height + 10, (number, shape.name)


def test_the_deck_file_still_reads_back_with_its_provenance_line(tmp_path):
    prs, report, insights, path = saved(tmp_path)
    text = dump_deck(path)
    assert text.count("--- slide ") == 9
    assert text.endswith("--- provenance (file property, not on a slide): " + insights.label)
    assert prs.core_properties.title == "Monthly performance report, October 2026"


# ---------------------------------------------------------------- the narrative check over the new figures

def test_every_line_the_template_writes_about_ads_and_search_is_made_of_computed_figures():
    report = report_with(web=True)
    facts = report_facts(report)
    insights = template_insights(report)                   # raises if a line quotes a figure that is not a fact
    assert insights.source == "template"
    for line in insights.lines:
        assert ungrounded_numbers(line, facts) == [], line
    joined = "\n".join(insights.lines)
    assert "Ad spend total was EUR 180.00 in October, 20.0% above September (EUR 150.00)" in joined
    assert "Return on ad spend was 4.50 in October, up 35.0% from September (3.33)" in joined
    assert "Search clicks total was 20 in October, 33.3% below September (30)" in joined
    assert "Average position was 6.0 in October, down 7.7% from September (6.5)" in joined
    assert "2 days have no ad data (Oct 10, Oct 11), so the ad totals are understated." in joined
    assert "2 days have no search data (Oct 10, Oct 11), so the search totals are understated." in joined


def test_the_new_facts_are_in_the_same_shape_as_the_stores_and_carry_no_year():
    facts = report_facts(report_with())
    for label in ("Ad spend", "Impressions", "Clicks", "Click through rate", "Cost per click", "Purchases", "Purchase value",
                  "Return on ad spend", "Search clicks", "Search impressions", "Search click through rate", "Average position"):
        assert label in facts and label + " change percent" in facts, label
    assert facts["Ad spend"] == "EUR 180.00" and facts["ad account currency"] == "EUR" and facts["dates with no ad data"] == "Oct 10, Oct 11"
    assert facts["Search click through rate"] == "10.00%" and facts["dates with no search data"] == "Oct 10, Oct 11"
    assert facts["purchase attribution window"] == "7 day click"
    assert not any("2026" in str(v) for v in facts.values())


def test_provisional_dates_are_facts_only_when_there_are_provisional_days():
    plain = report_facts(report_with())
    assert "provisional ad dates" not in plain and "provisional search dates" not in plain
    soon = report_facts(report_with(as_of=date(2026, 11, 2), paid_covered=[(date(2026, 9, 1), date(2026, 11, 1))],
                                    search_covered=[(date(2026, 9, 1), date(2026, 11, 1))]))
    assert soon["provisional ad dates"] == "Oct 26, Oct 27, Oct 28, Oct 29, Oct 30, Oct 31" and soon["provisional search dates"] == "Oct 30, Oct 31"


@pytest.mark.parametrize("sentence", [
    "Ad spend was EUR 180.00 in October.",
    "Ad spend was EUR 180.00 in October, up 20.0% from September (EUR 150.00).",
    "Click through rate was 4.38% in October.",
    "Cost per click was EUR 2.57 in October.",
    "Purchases were 6 in October, down 14.3% from September (7).",
    "Purchase value was EUR 540.00 in October.",
    "Return on ad spend was 4.50 in October.",
    "Clicks were 70 in October.",
    "Impressions were 1,600 in October.",
    "Search clicks were 20 in October, down 33.3% from September (30).",
    "Search impressions were 200 in October.",
    "Search click through rate was 10.00% in October.",
    "Average position was 6.0 in October, down 7.7% from September (6.5).",
    "Purchases for Oct 10 may still rise.",
])
def test_a_new_figure_beside_its_own_metric_in_its_own_currency_passes(sentence):
    assert ungrounded_numbers(sentence, report_facts(report_with())) == []


@pytest.mark.parametrize("sentence,flagged", [
    ("Ad spend was $180.00 in October.", "$180.00"),                    # the ad account's currency is EUR
    ("Ad spend was 180.00 in October.", "180.00"),                       # no currency at all
    ("Ad spend was EUR 181.00 in October.", "EUR 181.00"),               # not computed
    ("Ad spend was $210.00 in October.", "$210.00"),                     # the store's net sales figure
    ("Net sales were EUR 180.00 in October.", "EUR 180.00"),             # the ad spend figure beside the store's metric
    ("Clicks were 20 in October.", "20"),                                # the search clicks figure beside the ad clicks metric
    ("Search clicks were 70 in October.", "70"),                         # the ad clicks figure beside the search metric
    ("Click through rate was 4.4% in October.", "4.4%"),                 # a different rounding
    ("Click through rate was 10.00% in October.", "10.00%"),             # the search rate beside the ad metric
    ("Return on ad spend was 4.5 in October.", "4.5"),
    ("Return on ad spend was 3.00 in October.", "3.00"),                 # what dividing by all the spend would have said
    ("Average position was 6.1 in October.", "6.1"),
    ("Average position was 5.5 in October.", "5.5"),                     # the unweighted mean of 5.0 and 6.0, not a fact
    ("Search impressions were 1,600 in October.", "1,600"),
    ("Purchases for Oct 12 may still rise.", "Oct 12"),                  # not a date the facts hold
    ("Spend rose 25.0% in October.", "25.0%"),
    ("Conversions were 6 in October.", "6"),                             # no such metric in the facts
    ("It rose 20.0% from September.", "20.0%"),                          # a sentence that names no metric cannot carry a figure
])
def test_a_new_figure_beside_the_wrong_metric_in_the_wrong_currency_or_not_computed_is_refused(sentence, flagged):
    assert flagged in ungrounded_numbers(sentence, report_facts(report_with()))


def test_the_stores_own_metrics_still_refuse_each_others_figures_with_ads_and_search_in_the_facts():
    facts = report_facts(report_with(web=True))
    assert ungrounded_numbers("Net sales were $210.00 in October.", facts) == []
    assert "EUR 180.00" in ungrounded_numbers("Net sales were EUR 180.00 in October.", facts)
    assert ungrounded_numbers("Orders were 20 in October.", facts) == []
    assert "20" not in ungrounded_numbers("Search clicks were 20 in October.", facts)
    assert ungrounded_numbers("Gross sales were $240.00 in October.", facts) == []


GOOD_MODEL = "\n".join([
    "Ad spend was EUR 180.00 in October, up 20.0% from September (EUR 150.00).",
    "Return on ad spend was 4.50 in October, up 35.0% from September (3.33).",
    "Search clicks were 20 in October, down 33.3% from September (30).",
    "Average position was 6.0 in October, down 7.7% from September (6.5).",
    "2 days have no ad data, so the ad totals are understated.",
])


def replying(content):
    def handler(request):
        return httpx.Response(200, json={"choices": [{"message": {"content": content}}]})
    return ReportInsightWriter("test-key", client=httpx.AsyncClient(transport=httpx.MockTransport(handler)))


def test_a_model_reply_about_ads_and_search_is_used_when_every_figure_is_a_fact():
    insights = run(write_insights(report_with(), replying(GOOD_MODEL)))
    assert insights.source == "model" and list(insights.lines) == GOOD_MODEL.splitlines()


@pytest.mark.parametrize("old,new", [("EUR 180.00", "$180.00"), ("4.50", "4.5"), ("20.0%", "25.0%"), ("(30)", "(300)"),
                                     ("Search clicks were 20", "Search clicks were 70")])
def test_a_model_reply_with_one_wrong_new_figure_falls_back_to_the_template(old, new):
    bad = GOOD_MODEL.replace(old, new, 1)
    insights = run(write_insights(report_with(), replying(bad)))
    assert insights.source == "template_fallback" and insights.rejected_reason
    assert any(line.startswith("Ad spend total was EUR 180.00") for line in insights.lines)


def test_the_model_is_sent_the_ad_and_search_facts_and_the_prompt_names_the_new_sources():
    seen = []

    def handler(request):
        seen.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"message": {"content": GOOD_MODEL}}]})

    writer = ReportInsightWriter("test-key", client=httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    run(write_insights(report_with(), writer))
    user = seen[0]["messages"][1]["content"]
    assert "Ad spend: EUR 180.00" in user and "ad account currency: EUR" in user and "dates with no search data: Oct 10, Oct 11" in user
    assert "observed_on" not in user and "value_numeric" not in user and "2026" not in user
    prompt = seen[0]["messages"][0]["content"]
    assert "paid media" in prompt and "search" in prompt and "currency" in prompt


def test_check_model_text_refuses_a_conversions_claim_for_ads_because_no_such_metric_is_in_the_facts():
    lines, reason = check_model_text("Conversions were 6 in October.", report_facts(report_with()))
    assert lines == [] and "6" in reason


# ---------------------------------------------------------------- the whole sentences, in each state of a month

def test_the_narrative_lines_for_ads_and_search_are_these_exact_sentences_in_a_partial_month():
    lines = list(template_insights(report_with()).lines)
    assert lines[-6:] == [
        "Ad spend total was EUR 180.00 in October, 20.0% above September (EUR 150.00), but October is missing 2 days of ad data: "
        "per reported day, ad spend is up 24.1% (EUR 6.21 against EUR 5.00).",
        "Return on ad spend was 4.50 in October, up 35.0% from September (3.33), and October is missing 2 days of ad data.",
        "2 days have no ad data (Oct 10, Oct 11), so the ad totals are understated.",
        "Search clicks total was 20 in October, 33.3% below September (30), but October is missing 2 days of search data: "
        "per reported day, search clicks are down 31.0% (0.69 against 1.00).",
        "Average position was 6.0 in October, down 7.7% from September (6.5), and October is missing 2 days of search data.",
        "2 days have no search data (Oct 10, Oct 11), so the search totals are understated.",
    ]


def test_the_narrative_lines_for_a_complete_month_name_no_gap_and_a_single_missing_day_is_singular():
    full = [(date(2026, 9, 1), date(2026, 11, 1))]
    lines = list(template_insights(report_with(covered=full, paid_covered=full, search_covered=full)).lines)
    assert "Ad spend was EUR 180.00 in October, up 20.0% from September (EUR 150.00)." in lines
    assert "Return on ad spend was 4.50 in October, up 35.0% from September (3.33)." in lines
    assert "Search clicks were 20 in October, down 33.3% from September (30)." in lines
    assert not [ln for ln in lines if "no ad data" in ln or "no search data" in ln]
    one = [(date(2026, 9, 1), date(2026, 10, 10)), (date(2026, 10, 11), date(2026, 11, 1))]
    text = "\n".join(template_insights(report_with(paid_covered=one)).lines)
    assert "1 day has no ad data (Oct 10), so the ad totals are understated." in text


def test_a_first_month_and_a_source_with_nothing_stored_are_said_plainly():
    rows = [r for r in paid_both() if r["observed_on"].month == 10]
    lines = list(template_insights(report_with(search=False, paid_rows=rows, paid_covered=[(date(2026, 10, 1), date(2026, 11, 1))])).lines)
    assert "Ad spend was EUR 180.00 in October." in lines and "Return on ad spend was 4.50 in October." in lines
    empty = list(template_insights(report_with(search=False, paid_rows=[], paid_covered=[])).lines)
    assert "No ad data is stored for October yet." in empty
    no_roas = [r for r in paid_both() if r["metric_key"] not in ("ad_purchases", "ad_purchase_value")]
    text = "\n".join(template_insights(report_with(search=False, paid_rows=no_roas)).lines)
    assert "Click through rate was 4.38% in October" in text and "Return on ad spend" not in text, "no return to state: the rate is the second line"
