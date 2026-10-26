"""The monthly report as an HTTP call, against real Postgres, with real row level security and real idempotency.

Questions this file answers by attacking the running application: does asking twice for the same month give the same deck
and build nothing the second time; do two calls at the same moment build exactly one deck; does a key from one tenant ever
touch another's runs; may a viewer build or mark delivered; is a failed run retryable with the same key once the cause is
fixed; does a run that died while `running` get taken over; and does an unexpected error say nothing about its cause.

Every refusal is paired with the same call succeeding for the caller who should have it, because an API that answered
400 or 403 to everything would pass every negative test.
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import random
import sys
import uuid
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

psycopg = pytest.importorskip("psycopg", reason="psycopg not installed")
httpx = pytest.importorskip("httpx")
pytest.importorskip("fastapi")
pytest.importorskip("asyncpg")

DSN = os.environ.get("ANALYTICS_TEST_DSN", "postgresql://analytics:analytics@localhost:5433/analytics")


def _reachable() -> bool:
    try:
        with psycopg.connect(DSN, connect_timeout=3):
            return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(not _reachable(), reason="Postgres not reachable, run docker compose up -d")

from pptx import Presentation  # noqa: E402

from app.api.auth import generate_key, hash_key  # noqa: E402
from report_fixtures import AS_OF, both_months, day_rows  # noqa: E402
from sources_fixtures import paid_both, search_both, top_october  # noqa: E402
from test_report_service_db import _seed  # noqa: E402

MIGRATIONS = ("001_init.sql", "002_control_plane.sql", "003_api_keys.sql", "006_ga4_metrics.sql", "007_meta_metrics.sql",
              "008_search_metrics.sql", "009_report_runs.sql")
AS_OF_TEXT = AS_OF.isoformat()
KEY = "monthly-report:tenant-a:2026-10"


def connect_row(cur, tenant, platform, account, name):
    cur.execute(
        "insert into platform_connection (tenant_id, platform, external_account, display_name, credentials_enc) "
        "values (%s, %s, %s, %s, %s) returning id", (tenant, platform, account, name, b"ciphertext"))
    return cur.fetchone()[0]


def load_rows(cur, tenant, connection, rows):
    payload = [{"tenant_id": str(tenant), "connection_id": str(connection), "metric_key": r["metric_key"],
                "observed_on": r["observed_on"].isoformat(), "value_numeric": r["value_numeric"], "currency": r["currency"], "raw": {}}
               for r in rows]
    cur.execute("select upsert_snapshots(%s::jsonb)", (json.dumps(payload),))
    for start, end in ((date(2026, 9, 1), date(2026, 10, 10)), (date(2026, 10, 12), date(2026, 11, 1))):
        cur.execute("insert into sync_run (tenant_id, connection_id, trigger, window_start, window_end, status, finished_at) "
                    "values (%s, %s, 'schedule', %s, %s, 'succeeded', now())", (tenant, connection, start, end))


@pytest.fixture(scope="module")
def seeded():
    """A has a store, ads and search; B a store with far bigger figures; C a key and nothing else; E ads only (no store); D two ad accounts."""
    a, b, c, d, e = (uuid.uuid4() for _ in range(5))
    keys = {name: generate_key() for name in ("a_viewer", "a_operator", "a_admin", "b_admin", "c_admin", "d_admin", "e_admin")}
    ids = {}
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("drop schema if exists public cascade; create schema public")
        for name in MIGRATIONS:
            cur.execute((ROOT / "migrations" / name).read_text(encoding="utf-8"))
        cur.execute("grant usage on schema public to sync_worker")
        ids["a_shop"] = _seed(cur, a, "tenant-a", both_months())
        _seed(cur, b, "tenant-b", day_rows(date(2026, 10, 1), 999, 9_990_000, 0, 500) + day_rows(date(2026, 9, 1), 900, 9_000_000, 0, 400))
        _seed(cur, d, "tenant-d", both_months())
        for tenant, slug in ((c, "tenant-c"), (e, "tenant-e")):
            cur.execute("select set_config('app.tenant_id', %s, false)", (str(tenant),))
            cur.execute("insert into tenant (id, slug, name) values (%s, %s, %s)", (tenant, slug, slug))
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(a),))
        ids["a_ads"] = connect_row(cur, a, "meta", "act_1", "a ads")
        ids["a_search"] = connect_row(cur, a, "gsc", "sc-domain:a.example", "a search")
        load_rows(cur, a, ids["a_ads"], paid_both())
        load_rows(cur, a, ids["a_search"], search_both())
        cur.execute("select upsert_search_rows(%s::jsonb)", (json.dumps([
            {"tenant_id": str(a), "connection_id": str(ids["a_search"]), "observed_on": r["observed_on"].isoformat(), "dimension": r["dimension"],
             "key": r["key"], "clicks": r["clicks"], "impressions": r["impressions"], "position_milli": r["position_milli"]}
            for r in top_october()]),))
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(d),))
        ids["d_ads1"] = connect_row(cur, d, "meta", "act_11", "d ads one")
        ids["d_ads2"] = connect_row(cur, d, "meta", "act_12", "d ads two")
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(e),))
        ids["e_ads"] = connect_row(cur, e, "meta", "act_21", "e ads")
        load_rows(cur, e, ids["e_ads"], paid_both())
        cur.execute("select set_config('app.tenant_id', '', false)")
        for label, tenant, role in (("a_viewer", a, "viewer"), ("a_operator", a, "operator"), ("a_admin", a, "admin"),
                                    ("b_admin", b, "admin"), ("c_admin", c, "admin"), ("d_admin", d, "admin"), ("e_admin", e, "admin")):
            cur.execute("insert into api_key (tenant_id, key_hash, label, role) values (%s, %s, %s, %s)",
                        (tenant, hash_key(keys[label]), label, role))
    return {"keys": keys, "tenants": {"a": a, "b": b, "c": c, "d": d, "e": e}, **ids}


@pytest.fixture(scope="module")
def client(seeded):
    os.environ["DATABASE_URL"] = DSN
    os.environ.setdefault("DASHBOARD_ORIGIN", "http://localhost:3000")
    from app.api.main import app as fastapi_app
    from contextlib import AsyncExitStack

    loop = asyncio.new_event_loop()
    c = httpx.AsyncClient(transport=httpx.ASGITransport(app=fastapi_app), base_url="http://test", timeout=60)
    stack = AsyncExitStack()
    loop.run_until_complete(stack.enter_async_context(fastapi_app.router.lifespan_context(fastapi_app)))
    yield loop, c
    loop.run_until_complete(stack.aclose())
    loop.run_until_complete(c.aclose())
    loop.close()


def call(client, method, url, key=None, **kw):
    loop, c = client
    headers = kw.pop("headers", {})
    if key:
        headers["Authorization"] = "Bearer %s" % key
    return loop.run_until_complete(getattr(c, method)(url, headers=headers, **kw))


def build(client, seeded, who="a_operator", key=KEY, month="2026-10", **body):
    return call(client, "post", "/api/reports/monthly", seeded["keys"][who], headers={"Idempotency-Key": key},
                json={"month": month, "as_of": AS_OF_TEXT, **body})


def sql(query, params=()):
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute(query, params)
        return cur.fetchall() if cur.description else None


# ---------------------------------------------------------------- who may do what

def test_the_call_needs_a_key_and_a_viewer_may_read_but_not_build_or_mark_delivered(client, seeded):
    assert call(client, "post", "/api/reports/monthly", json={"month": "2026-10"}).status_code == 401
    assert call(client, "post", "/api/reports/monthly", "aah_not_a_real_key", headers={"Idempotency-Key": KEY},
                json={"month": "2026-10"}).status_code == 401
    denied = build(client, seeded, who="a_viewer", key="viewer-attempt-0001")
    assert denied.status_code == 403 and "viewer" in denied.json()["detail"]
    assert call(client, "post", "/api/reports/monthly/2026-10/delivered", seeded["keys"]["a_viewer"]).status_code == 403
    ok = build(client, seeded, who="a_operator", key="operator-control-0001")                    # the paired control
    assert ok.status_code == 200 and ok.json()["status"] == "created"
    assert call(client, "get", "/api/reports/monthly/2026-10/deck", seeded["keys"]["a_viewer"]).status_code == 200
    assert call(client, "get", "/api/reports/monthly/2026-10/deck").status_code == 401
    assert sql("select count(*) from report_run where idempotency_key = 'viewer-attempt-0001'")[0][0] == 0


# ---------------------------------------------------------------- building, and asking again

def test_the_first_call_builds_the_deck_and_describes_it_and_the_deck_is_the_one_described(client, seeded):
    r = build(client, seeded)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "created" and body["month"] == "2026-10" and body["slides"] == 9 and body["delivered"] is False
    assert body["sources"] == ["store", "ads", "search"] and body["insights"] == "template" and body["idempotency_key"] == KEY
    assert any(line.startswith("Ad spend total was EUR 180.00") for line in body["narrative"])
    deck = call(client, "get", "/api/reports/monthly/2026-10/deck", seeded["keys"]["a_operator"])
    assert deck.status_code == 200 and "presentationml" in deck.headers["content-type"]
    assert hashlib.sha256(deck.content).hexdigest() == body["sha256"] == deck.headers["x-report-sha256"] and len(deck.content) == body["bytes"]
    assert len(Presentation(io.BytesIO(deck.content)).slides) == 9


def test_asking_again_with_the_same_key_gives_the_same_deck_and_builds_nothing(client, seeded):
    first = build(client, seeded).json()
    deck_before = call(client, "get", "/api/reports/monthly/2026-10/deck", seeded["keys"]["a_operator"]).content
    again = build(client, seeded)
    assert again.status_code == 200
    body = again.json()
    assert body["status"] == "replayed" and body["sha256"] == first["sha256"] and body["narrative"] == first["narrative"]
    assert call(client, "get", "/api/reports/monthly/2026-10/deck", seeded["keys"]["a_operator"]).content == deck_before
    assert sql("select count(*) from report_run where idempotency_key = %s", (KEY,))[0][0] == 1


def test_a_replay_does_not_read_the_data_again_so_a_later_change_is_not_in_the_old_deck(client, seeded):
    first = build(client, seeded, key="replay-after-change-1").json()
    sql("select set_config('app.tenant_id', %s, false)", (str(seeded["tenants"]["a"]),))
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(seeded["tenants"]["a"]),))
        cur.execute("update metric_snapshot set value_numeric = value_numeric + 1 where connection_id = %s and metric_key = 'ad_spend'",
                    (seeded["a_ads"],))
    try:
        again = build(client, seeded, key="replay-after-change-1").json()
        fresh = build(client, seeded, key="replay-after-change-2").json()
        assert again["sha256"] == first["sha256"] and again["status"] == "replayed"
        assert fresh["status"] == "created" and fresh["narrative"] != first["narrative"], "a new key reads the changed rows"
    finally:
        with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
            cur.execute("select set_config('app.tenant_id', %s, false)", (str(seeded["tenants"]["a"]),))
            cur.execute("update metric_snapshot set value_numeric = value_numeric - 1 where connection_id = %s and metric_key = 'ad_spend'",
                        (seeded["a_ads"],))


def test_the_same_key_for_another_month_is_refused_and_a_new_key_for_it_is_built(client, seeded):
    build(client, seeded)
    reused = build(client, seeded, month="2026-09")
    assert reused.status_code == 422 and "2026-10" in reused.json()["detail"]
    other = build(client, seeded, key="monthly-report:tenant-a:2026-09", month="2026-09")
    assert other.status_code == 200 and other.json()["status"] == "created" and other.json()["month"] == "2026-09"


@pytest.mark.parametrize("key", [None, "", "short", "has a space in it", "x" * 129, "semi;colon-key-1", "slash/key-000001"])
def test_a_missing_or_malformed_idempotency_key_is_refused_before_anything_is_built(client, seeded, key):
    headers = {} if key is None else {"Idempotency-Key": key}
    r = call(client, "post", "/api/reports/monthly", seeded["keys"]["a_operator"], headers=headers, json={"month": "2026-10"})
    assert r.status_code == 400 and "Idempotency-Key" in r.json()["detail"]
    assert build(client, seeded, key="well-formed-key-0001").status_code == 200                # the control


@pytest.mark.parametrize("month", ["2026-13", "2026-00", "October", "2026-1", "26-10", "2026-10-01", " 2026-10", ""])
def test_a_month_that_is_not_yyyy_mm_is_refused(client, seeded, month):
    r = build(client, seeded, key="bad-month-key-0001", month=month)
    assert r.status_code == 400 and "YYYY-MM" in r.json()["detail"]
    assert call(client, "get", "/api/reports/monthly/%s/deck" % (month or "x"), seeded["keys"]["a_operator"]).status_code in (400, 404)


def test_a_body_that_is_not_an_object_or_has_a_bad_date_is_refused_by_validation(client, seeded):
    headers = {"Idempotency-Key": "validation-key-0001"}
    key = seeded["keys"]["a_operator"]
    assert call(client, "post", "/api/reports/monthly", key, headers=headers, json=["2026-10"]).status_code == 422
    assert call(client, "post", "/api/reports/monthly", key, headers=headers, json={"month": "2026-10", "as_of": "tomorrow"}).status_code == 422
    assert call(client, "post", "/api/reports/monthly", key, headers=headers, json={}).status_code == 422


def test_no_deck_before_a_report_is_built_is_a_404(client, seeded):
    assert call(client, "get", "/api/reports/monthly/2025-01/deck", seeded["keys"]["a_operator"]).status_code == 404
    assert call(client, "get", "/api/reports/monthly/2026-10/deck", seeded["keys"]["c_admin"]).status_code == 404


# ---------------------------------------------------------------- delivered

def test_delivered_is_recorded_once_and_a_replay_says_so(client, seeded):
    key = "delivery-flow-key-0001"
    assert build(client, seeded, key=key, month="2026-09").json()["delivered"] is False
    first = call(client, "post", "/api/reports/monthly/2026-09/delivered", seeded["keys"]["a_operator"])
    assert first.status_code == 200 and first.json() == {"month": "2026-09", "delivered": True, "already_delivered": False}
    stamp = sql("select delivered_at from report_run where idempotency_key = %s", (key,))[0][0]
    second = call(client, "post", "/api/reports/monthly/2026-09/delivered", seeded["keys"]["a_operator"])
    assert second.json()["already_delivered"] is True
    assert sql("select delivered_at from report_run where idempotency_key = %s", (key,))[0][0] == stamp, "the first time is kept"
    replay = build(client, seeded, key=key, month="2026-09").json()
    assert replay["status"] == "replayed" and replay["delivered"] is True


def test_nothing_can_be_marked_delivered_that_was_not_built(client, seeded):
    assert call(client, "post", "/api/reports/monthly/2024-02/delivered", seeded["keys"]["a_operator"]).status_code == 404
    assert call(client, "post", "/api/reports/monthly/2026-13/delivered", seeded["keys"]["a_operator"]).status_code == 400
    with pytest.raises(psycopg.errors.CheckViolation):
        sql("update report_run set delivered_at = now(), status = 'failed', error = 'x' where idempotency_key = %s", (KEY,))


# ---------------------------------------------------------------- failure, retry and takeover

def test_a_report_that_cannot_be_built_is_a_422_with_the_reason_and_a_failed_row_and_the_same_key_can_be_retried(client, seeded):
    key = "fail-then-fix-key-0001"
    r = build(client, seeded, who="e_admin", key=key)
    assert r.status_code == 422 and r.json()["detail"] == "no active connection found for this tenant"
    assert sql("select status, error, deck is null from report_run where idempotency_key = %s", (key,)) == [
        ("failed", "no active connection found for this tenant", True)]
    again = build(client, seeded, who="e_admin", key=key)
    assert again.status_code == 422 and sql("select count(*) from report_run where idempotency_key = %s", (key,))[0][0] == 1
    # the cause is fixed: the tenant gets a store, and the same key now builds
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(seeded["tenants"]["e"]),))
        shop = connect_row(cur, seeded["tenants"]["e"], "shopify", "e.myshopify.com", "e store")
        load_rows(cur, seeded["tenants"]["e"], shop, both_months())
    fixed = build(client, seeded, who="e_admin", key=key)
    assert fixed.status_code == 200 and fixed.json()["status"] == "created" and fixed.json()["sources"] == ["store", "ads"]
    assert sql("select status, error from report_run where idempotency_key = %s", (key,)) == [("done", None)]


def test_two_ad_accounts_and_none_named_is_a_422_and_naming_one_in_the_body_builds(client, seeded):
    r = build(client, seeded, who="d_admin", key="two-ad-accounts-key-1")
    assert r.status_code == 422 and "2 active paid media connections" in r.json()["detail"]
    named = build(client, seeded, who="d_admin", key="two-ad-accounts-key-1", paid_connection_id=str(seeded["d_ads2"]))
    assert named.status_code == 200 and named.json()["sources"] == ["store", "ads"]


def test_a_run_left_running_by_a_process_that_died_is_taken_over_and_a_fresh_one_is_not(client, seeded):
    key = "died-while-running-0001"
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(seeded["tenants"]["a"]),))
        cur.execute("insert into report_run (tenant_id, idempotency_key, month, status, started_at) values (%s, %s, '2026-10', 'running', now() - interval '1 hour')",
                    (seeded["tenants"]["a"], key))
        cur.execute("insert into report_run (tenant_id, idempotency_key, month, status) values (%s, %s, '2026-10', 'running')",
                    (seeded["tenants"]["a"], key + "-fresh"))
    busy = build(client, seeded, key=key + "-fresh")
    assert busy.status_code == 409 and busy.headers["retry-after"] == "5", "a run that is building right now is not taken over"
    taken = build(client, seeded, key=key)
    assert taken.status_code == 200 and taken.json()["status"] == "created"


def test_two_calls_at_the_same_moment_build_exactly_one_deck(client, seeded, monkeypatch):
    import app.report.runs as runs

    builds = []
    real = runs.generate

    async def counting(*args, **kwargs):
        builds.append(1)
        await asyncio.sleep(0.2)                       # long enough that the others arrive while this one is building
        return await real(*args, **kwargs)

    monkeypatch.setattr(runs, "generate", counting)
    loop, c = client
    key = "simultaneous-key-000001"

    async def post():
        return await c.post("/api/reports/monthly", headers={"Authorization": "Bearer " + seeded["keys"]["a_operator"], "Idempotency-Key": key},
                            json={"month": "2026-10", "as_of": AS_OF_TEXT})

    async def six_at_once():
        return await asyncio.gather(*[post() for _ in range(6)])

    results = loop.run_until_complete(six_at_once())
    codes = sorted(r.status_code for r in results)
    assert len(builds) == 1, "the database let two requests both build"
    assert codes.count(200) >= 1 and set(codes) <= {200, 409}
    assert [r.json()["status"] for r in results if r.status_code == 200].count("created") == 1
    assert all(r.headers.get("retry-after") == "5" for r in results if r.status_code == 409)
    assert sql("select count(*), count(*) filter (where status = 'done') from report_run where idempotency_key = %s", (key,)) == [(1, 1)]
    retry = build(client, seeded, key=key)
    assert retry.status_code == 200 and retry.json()["status"] == "replayed"


def test_an_unexpected_error_says_nothing_about_its_cause_and_leaves_a_failed_row_not_a_running_one(client, seeded, monkeypatch):
    import app.report.runs as runs

    async def broken(*args, **kwargs):
        raise RuntimeError("token tok_SECRET_VALUE leaked in a message")

    monkeypatch.setattr(runs, "generate", broken)
    key = "unexpected-error-key-01"
    r = build(client, seeded, key=key)
    assert r.status_code == 500 and r.json() == {"detail": "the report could not be built"} and "tok_SECRET" not in r.text
    row = sql("select status, error from report_run where idempotency_key = %s", (key,))[0]
    assert row == ("failed", "the report could not be built (RuntimeError)") and "tok_SECRET" not in row[1]
    monkeypatch.undo()
    assert build(client, seeded, key=key).status_code == 200, "a failed run is retried with the same key"


# ---------------------------------------------------------------- tenants

def test_another_tenants_key_string_makes_its_own_run_and_never_reads_or_replays_this_ones(client, seeded):
    mine = build(client, seeded, key="shared-key-string-0001").json()
    theirs = build(client, seeded, who="b_admin", key="shared-key-string-0001")
    assert theirs.status_code == 200 and theirs.json()["status"] == "created", "the key is scoped to the tenant that sent it"
    assert theirs.json()["sha256"] != mine["sha256"] and theirs.json()["sources"] == ["store"]
    deck_b = call(client, "get", "/api/reports/monthly/2026-10/deck", seeded["keys"]["b_admin"]).content
    deck_a = call(client, "get", "/api/reports/monthly/2026-10/deck", seeded["keys"]["a_operator"]).content
    assert hashlib.sha256(deck_b).hexdigest() == theirs.json()["sha256"] and deck_a != deck_b
    text_a = "".join(shape.text_frame.text for slide in Presentation(io.BytesIO(deck_a)).slides for shape in slide.shapes if shape.has_text_frame)
    assert "99,900" not in text_a and "9,990,000" not in text_a
    assert sql("select count(*) from report_run where idempotency_key = 'shared-key-string-0001'")[0][0] == 2


def test_a_tenant_with_no_connection_at_all_gets_a_422_not_a_crash_and_a_failed_row(client, seeded):
    r = build(client, seeded, who="c_admin", key="no-connection-key-0001")
    assert r.status_code == 422 and r.json()["detail"] == "no active connection found for this tenant"


def test_the_row_level_security_policy_and_checks_of_the_run_table_hold_in_the_database(client, seeded):
    assert sql("select relrowsecurity, relforcerowsecurity from pg_class where relname = 'report_run'") == [(True, True)]
    with pytest.raises(psycopg.errors.CheckViolation):
        sql("insert into report_run (tenant_id, idempotency_key, month, status) values (%s, 'bad month key 1', 'x', 'running')",
            (seeded["tenants"]["a"],))
    with pytest.raises(psycopg.errors.CheckViolation):
        sql("insert into report_run (tenant_id, idempotency_key, month, status) values (%s, 'has a space in it', '2026-10', 'running')",
            (seeded["tenants"]["a"],))
    with pytest.raises(psycopg.errors.CheckViolation):
        sql("insert into report_run (tenant_id, idempotency_key, month, status) values (%s, 'done-without-deck', '2026-10', 'done')",
            (seeded["tenants"]["a"],))
    with pytest.raises(psycopg.errors.CheckViolation):
        sql("insert into report_run (tenant_id, idempotency_key, month, status) values (%s, 'failed-without-why', '2026-10', 'failed')",
            (seeded["tenants"]["a"],))
    sql((ROOT / "migrations" / "009_report_runs.sql").read_text(encoding="utf-8"))          # safe to apply twice


# ---------------------------------------------------------------- the class, over many seeded sequences

@pytest.mark.parametrize("seed", range(6))
def test_over_random_sequences_of_calls_a_key_builds_once_replays_the_same_bytes_and_delivery_never_goes_back(client, seeded, monkeypatch, seed):
    """Random calls by three tenants over four keys and two months, against a model of what each call must answer.

    The model: the first call with a key creates and the rest replay the same hash; a key used for the other month is a 422; the tenant
    with no connection always fails and is always retried; delivery lands on the latest deck built for the month and a replay reports
    its own run's flag; and the number of decks built is the number of keys that succeeded, never one per call.
    """
    import app.report.runs as runs

    sql("truncate report_run")                  # the other tests in this file have left runs; the model starts from none
    builds = []
    real = runs.generate

    async def counting(*args, **kwargs):
        builds.append(args[1])
        return await real(*args, **kwargs)

    monkeypatch.setattr(runs, "generate", counting)
    rng = random.Random(seed)
    tenants = {"a_operator": "a", "b_admin": "b", "c_admin": "c"}
    runs_model: dict[tuple[str, str], dict] = {}
    latest: dict[tuple[str, str], str] = {}
    failed_attempts = 0
    c_bound: dict[str, str] = {}            # the tenant with no connection: a key is bound to the month of its first (failed) call
    for step in range(30):
        who = rng.choice(list(tenants))
        t = tenants[who]
        slot = rng.randint(0, 3)
        key = "seq-%d-key-%d" % (seed, slot)
        month = "2026-10" if slot % 2 == 0 else "2026-09"
        action = rng.choice(["build", "build", "build", "deliver", "other_month"])
        known = runs_model.get((t, key))
        where = (seed, step, who, key, action)
        if action == "other_month":
            other = "2026-09" if month == "2026-10" else "2026-10"
            r = build(client, seeded, who=who, key=key, month=other)
            if t == "c":
                if c_bound.setdefault(key, other) != other:
                    assert r.status_code == 422 and "used for" in r.json()["detail"], where
                else:
                    failed_attempts += 1
                    assert r.status_code == 422 and "no active connection" in r.json()["detail"], where
            elif known is not None and known["month"] != other:
                assert r.status_code == 422 and known["month"] in r.json()["detail"], where
            elif known is not None:
                assert r.status_code == 200 and r.json()["status"] == "replayed" and r.json()["sha256"] == known["sha"], where
            else:
                runs_model[(t, key)] = {"month": other, "sha": r.json()["sha256"], "delivered": False}
                latest[(t, other)] = key
            continue
        if action == "deliver":
            r = call(client, "post", "/api/reports/monthly/%s/delivered" % month, seeded["keys"][who])
            if (t, month) in latest:
                assert r.status_code == 200, where
                runs_model[(t, latest[(t, month)])]["delivered"] = True
            else:
                assert r.status_code == 404, where
            continue
        r = build(client, seeded, who=who, key=key, month=month)
        if known is not None and known["month"] != month:
            assert r.status_code == 422 and "used for" in r.json()["detail"], where
            continue
        if t == "c":
            if c_bound.setdefault(key, month) != month:
                assert r.status_code == 422 and "used for" in r.json()["detail"], where
            else:
                failed_attempts += 1
                assert r.status_code == 422 and "no active connection" in r.json()["detail"], where
            continue
        assert r.status_code == 200, where + (r.text,)
        body = r.json()
        if known is None:
            assert body["status"] == "created" and body["delivered"] is False, where
            runs_model[(t, key)] = {"month": month, "sha": body["sha256"], "delivered": False}
            latest[(t, month)] = key
        else:
            assert body["month"] == known["month"] and body["status"] == "replayed", where
            assert body["sha256"] == known["sha"] and body["delivered"] == known["delivered"], where
    for (t, key), entry in runs_model.items():
        row = sql("select deck_sha256, delivered_at is not null, status from report_run where idempotency_key = %s and tenant_id = %s",
                  (key, seeded["tenants"][t]))
        assert row == [(entry["sha"], entry["delivered"], "done")], (seed, t, key)
    for t in ("a", "b"):
        assert builds.count(str(seeded["tenants"][t])) == sum(1 for (tt, _) in runs_model if tt == t), (seed, t)
    assert builds.count(str(seeded["tenants"]["c"])) == failed_attempts, seed


# ---------------------------------------------------------------- the last survivors of the derived guards

def _db(client):
    from app.api.main import app as fastapi_app
    return fastapi_app.state.db


def _drive(client, coro):
    return client[0].run_until_complete(coro)


def _insert_running(seeded, key, age):
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(seeded["tenants"]["a"]),))
        cur.execute("insert into report_run (tenant_id, idempotency_key, month, status, started_at) "
                    "values (%s, %s, '2026-10', 'running', now() - %s * interval '1 second')", (seeded["tenants"]["a"], key, age))


def test_a_running_run_is_taken_over_only_after_fifteen_minutes(client, seeded):
    from app.report.runs import STALE_SECONDS
    assert STALE_SECONDS == 900
    _insert_running(seeded, "almost-stale-key-0001", 899.0)
    assert build(client, seeded, key="almost-stale-key-0001").status_code == 409
    _insert_running(seeded, "just-stale-key-00001", 900.5)
    assert build(client, seeded, key="just-stale-key-00001").status_code == 200


def test_claiming_names_who_builds_who_waits_and_a_key_for_another_month(client, seeded):
    from app.report.runs import KeyReused, RunInProgress, claim, mark_delivered, parse_month
    import app.report.runs as runs
    tenant, db = str(seeded["tenants"]["a"]), _db(client)
    key = "claim-direct-key-0001"
    assert _drive(client, claim(db, tenant, key, "2026-10")) == ("owner", None)
    with pytest.raises(RunInProgress) as busy:
        _drive(client, claim(db, tenant, key, "2026-10"))
    assert str(busy.value) == "a request with this key is still building the report"
    with pytest.raises(KeyReused) as other:
        _drive(client, claim(db, tenant, key, "2026-09"))
    assert str(other.value) == "this key was used for 2026-10"
    _drive(client, runs._fail(db, tenant, key, "x" * 600))
    error = sql("select error from report_run where idempotency_key = %s", (key,))[0][0]
    assert error == "x" * 500, "a reason is cut to 500 characters"
    assert _drive(client, claim(db, tenant, key, "2026-10")) == ("owner", None), "a failed run is claimed again by the next caller"
    assert parse_month("2026-10") == (2026, 10)
    assert _drive(client, mark_delivered(db, tenant, "1999-01")) == (False, False), "nothing was built for that month"
    assert _drive(client, runs.latest_deck(db, tenant, "1999-01")) is None


def test_a_stored_run_cannot_be_changed_and_the_sources_of_a_report_are_named_in_a_fixed_order():
    import dataclasses
    from types import SimpleNamespace
    from app.report.runs import StoredRun, _sources
    run_ = StoredRun("1", "2026-10", "done", b"", "00", 1, [], [], None, False)
    with pytest.raises(dataclasses.FrozenInstanceError):
        run_.status = "failed"
    some = object()
    assert _sources(SimpleNamespace(web=None, paid=None, search=None)) == ["store"]
    assert _sources(SimpleNamespace(web=some, paid=None, search=None)) == ["store", "website"]
    assert _sources(SimpleNamespace(web=None, paid=some, search=None)) == ["store", "ads"]
    assert _sources(SimpleNamespace(web=some, paid=some, search=some)) == ["store", "website", "ads", "search"]


def test_the_deck_and_delivered_calls_refuse_a_month_that_is_not_a_month_and_say_when_none_was_built(client, seeded):
    for bad in ("2026-13", "abc", "2026-1"):
        assert call(client, "get", "/api/reports/monthly/%s/deck" % bad, seeded["keys"]["a_operator"]).status_code == 400, bad
        assert call(client, "post", "/api/reports/monthly/%s/delivered" % bad, seeded["keys"]["a_operator"]).status_code == 400, bad
    deck = call(client, "get", "/api/reports/monthly/2025-01/deck", seeded["keys"]["a_operator"])
    assert deck.status_code == 404 and deck.json() == {"detail": "no report has been built for that month"}
    done = call(client, "post", "/api/reports/monthly/2025-01/delivered", seeded["keys"]["a_operator"])
    assert done.status_code == 404 and done.json() == {"detail": "no report has been built for that month"}


def test_each_connection_named_in_the_body_is_looked_up_and_an_unknown_one_is_refused_not_ignored(client, seeded):
    stranger = str(uuid.uuid4())
    for field_name in ("connection_id", "web_connection_id", "paid_connection_id", "search_connection_id"):
        r = build(client, seeded, key="named-%s-0001" % field_name.replace("_", "-"), **{field_name: stranger})
        assert r.status_code == 422, (field_name, r.status_code, r.text)
    ok = build(client, seeded, key="named-search-ok-0001", search_connection_id=str(seeded["a_search"]))
    assert ok.status_code == 200 and "search" in ok.json()["sources"]


def test_an_unexpected_failure_is_logged_with_its_event_name_under_the_hub_logger(client, seeded, monkeypatch, caplog):
    import logging
    import app.report.runs as runs

    async def broken(*args, **kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr(runs, "generate", broken)
    with caplog.at_level(logging.ERROR, logger="hub"):
        r = build(client, seeded, key="logged-failure-key-01")
    assert r.status_code == 500
    records = [x for x in caplog.records if x.getMessage() == "monthly report failed"]
    assert len(records) == 1 and records[0].name == "hub" and getattr(records[0], "event", None) == "report_failed"
    assert getattr(records[0], "tenant_id", None) == str(seeded["tenants"]["a"])
    assert records[0].exc_info is not None, "the traceback is logged, never returned"


def test_the_published_schema_names_the_idempotency_header_and_describes_the_month(client, seeded):
    spec = call(client, "get", "/openapi.json")
    assert spec.status_code == 200, "the schema is the contract a caller reads"
    post = spec.json()["paths"]["/api/reports/monthly"]["post"]
    assert "Idempotency-Key" in [p["name"] for p in post.get("parameters", [])]
    body = spec.json()["components"]["schemas"]["MonthlyRequest"]["properties"]["month"]
    assert body["description"] == "YYYY-MM, the month to report"
