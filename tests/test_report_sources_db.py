"""Ads and search through the real sync runner into real Postgres, out through the report and the commands.

The whole path against the stand-in servers (never a live ad account or property): MetaConnector and GSCConnector over
HTTP, the SyncRunner as the sync_worker role, the stored rows, then the report as the request role with row level
security. It proves what only a database can: a tenant's ads and search rows reach only its own deck, a failed run is
not coverage for an ad or a search day, the report refuses two ad accounts or two properties unless one is named, a
store is still required, a month in two ad currencies is refused, and the three commands print what ran and never a token.

Skips cleanly when Postgres is not running (`docker compose up -d`). Like the other database suites it drops and
recreates the `public` schema of the database it points at.
"""
from __future__ import annotations

import asyncio
import os
import subprocess
import sys
import uuid
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

psycopg = pytest.importorskip("psycopg", reason="psycopg not installed")
pytest.importorskip("asyncpg")

DSN = os.environ.get("ANALYTICS_TEST_DSN", "postgresql://analytics:analytics@localhost:5433/analytics")


def _reachable() -> bool:
    try:
        with psycopg.connect(DSN, connect_timeout=3):
            return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(not _reachable(), reason="Postgres not reachable, run docker compose up -d")

import httpx  # noqa: E402
from pptx import Presentation  # noqa: E402

from app.connectors.gsc import GSCConnector  # noqa: E402
from app.connectors.gsc_fake import FakeGSC  # noqa: E402
from app.connectors.meta import MetaConnector  # noqa: E402
from app.connectors.meta_fake import FakeMeta  # noqa: E402
from app.db import Database  # noqa: E402
from app.report.deck import dump_deck  # noqa: E402
from app.report.monthly import ReportError  # noqa: E402
from app.report.service import generate, load_inputs  # noqa: E402
from app.sync.runner import SyncRunner  # noqa: E402
from report_fixtures import AS_OF, both_months  # noqa: E402
from test_report_service_db import _seed  # noqa: E402

TOKEN = "tok-sources-db-SECRET"
AD_ACCOUNT, OTHER_AD_ACCOUNT, SECOND_AD_ACCOUNT = "880001", "880002", "880003"
SITE, OTHER_SITE, SECOND_SITE = "sc-domain:a.example", "sc-domain:b.example", "sc-domain:c.example"
MIGRATIONS = ("001_init.sql", "002_control_plane.sql", "006_ga4_metrics.sql", "007_meta_metrics.sql", "008_search_metrics.sql")

ADS = {
    date(2026, 9, 1): {"spend": "100.00", "impressions": 1000, "clicks": 50, "purchases": 5, "purchase_value": "400.00"},
    date(2026, 9, 2): {"spend": "50.00", "impressions": 500, "clicks": 20, "purchases": 2, "purchase_value": "100.00"},
    date(2026, 10, 1): {"spend": "120.00", "impressions": 1200, "clicks": 60, "purchases": 6, "purchase_value": "540.00"},
    date(2026, 10, 2): {"spend": "60.00", "impressions": 400, "clicks": 10, "purchases": None, "purchase_value": None},
}
SEARCH = {date(2026, 9, 1): (10, 100, 7.5), date(2026, 9, 2): (20, 200, 6.0), date(2026, 10, 1): (12, 150, 5.0), date(2026, 10, 2): (8, 50, 9.0)}
QUERIES = {date(2026, 10, 1): [("shoes", 6, 60, 5.0), ("boots", 4, 40, 10.0)], date(2026, 10, 2): [("shoes", 3, 20, 7.0), ("socks", 5, 30, 4.0)]}
PAGES = {date(2026, 10, 1): [("/a", 10, 100, 7.0)], date(2026, 10, 2): [("/a", 4, 40, 6.0), ("/b", 4, 10, 2.0)]}
#: The runs the world makes for ads and search: Oct 10 and 11 are covered by none (a failed run touches them, which is not coverage).
WINDOWS = ((date(2026, 9, 1), date(2026, 10, 10)), (date(2026, 10, 12), date(2026, 11, 1)))
TITLE_ADS = ("October ad spend total is 20.0% above September, but October is missing 2 days of ad data: per reported day it is up 24.1%")


def run(coro):
    return asyncio.run(coro)


@pytest.fixture(scope="module")
def servers():
    with FakeMeta(token=TOKEN) as meta, FakeGSC(token=TOKEN, today=AS_OF) as gsc:
        yield meta, gsc


def connect_row(cur, tenant, platform, account, name):
    cur.execute(
        "insert into platform_connection (tenant_id, platform, external_account, display_name, credentials_enc) "
        "values (%s, %s, %s, %s, %s) returning id", (tenant, platform, account, name, b"ciphertext"))
    return cur.fetchone()[0]


@pytest.fixture(scope="module")
def world(servers):
    """A has a store, ads and search; B only ads and search with huge figures; C a store; D two ad accounts and two properties; E ads
    only; G a store whose ad rows name two currencies in October; H nothing at all."""
    meta, gsc = servers
    a, b, c, d, e, g, h = (uuid.uuid4() for _ in range(7))
    ids = {}
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("drop schema if exists public cascade; create schema public")
        for name in MIGRATIONS:
            cur.execute((ROOT / "migrations" / name).read_text(encoding="utf-8"))
        cur.execute("grant usage on schema public to sync_worker")
        ids["a_shop"] = _seed(cur, a, "tenant-a", both_months())
        ids["c_shop"] = _seed(cur, c, "tenant-c", both_months())
        ids["d_shop"] = _seed(cur, d, "tenant-d", both_months())
        ids["g_shop"] = _seed(cur, g, "tenant-g", both_months())
        for tenant, label in ((b, "tenant-b"), (e, "tenant-e"), (h, "tenant-h")):
            cur.execute("select set_config('app.tenant_id', %s, false)", (str(tenant),))
            cur.execute("insert into tenant (id, slug, name) values (%s, %s, %s)", (tenant, label, label))
        for key, tenant, platform, account, name in (
            ("a_ads", a, "meta", "act_" + AD_ACCOUNT, "a ads"), ("a_search", a, "gsc", SITE, "a search"),
            ("b_ads", b, "meta", "act_" + OTHER_AD_ACCOUNT, "b ads"), ("b_search", b, "gsc", OTHER_SITE, "b search"),
            ("d_ads1", d, "meta", "act_1", "d ads one"), ("d_ads2", d, "meta", "act_2", "d ads two"),
            ("d_search1", d, "gsc", "sc-domain:one.example", "d search one"), ("d_search2", d, "gsc", "sc-domain:two.example", "d search two"),
            ("e_ads", e, "meta", "act_3", "e ads"), ("g_ads", g, "meta", "act_4", "g ads"),
        ):
            cur.execute("select set_config('app.tenant_id', %s, false)", (str(tenant),))
            ids[key] = connect_row(cur, tenant, platform, account, name)
        # G's October ad rows name two currencies. The worker connection is a superuser here, so only the schema's own checks apply.
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(g),))
        for day, currency in ((date(2026, 10, 1), "EUR"), (date(2026, 10, 2), "GBP")):
            cur.execute("insert into metric_snapshot (tenant_id, connection_id, metric_key, observed_on, value_numeric, currency) "
                        "values (%s, %s, 'ad_spend', %s, 1000, %s)", (g, ids["g_ads"], day, currency))
    meta.set_account(AD_ACCOUNT, ADS, currency="EUR")
    meta.set_account(OTHER_AD_ACCOUNT, {date(2026, 10, 1): {"spend": "99999.00", "impressions": 99999999, "clicks": 9999999, "purchases": 9999,
                                                            "purchase_value": "99999.00"}}, currency="USD")
    gsc.set_site(SITE, SEARCH, queries=QUERIES, pages=PAGES)
    gsc.set_site(OTHER_SITE, {date(2026, 10, 1): (9999999, 99999999, 1.0)}, queries={date(2026, 10, 1): [("secret competitor query", 9999999, 99999999, 1.0)]})
    full = {"a": a, "b": b, "c": c, "d": d, "e": e, "g": g, "h": h, **ids}
    for start, end in WINDOWS:
        assert run(sync_ads(full, "a_ads", "a", meta, start, end)).status == "succeeded"
        assert run(sync_search(full, "a_search", "a", gsc, start, end)).status == "succeeded"
    # a failed run over the hole, for each source: it must not count as coverage
    assert run(sync_ads(full, "a_ads", "a", meta, date(2026, 10, 10), date(2026, 10, 12), token="wrong")).status == "failed"
    assert run(sync_search(full, "a_search", "a", gsc, date(2026, 10, 10), date(2026, 10, 12), token="wrong")).status == "failed"
    run(sync_ads(full, "b_ads", "b", meta, date(2026, 10, 1), date(2026, 10, 2), account=OTHER_AD_ACCOUNT))
    run(sync_search(full, "b_search", "b", gsc, date(2026, 10, 1), date(2026, 10, 2), site=OTHER_SITE))
    return full


def _client():
    return httpx.AsyncClient(trust_env=False, verify=False, timeout=10)


async def _sync(world, key, tenant_key, connector, start, end):
    db = Database(DSN, max_size=2, role="sync_worker")
    await db.connect()
    try:
        return await SyncRunner(db).run_once(tenant_id=str(world[tenant_key]), connection_id=str(world[key]), connector=connector,
                                             window_start=start, window_end=end, trigger="backfill")
    finally:
        await connector.aclose()
        await db.close()


async def sync_ads(world, key, tenant_key, server, start, end, account=AD_ACCOUNT, token=TOKEN):
    async def no_sleep(seconds):
        return None
    return await _sync(world, key, tenant_key, MetaConnector(ad_account_id=account, access_token=token, base_url=server.url,
                                                              client=_client(), sleep=no_sleep), start, end)


async def sync_search(world, key, tenant_key, server, start, end, site=SITE, token=TOKEN):
    async def no_sleep(seconds):
        return None
    return await _sync(world, key, tenant_key, GSCConnector(site_url=site, access_token=token, base_url=server.url,
                                                             client=_client(), sleep=no_sleep), start, end)


async def _inputs(world, tenant_key, **kw):
    db = Database(DSN, max_size=2)
    await db.connect()
    try:
        return await load_inputs(db, str(world[tenant_key]), 2026, 10, **kw)
    finally:
        await db.close()


async def _generate(world, tenant_key, tmp_path, **kw):
    db = Database(DSN, max_size=2)
    await db.connect()
    try:
        return await generate(db, str(world[tenant_key]), 2026, 10, tmp_path / "deck.pptx", as_of=AS_OF, **kw)
    finally:
        await db.close()


# ---------------------------------------------------------------- from the stand-in servers to the deck

def test_from_the_stand_ins_to_the_deck_the_figures_are_the_ones_computed_by_hand(world, tmp_path):
    path, report, _ = run(_generate(world, "a", tmp_path, connection_id=str(world["a_shop"])))
    prs = Presentation(str(path))
    assert len(prs.slides) == 9
    assert prs.slides[5].shapes.title.text_frame.text == TITLE_ADS
    assert report.paid.kpi("ad_spend").current_display == "EUR 180.00" and report.paid.kpi("roas").current_display == "4.50"
    assert report.paid.kpi("ctr").current_display == "4.38%" and report.paid.kpi("cpc").current_display == "EUR 2.57"
    assert report.paid.current.missing_days == [date(2026, 10, 10), date(2026, 10, 11)]
    assert report.search.kpi("search_clicks").current_display == "20" and report.search.kpi("position").current_display == "6.0"
    assert [(r.key, r.clicks) for r in report.search.top_queries] == [("shoes", 9), ("socks", 5), ("boots", 4)]
    assert [(r.key, r.clicks) for r in report.search.top_pages] == [("/a", 14), ("/b", 4)]
    assert report.search.current.missing_days == [date(2026, 10, 10), date(2026, 10, 11)]
    assert report.kpi("net_sales").current_display == "$210.00", "the store half is the store's own, in dollars beside the ads' euros"
    assert report.currency == "USD" and report.paid.currency == "EUR"


def test_another_tenants_ads_and_search_never_reach_this_tenants_report(world, tmp_path):
    path, _, _ = run(_generate(world, "a", tmp_path))
    text = dump_deck(path)
    for leaked in ("99,999", "9,999,999", "99999999", "secret competitor", "USD"):
        assert leaked not in text, leaked
    inputs = run(_inputs(world, "a"))
    assert {r["value_numeric"] for r in inputs.paid.rows if r["metric_key"] == "ad_clicks"} == {50, 20, 60, 10}
    assert "secret competitor query" not in {r["key"] for r in inputs.top_rows}


def test_a_failed_ad_or_search_run_over_a_hole_is_not_coverage_so_those_days_stay_missing(world):
    inputs = run(_inputs(world, "a"))
    for part in (inputs.paid, inputs.search):
        assert (date(2026, 10, 10), date(2026, 10, 12)) not in part.covered
        assert (date(2026, 10, 12), date(2026, 11, 1)) in part.covered
    failed = None
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("select count(*) from sync_run where connection_id in (%s, %s) and status = 'failed' and error_code = 'auth_expired'",
                    (world["a_ads"], world["a_search"]))
        failed = cur.fetchone()[0]
    assert failed == 2


def test_the_inputs_hold_the_top_rows_of_the_month_only_and_none_for_a_tenant_without_search(world):
    inputs = run(_inputs(world, "a"))
    assert len(inputs.top_rows) == 7 and {r["observed_on"] for r in inputs.top_rows} == {date(2026, 10, 1), date(2026, 10, 2)}
    assert all(r["position_milli"] > 0 for r in inputs.top_rows)
    store_only = run(_inputs(world, "c"))
    assert store_only.paid is None and store_only.search is None and store_only.top_rows is None


def test_a_store_only_tenant_gets_the_seven_slide_deck_exactly_as_before(world, tmp_path):
    path, report, _ = run(_generate(world, "c", tmp_path))
    assert report.paid is None and report.search is None and len(Presentation(str(path)).slides) == 7


# ---------------------------------------------------------------- the refusals

def test_two_ad_accounts_or_two_properties_with_no_name_are_refused_and_naming_one_works(world):
    with pytest.raises(ReportError, match="2 active paid media connections, name one with --paid-connection"):
        run(_inputs(world, "d"))
    with pytest.raises(ReportError, match="2 active search connections, name one with --search-connection"):
        run(_inputs(world, "d", paid_connection_id=str(world["d_ads1"])))
    inputs = run(_inputs(world, "d", paid_connection_id=str(world["d_ads1"]), search_connection_id=str(world["d_search2"])))
    assert inputs.paid.label == "d ads one" and inputs.search.label == "d search two" and inputs.label == "tenant-d store"


def test_an_id_from_another_tenant_or_another_kind_of_connection_is_not_found(world):
    with pytest.raises(ReportError, match="no active paid media connection found with that id"):
        run(_inputs(world, "a", paid_connection_id=str(world["b_ads"])))
    with pytest.raises(ReportError, match="no active paid media connection found with that id"):
        run(_inputs(world, "a", paid_connection_id=str(world["a_shop"])))
    with pytest.raises(ReportError, match="no active search connection found with that id"):
        run(_inputs(world, "a", search_connection_id=str(world["a_ads"])))
    with pytest.raises(ReportError, match="no active connection"):
        run(_inputs(world, "a", connection_id=str(world["a_ads"])))


def test_a_tenant_with_ads_but_no_store_or_with_no_connection_at_all_gets_no_report(world):
    for tenant in ("e", "h"):
        with pytest.raises(ReportError, match="no active connection found for this tenant"):
            run(_inputs(world, tenant))


def test_a_month_of_ad_rows_in_two_currencies_is_refused_not_converted(world, tmp_path):
    with pytest.raises(ReportError, match="more than one currency in October 2026: EUR, GBP"):
        run(_generate(world, "g", tmp_path))


# ---------------------------------------------------------------- the commands

def _cli(args, extra_env=None):
    env = dict(os.environ, DATABASE_URL=DSN, **(extra_env or {}))
    for name in ("GROQ_API_KEY", "LLM_API_KEY", "META_ACCESS_TOKEN", "META_AD_ACCOUNT_ID", "META_API_BASE_URL", "GSC_ACCESS_TOKEN",
                 "GSC_SITE_URL", "GSC_API_BASE_URL"):
        if not extra_env or name not in extra_env:
            env.pop(name, None)
    return subprocess.run([sys.executable, str(ROOT / "scripts" / args[0]), *args[1:]], capture_output=True, text=True,
                          encoding="utf-8", timeout=90, env=env, cwd=ROOT)


def test_the_report_command_prints_that_ads_and_search_are_in_the_deck_and_refuses_a_wrong_id(world, tmp_path):
    done = _cli(["monthly_report.py", "--month", "2026-10", "--tenant", str(world["a"]), "--connection", str(world["a_shop"]),
                 "--paid-connection", str(world["a_ads"]), "--search-connection", str(world["a_search"]),
                 "--out", str(tmp_path / "cli.pptx"), "--as-of", AS_OF.isoformat(), "--dump"])
    assert done.returncode == 0, done.stderr
    assert "(October 2026 with ad, search figures, insights: template)" in done.stdout
    assert TITLE_ADS in done.stdout and "Top queries" in done.stdout
    bad = _cli(["monthly_report.py", "--month", "2026-10", "--tenant", str(world["a"]), "--paid-connection", str(world["a_shop"]),
                "--out", str(tmp_path / "x.pptx"), "--as-of", AS_OF.isoformat()])
    assert bad.returncode == 2 and "no report: no active paid media connection found with that id" in bad.stderr
    two = _cli(["monthly_report.py", "--month", "2026-10", "--tenant", str(world["g"]), "--out", str(tmp_path / "y.pptx"),
                "--as-of", AS_OF.isoformat()])
    assert two.returncode == 2 and "more than one currency" in two.stderr and "Traceback" not in two.stderr


def test_the_ad_sync_command_runs_against_the_stand_in_and_never_prints_the_token(world, servers):
    meta, _ = servers
    env = {"META_AD_ACCOUNT_ID": AD_ACCOUNT, "META_ACCESS_TOKEN": TOKEN, "META_API_BASE_URL": meta.url}
    done = _cli(["sync_meta.py", "--tenant", str(world["a"]), "--connection", str(world["a_ads"]), "--days", "3", "--end", "2026-09-04"], env)
    assert done.returncode == 0, done.stderr + done.stdout
    assert "sync succeeded: window 2026-09-01 to 2026-09-04, 10 rows written, 1 api calls" in done.stdout
    assert TOKEN not in done.stdout + done.stderr
    wrong = _cli(["sync_meta.py", "--tenant", str(world["b"]), "--connection", str(world["a_ads"]), "--days", "3", "--end", "2026-09-04"], env)
    assert wrong.returncode == 2 and "no active meta connection with that id" in wrong.stderr
    store = _cli(["sync_meta.py", "--tenant", str(world["a"]), "--connection", str(world["a_shop"]), "--days", "3"], env)
    assert store.returncode == 2
    missing = _cli(["sync_meta.py", "--tenant", str(world["a"]), "--connection", str(world["a_ads"])], {"META_AD_ACCOUNT_ID": AD_ACCOUNT})
    assert missing.returncode == 2 and "META_ACCESS_TOKEN" in missing.stderr and "Traceback" not in missing.stderr
    failed = _cli(["sync_meta.py", "--tenant", str(world["a"]), "--connection", str(world["a_ads"]), "--days", "3", "--end", "2026-09-04"],
                  {**env, "META_ACCESS_TOKEN": "wrong-one"})
    assert failed.returncode == 1 and "error auth_expired" in failed.stdout and "wrong-one" not in failed.stdout + failed.stderr


def test_the_search_sync_command_runs_against_the_stand_in_and_never_prints_the_token(world, servers):
    _, gsc = servers
    env = {"GSC_SITE_URL": SITE, "GSC_ACCESS_TOKEN": TOKEN, "GSC_API_BASE_URL": gsc.url}
    done = _cli(["sync_gsc.py", "--tenant", str(world["a"]), "--connection", str(world["a_search"]), "--days", "3", "--end", "2026-09-04"], env)
    assert done.returncode == 0, done.stderr + done.stdout
    assert "sync succeeded: window 2026-09-01 to 2026-09-04, 6 rows written, 6 api calls" in done.stdout
    assert TOKEN not in done.stdout + done.stderr
    none = _cli(["sync_gsc.py", "--tenant", str(world["a"]), "--connection", str(world["a_search"]), "--days", "3", "--end", "2026-09-04",
                 "--top", "0"], env)
    assert none.returncode == 0 and "2 api calls" in none.stdout, "no top rows asked for, so only the date query and its empty page"
    wrong = _cli(["sync_gsc.py", "--tenant", str(world["a"]), "--connection", str(world["a_ads"]), "--days", "3"], env)
    assert wrong.returncode == 2 and "no active gsc connection with that id" in wrong.stderr
    bad_top = _cli(["sync_gsc.py", "--tenant", str(world["a"]), "--connection", str(world["a_search"]), "--top", "99999"], env)
    assert bad_top.returncode == 2
    failed = _cli(["sync_gsc.py", "--tenant", str(world["a"]), "--connection", str(world["a_search"]), "--days", "3", "--end", "2026-09-04"],
                  {**env, "GSC_ACCESS_TOKEN": "wrong-one"})
    assert failed.returncode == 1 and "error auth_expired" in failed.stdout and "wrong-one" not in failed.stdout + failed.stderr


# ---------------------------------------------------------------- the last survivors of the derived guards

def test_two_stores_and_no_name_are_refused_with_the_count_and_the_flag_and_naming_one_works(world):
    with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
        cur.execute("select set_config('app.tenant_id', %s, false)", (str(world["c"]),))
        second = connect_row(cur, world["c"], "shopify", "second-store.example", "c second store")
    try:
        with pytest.raises(ReportError) as exc:
            run(_inputs(world, "c"))
        assert str(exc.value) == "this tenant has 2 active connections, name one with --connection"
        named = run(_inputs(world, "c", connection_id=str(world["c_shop"])))
        assert named.label == "tenant-c store"
    finally:
        with psycopg.connect(DSN, autocommit=True) as conn, conn.cursor() as cur:
            cur.execute("select set_config('app.tenant_id', %s, false)", (str(world["c"]),))
            cur.execute("delete from platform_connection where id = %s", (second,))


def test_a_deck_that_cannot_be_written_says_where_and_why_and_names_the_error_type_only(world, tmp_path):
    (tmp_path / "a-file").write_text("not a folder", encoding="utf-8")
    missing = tmp_path / "a-file" / "deck.pptx"          # a folder that is really a file: the deck cannot go there

    async def go():
        db = Database(DSN, max_size=2)
        await db.connect()
        try:
            return await generate(db, str(world["a"]), 2026, 10, missing, as_of=AS_OF)
        finally:
            await db.close()
    with pytest.raises(ReportError) as exc:
        run(go())
    text = str(exc.value)
    assert text.startswith("the deck could not be saved to %s (" % missing) and text.endswith("); close it if it is open")
    assert "Error)" in text and "Traceback" not in text


# ---------------------------------------------------------------- the shared sync command body, argument by argument

def _gsc_env(gsc):
    return {"GSC_SITE_URL": SITE, "GSC_ACCESS_TOKEN": TOKEN, "GSC_API_BASE_URL": gsc.url}


def test_the_sync_commands_explain_their_arguments_and_need_a_tenant_and_a_connection(world):
    helped = _cli(["sync_gsc.py", "--help"])
    assert helped.returncode == 0
    words = " ".join(helped.stdout.split())
    assert "--tenant TENANT tenant uuid" in words and "--connection CONNECTION the gsc connection uuid" in words
    assert "how many days back from the end (default 30)" in words and "first day NOT pulled, YYYY-MM-DD" in words
    no_tenant = _cli(["sync_gsc.py", "--connection", str(world["a_search"])])
    assert no_tenant.returncode == 2 and "--tenant" in no_tenant.stderr and "required" in no_tenant.stderr
    no_connection = _cli(["sync_gsc.py", "--tenant", str(world["a"])])
    assert no_connection.returncode == 2 and "--connection" in no_connection.stderr and "required" in no_connection.stderr


def test_the_window_defaults_to_thirty_days_and_days_must_be_one_to_3650(world, servers):
    _, gsc = servers
    base = ["sync_gsc.py", "--tenant", str(world["a"]), "--connection", str(world["a_search"]), "--end", "2026-09-04", "--top", "0"]
    default = _cli(base, _gsc_env(gsc))
    assert default.returncode == 0 and "window 2026-08-05 to 2026-09-04" in default.stdout, default.stdout + default.stderr
    for bad in ("0", "3651", "-5"):
        refused = _cli(base + ["--days", bad], _gsc_env(gsc))
        assert refused.returncode == 2 and "--days must be 1 to 3650" in refused.stderr, bad
    for edge, start in (("1", "2026-09-03"), ("3650", "2016-09-06")):
        ok = _cli(base + ["--days", edge], _gsc_env(gsc))
        assert ok.returncode == 0 and "window %s to 2026-09-04" % start in ok.stdout, (edge, ok.stdout + ok.stderr)


def test_without_a_database_url_or_with_an_unreachable_database_the_command_says_so_and_exits_two(world, servers):
    _, gsc = servers
    args = ["sync_gsc.py", "--tenant", str(world["a"]), "--connection", str(world["a_search"]), "--days", "3", "--end", "2026-09-04"]
    env = dict(os.environ, **_gsc_env(gsc))
    env.pop("DATABASE_URL", None)
    nodb = subprocess.run([sys.executable, str(ROOT / "scripts" / args[0]), *args[1:]], capture_output=True, text=True,
                          encoding="utf-8", timeout=90, env=env, cwd=ROOT)
    assert nodb.returncode == 2 and nodb.stderr.startswith("no sync: ") and "DATABASE_URL" in nodb.stderr and "Traceback" not in nodb.stderr
    env["DATABASE_URL"] = "postgresql://analytics:analytics@127.0.0.1:1/analytics"
    down = subprocess.run([sys.executable, str(ROOT / "scripts" / args[0]), *args[1:]], capture_output=True, text=True,
                          encoding="utf-8", timeout=90, env=env, cwd=ROOT)
    assert down.returncode == 2 and "no sync: the database could not be used (" in down.stderr
    assert "Traceback" not in down.stderr and "127.0.0.1" not in down.stderr and "analytics:analytics" not in down.stderr


@pytest.mark.parametrize("outcome_kwargs,code,words", [
    ({"claimed": False}, 1, "another sync holds this connection, nothing was run"),
    ({"claimed": True, "status": "partial", "rows_written": 4, "api_calls": 2, "throttle_waits": 1}, 0, "sync partial: window 2026-09-01 to 2026-09-04, 4 rows written, 2 api calls, 1 waits"),
    ({"claimed": True, "status": "succeeded"}, 0, "sync succeeded: window"),
    ({"claimed": True, "status": "failed", "error_code": "auth_expired", "error_detail": "the token ended"}, 1, "error auth_expired: the token ended"),
])
def test_the_exit_code_and_the_line_printed_say_what_the_run_did(world, servers, monkeypatch, capsys, outcome_kwargs, code, words):
    import argparse
    from datetime import date as _date
    import app.sync.cli as cli
    from app.sync.runner import RunOutcome

    class Closes:
        closed = False

        async def aclose(self):
            Closes.closed = True

    class FakeRunner:
        def __init__(self, db):
            pass

        async def run_once(self, **kw):
            return RunOutcome("run-1", **outcome_kwargs)
    monkeypatch.setenv("DATABASE_URL", DSN)
    monkeypatch.setattr(cli, "SyncRunner", FakeRunner)
    args = argparse.Namespace(tenant=str(world["a"]), connection=str(world["a_search"]), days=3, end=_date(2026, 9, 4))
    got = run(cli.run_connection_sync(args, platform="gsc", connector=Closes(), default_end=_date(2026, 1, 1)))
    assert got == code and Closes.closed, "the connector is closed whatever happened"
    assert words in capsys.readouterr().out
